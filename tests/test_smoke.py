from __future__ import annotations

import ast
import importlib
import json
import subprocess
import sys
import tempfile
import threading
import unittest
import os
import shutil
from contextlib import ExitStack
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from unittest import mock


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def package_dir() -> Path:
    candidates = [path for path in PROJECT_ROOT.iterdir() if path.is_dir() and path.name.startswith("WinPython_")]
    if len(candidates) != 1:
        raise AssertionError(f"expected one WinPython package directory, found {len(candidates)}")
    return candidates[0]


def run_powershell_contract(source: str) -> subprocess.CompletedProcess[str]:
    system_powershell = (
        Path(os.environ.get("SystemRoot", "C:/Windows"))
        / "System32" / "WindowsPowerShell" / "v1.0" / "powershell.exe"
    )
    process_env = os.environ.copy()
    if system_powershell.is_file():
        process_env["PSModulePath"] = os.pathsep.join(
            [str(system_powershell.parent / "Modules"), process_env.get("PSModulePath", "")]
        )
    with tempfile.TemporaryDirectory() as temp_dir:
        contract_path = Path(temp_dir) / "contract.ps1"
        contract_path.write_text(source, encoding="utf-8-sig")
        contract_literal = str(contract_path).replace("'", "''")
        return subprocess.run(
            [
                str(system_powershell) if system_powershell.is_file() else "powershell",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-Command",
                "[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding($false); "
                "$OutputEncoding = [Console]::OutputEncoding; "
                f"& '{contract_literal}'",
            ],
            cwd=package_dir(),
            env=process_env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
            check=False,
        )


def rest_time_module():
    root = package_dir()
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    return importlib.import_module("rest_time_automation")


def duty_rehearsal_module():
    root = package_dir()
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    return importlib.import_module("duty_rehearsal")


def legacy_tk_path(*parts: str) -> Path:
    return package_dir() / "legacy_tk" / Path(*parts)


def legacy_tk_module(name: str):
    root = package_dir()
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    return importlib.import_module(f"legacy_tk.{name}")


def duty_gui_module():
    return legacy_tk_module("duty_gui")


def package_module(name: str):
    root = package_dir()
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    return importlib.import_module(name)


def legacy_duty_sheet_module():
    legacy_dir = package_dir() / "duty_sheet_legacy"
    if str(legacy_dir) not in sys.path:
        sys.path.insert(0, str(legacy_dir))
    spec = importlib.util.spec_from_file_location("legacy_sinposmart_1", legacy_dir / "sinposmart_1.py")
    if spec is None or spec.loader is None:
        raise AssertionError("cannot load legacy duty sheet module")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class DailyVehicleVerificationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        script = package_dir() / "daily_vehicle_legacy" / "automation" / "ppe_selenium_daily.py"
        spec = importlib.util.spec_from_file_location("_vehicle_verification_smoke", script)
        cls.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.module)

    def test_empty_grid_is_unconfirmed_instead_of_success(self) -> None:
        driver, wait, _old, _new, actions = self._query_fixture([{"ready": True, "failed": False}])
        driver.find_elements.side_effect = lambda *_args: []
        with self.assertRaisesRegex(RuntimeError, "empty_list"):
            self.module.query_grid_rows(driver, wait, "maintain")
        self.assertEqual(actions[-1], "cleanup")

    def _query_fixture(self, states):
        module = self.module
        old_row, new_row = object(), object()
        visible = [old_row]
        observed = iter(states)
        actions = []
        driver, wait = mock.Mock(), mock.Mock()

        def execute(_script, action=None, token=None, *_args):
            if _script == module.GRID_PAGE_READY_SCRIPT:
                return True
            actions.append(action)
            if action == "state":
                state = next(observed, states[-1])
                if isinstance(state, Exception):
                    raise state
                if state.get("ready"):
                    visible[:] = [new_row]
                return {"token": token, **state}
            return True

        driver.execute_script.side_effect = execute
        driver.find_elements.side_effect = lambda *_args: list(visible)
        return driver, wait, old_row, new_row, actions

    def test_search_waits_for_current_binding_instead_of_using_old_rows(self) -> None:
        driver, wait, _old, new, actions = self._query_fixture([
            {"ready": False, "failed": False}, {"ready": True, "failed": False}])
        with mock.patch.object(self.module.time, "sleep"):
            rows = self.module.query_grid_rows(driver, wait, "maintain")
        self.assertEqual(rows, [new])
        self.assertEqual(driver.find_elements.call_count, 1)
        self.assertEqual(actions[:2], ["begin", "click"])
        self.assertGreaterEqual(actions.count("state"), 4)
        self.assertEqual(actions[-1], "cleanup")

    def test_source_change_while_reading_rows_is_unconfirmed(self) -> None:
        driver, wait, _old, _new, actions = self._query_fixture([
            {"ready": True, "failed": False}, {"ready": True, "failed": False},
            {"ready": False, "failed": True}])
        with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
            self.module.query_grid_rows(driver, wait, "maintain")
        driver.find_elements.assert_called_once()
        self.assertEqual(actions[-1], "cleanup")

    def test_search_error_and_missing_binding_are_unconfirmed_and_cleanup(self) -> None:
        for state in ({"ready": False, "failed": True}, {"ready": False, "failed": False}):
            with self.subTest(state=state):
                driver, wait, _old, _new, actions = self._query_fixture([state])
                with mock.patch.object(self.module.time, "monotonic", side_effect=[0, 0, 11]), \
                        mock.patch.object(self.module.time, "sleep"):
                    with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
                        self.module.query_grid_rows(driver, wait, "maintain")
                driver.find_elements.assert_not_called()
                self.assertEqual(actions[-1], "cleanup")

    def test_navigation_waits_for_current_document_and_query_initialization_before_begin(self) -> None:
        runner = r'''
import json
import sys
from selenium.common.exceptions import NoSuchElementException
from selenium.webdriver.support.ui import WebDriverWait
from PySide6.QtCore import QCoreApplication
from PySide6.QtQml import QJSEngine
from tests.test_smoke import DailyVehicleVerificationTests

DailyVehicleVerificationTests.setUpClass()
module = DailyVehicleVerificationTests.module
app = QCoreApplication([])
engine = QJSEngine()
engine.evaluate("""
var document = {readyState: 'complete'};
var location = {origin: 'https://fixture.invalid', pathname: '/CarMaintainCheck/Index', href: 'https://fixture.invalid/CarMaintainCheck/Index'};
var window = {document: document, location: location};
var resourceEntries = [{name: location.origin + '/CarMaintainCheck/QueryProc', responseEnd: 1}];
var performance = {getEntriesByType: kind => kind === 'resource' ? resourceEntries : []};
window.performance = performance;
var expectedPath = location.pathname, queryName = 'QueryData', requireSource = true;
function source(uid) {return {view: () => [{uid: uid}], bind: () => {}, unbind: () => {}};}
function makeGrid(uid) {const grid = {dataSource: source(uid), bind: () => {}, unbind: () => {}};
    grid.tbody = {find: () => [{getAttribute: () => grid.dataSource.view()[0].uid}]};
    grid.setDataSource = value => {grid.dataSource = value;}; return grid;}
var grid = makeGrid('old');
function installQuery() {window.jQuery = function() {return {data: () => grid};};
    window.jQuery.fn = {jquery: 'fixture'}; window.jQuery.active = 0;
    window.QueryData = function() {}; window.Query = function() {};}
function fixtureReady() {return location.pathname === expectedPath &&
    document.readyState === 'complete' && typeof window.jQuery === 'function' &&
    window.jQuery.active === 0 &&
    typeof window[queryName] === 'function' && (!requireSource || !!grid && !!grid.dataSource);}
function releaseDocument() {location.pathname = expectedPath;
    location.href = 'https://fixture.invalid' + expectedPath;
    document.readyState = 'complete'; installQuery(); grid = makeGrid('current');}
installQuery();
var fixtureButton = {click: () => {grid = makeGrid('searched');}};
""")
pending = sys.argv[1]
engine.evaluate({
    'new_document_scripts_loading': "document.readyState = 'loading'; delete window.jQuery; delete window.QueryData;",
    'old_document_search_still_visible': "location.pathname = '/CarEquipCheck/Index'; location.href = 'https://fixture.invalid/CarEquipCheck/Index';",
    'query_function_not_initialized': "delete window.QueryData;",
    'query_datasource_not_initialized': "grid.dataSource = null;",
    'initial_query_still_running': "window.jQuery.active = 1;",
    'equipment_without_initial_grid': "expectedPath = '/CarEquipCheck/Index'; location.pathname = expectedPath; queryName = 'Query'; requireSource = false; grid = null; resourceEntries = [];",
}[pending])

class Button:
    def is_displayed(self): return True
    def is_enabled(self): return True

class Driver:
    def __init__(self):
        self.readiness_polls = 0
        self.actions = []
        self.early = []
    @property
    def current_url(self): return engine.evaluate('location.href').toString()
    def find_element(self, by, selector):
        if selector == 'spinner-container':
            raise NoSuchElementException('fixture spinner has not been created')
        return Button()
    def find_elements(self, by, selector): return ['current-row']
    def execute_script(self, script, *args):
        observer = script == module.GRID_QUERY_OBSERVER_SCRIPT
        if observer:
            self.actions.append(args[0])
            if args[0] in ('begin', 'click') and not engine.evaluate('fixtureReady()').toBool():
                self.early.append(args[0])
        else:
            self.readiness_polls += 1
            if self.readiness_polls >= 2:
                engine.evaluate('releaseDocument()')
        values = [json.dumps(arg) if not isinstance(arg, Button) else 'fixtureButton' for arg in args]
        result = engine.evaluate('(function(){' + script + '}).apply(null,[' + ','.join(values) + ']);')
        if result.isError(): raise AssertionError(result.toString())
        return result.toVariant()

driver = Driver()
error = ''
try:
    context = 'equip' if pending == 'equipment_without_initial_grid' else 'maintain'
    rows = module.query_grid_rows(driver, WebDriverWait(driver, 0.2, poll_frequency=0.01), context)
except Exception as failure:
    rows, error = [], type(failure).__name__
print(json.dumps({'rows': rows, 'error': error, 'early': driver.early, 'actions': driver.actions,
                  'readiness_polls': driver.readiness_polls}))
'''
        for pending in ("new_document_scripts_loading", "old_document_search_still_visible",
                        "query_function_not_initialized", "query_datasource_not_initialized",
                        "initial_query_still_running",
                        "equipment_without_initial_grid"):
            with self.subTest(pending=pending):
                result = subprocess.run([sys.executable, "-B", "-c", runner, pending], cwd=PROJECT_ROOT,
                                        capture_output=True, text=True, encoding="utf-8", timeout=15,
                                        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                self.assertEqual(result.returncode, 0, result.stderr)
                evidence = json.loads(result.stdout.splitlines()[-1])
                self.assertEqual(evidence["early"], [], evidence)
                self.assertEqual(evidence["error"], "", evidence)
                self.assertEqual(evidence["rows"], ["current-row"], evidence)
                self.assertGreater(evidence["readiness_polls"], 0)
                self.assertEqual(evidence["actions"][0:2], ["begin", "click"])
                self.assertEqual(evidence["actions"][-1], "cleanup")

    def test_maintenance_page_requires_completed_matching_initial_query_resource(self) -> None:
        javascript = """
var document = {readyState: 'complete'};
var window = {location: {origin: 'https://fixture.invalid', pathname: '/CarMaintainCheck/Index'}};
var resourceEntries = [];
var performance = {getEntriesByType: kind => kind === 'resource' ? resourceEntries : []};
window.performance = performance;
var grid = {setDataSource: () => {}, dataSource: {view: () => []}};
window.jQuery = () => ({data: () => grid}); window.jQuery.active = 0;
window.QueryData = () => {}; window.Query = () => {};
function ready(requireGrid) {return (function() {
""" + self.module.GRID_PAGE_READY_SCRIPT + """
}).apply(null, [window.location.pathname, requireGrid ? 'QueryData' : 'Query', requireGrid]);}
var cases = {};
cases.queuedBeforeFirstRequest = ready(true);
resourceEntries = [{name: window.location.origin + '/CarMaintainCheck/QueryProc', responseEnd: 0}];
cases.incomplete = ready(true);
resourceEntries[0].responseEnd = -1; cases.negativeEnd = ready(true);
resourceEntries = [{name: 'https://other.invalid/CarMaintainCheck/QueryProc', responseEnd: 1}];
cases.wrongOrigin = ready(true);
resourceEntries = [{name: window.location.origin + '/CarMaintainCheck/QueryProcExtra', responseEnd: 1}];
cases.endpointSuffix = ready(true);
resourceEntries = [{name: window.location.origin + '/CarEquipCheck/QueryProc', responseEnd: 1}];
cases.wrongEndpoint = ready(true);
resourceEntries = [{name: window.location.origin + '/CarMaintainCheck/QueryProc?initial=1#response', responseEnd: 1}];
cases.matchingCompleted = ready(true);
window.jQuery.active = 1; cases.completedWhileAnotherRequestPending = ready(true);
window.jQuery.active = 0;
performance.getEntriesByType = () => {throw new Error('fixture unavailable timing');};
cases.unavailableTiming = ready(true);
performance.getEntriesByType = kind => kind === 'resource' ? resourceEntries : [];
resourceEntries = []; grid = null; window.location.pathname = '/CarEquipCheck/Index';
cases.equipmentWithoutInitialResource = ready(false);
JSON.stringify(cases);
"""
        runner = ("from PySide6.QtCore import QCoreApplication\nfrom PySide6.QtQml import QJSEngine\n"
                  "app=QCoreApplication([])\nengine=QJSEngine()\nresult=engine.evaluate(" + repr(javascript) + ")\n"
                  "assert not result.isError(),result.toString()\nprint(result.toString())\n")
        result = subprocess.run([sys.executable, "-B", "-c", runner], cwd=PROJECT_ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=15,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(result.returncode, 0, result.stderr)
        evidence = json.loads(result.stdout)
        for case, expected in {
            "queuedBeforeFirstRequest": False, "incomplete": False, "negativeEnd": False,
            "wrongOrigin": False, "endpointSuffix": False, "wrongEndpoint": False,
            "matchingCompleted": True, "completedWhileAnotherRequestPending": False,
            "unavailableTiming": False, "equipmentWithoutInitialResource": True,
        }.items():
            with self.subTest(case=case):
                self.assertEqual(evidence[case], expected, evidence)

    def test_cold_page_waits_for_initial_ajax_before_manual_query(self) -> None:
        # Each response completes between WebDriver commands, never inside one JS callback.
        runner = r'''
import json
import sys
from selenium.common.exceptions import NoSuchElementException
from selenium.webdriver.support.ui import WebDriverWait
from PySide6.QtCore import QCoreApplication
from PySide6.QtQml import QJSEngine
from tests.test_smoke import DailyVehicleVerificationTests

DailyVehicleVerificationTests.setUpClass()
module = DailyVehicleVerificationTests.module
app = QCoreApplication([])
engine = QJSEngine()
engine.evaluate("""
var document = {readyState: 'complete'};
var window = {location: {origin: 'https://fixture.invalid', pathname: '/CarMaintainCheck/Index'}};
var resourceEntries = [];
var performance = {getEntriesByType: kind => kind === 'resource' ? resourceEntries : []};
window.performance = performance;
function source(uid) {return {view: () => [{uid}], bind: () => {}, unbind: () => {}};}
var grid = {dataSource: source('initial-placeholder'), dom: 'initial-placeholder',
            bind: () => {}, unbind: () => {}, setDataSource: value => {grid.dataSource = value;}};
grid.tbody = {find: () => [{getAttribute: () => grid.dom}]};
window.jQuery = () => ({data: () => grid});
window.jQuery.active = 0;
var clock = 0, initialStarted = false;
var initialPending = true, manualPending = false, clickedWhileInitial = false, clickedAt = null;
function response(uid) {grid.setDataSource(source(uid)); grid.dom = uid; window.jQuery.active--;}
function initialResponse() {initialPending = false; response('auto-A');
    resourceEntries.push({name: window.location.origin + '/CarMaintainCheck/QueryProc?initial=1', responseEnd: clock});}
function advanceInitial() {clock += 250;
    if (!initialStarted && clock >= 500) {initialStarted = true; window.jQuery.active++;}
    if (initialPending && initialStarted && clock >= 750) initialResponse();}
function manualResponse() {manualPending = false; response('manual-B');}
window.QueryData = () => {clickedWhileInitial = initialPending; clickedAt = clock;
    manualPending = true; window.jQuery.active++;};
var fixtureButton = {click: () => window.QueryData()};
""")
if sys.argv[1] == 'running':
    engine.evaluate('initialStarted = true; window.jQuery.active = 1;')

class Button:
    def is_displayed(self): return True
    def is_enabled(self): return True

class Driver:
    def __init__(self):
        self.readiness_polls = 0
        self.readiness = []
        self.observations = []
        self.actions = []
    def find_element(self, by, selector):
        if selector == 'spinner-container':
            raise NoSuchElementException('fixture no visible loading overlay')
        return Button()
    def find_elements(self, by, selector): return [engine.evaluate('grid.dom').toString()]
    def execute_script(self, script, *args):
        if script == module.GRID_PAGE_READY_SCRIPT:
            self.readiness_polls += 1
            engine.evaluate('advanceInitial()')
        elif script == module.GRID_QUERY_OBSERVER_SCRIPT:
            self.actions.append(args[0])
            if args[0] == 'state':
                if engine.evaluate('initialPending').toBool():
                    engine.evaluate('advanceInitial()')
                elif engine.evaluate('manualPending').toBool():
                    engine.evaluate('manualResponse()')
        values = [json.dumps(arg) if not isinstance(arg, Button) else 'fixtureButton' for arg in args]
        result = engine.evaluate('(function(){' + script + '}).apply(null,[' + ','.join(values) + ']);')
        if result.isError(): raise AssertionError(result.toString())
        value = result.toVariant()
        if script == module.GRID_PAGE_READY_SCRIPT:
            self.readiness.append({'ready': value, 'clock': engine.evaluate('clock').toInt(),
                                   'active': engine.evaluate('window.jQuery.active').toInt()})
        if script == module.GRID_QUERY_OBSERVER_SCRIPT and args[0] == 'state':
            self.observations.append(value)
        return value

driver = Driver()
error = ''
try:
    rows = module.query_grid_rows(driver, WebDriverWait(driver, 0.2, poll_frequency=0.01), 'maintain')
except Exception as failure:
    rows, error = [], type(failure).__name__
print(json.dumps({'rows': rows, 'error': error, 'readiness_polls': driver.readiness_polls,
                  'clicked_while_initial': engine.evaluate('clickedWhileInitial').toBool(),
                  'clicked_at': engine.evaluate('clickedAt').toVariant(), 'readiness': driver.readiness,
                  'observations': driver.observations, 'actions': driver.actions}))
'''
        for initial in ("running", "queued_500ms"):
            with self.subTest(initial=initial):
                result = subprocess.run([sys.executable, "-B", "-c", runner, initial], cwd=PROJECT_ROOT,
                                        capture_output=True, text=True, encoding="utf-8", timeout=15,
                                        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                self.assertEqual(result.returncode, 0, result.stderr)
                evidence = json.loads(result.stdout.splitlines()[-1])
                self.assertEqual(evidence["error"], "", evidence)
                self.assertEqual(evidence["rows"], ["manual-B"], evidence)
                self.assertFalse(evidence["clicked_while_initial"], evidence)
                self.assertGreaterEqual(evidence["clicked_at"], 750, evidence)
                self.assertGreaterEqual(evidence["readiness_polls"], 3, evidence)
                if initial == "queued_500ms":
                    self.assertEqual(evidence["readiness"][0], {"ready": False, "clock": 250, "active": 0}, evidence)
                self.assertTrue(all(value["ready"] is True and value["failed"] is False
                                    for value in evidence["observations"]), evidence)
                self.assertEqual(evidence["actions"][-1], "cleanup")

    def test_query_verify_diagnostics_distinguish_fixed_failure_boundaries(self) -> None:
        from selenium.common.exceptions import JavascriptException, TimeoutException, WebDriverException
        from urllib3.exceptions import MaxRetryError

        module = self.module
        ready = {"ready": True, "failed": False}
        changed = {"ready": False, "failed": True}
        failures = (
            ("query_spinner_timeout", [ready], TimeoutException("token=fixture-secret"), None, 0),
            ("query_spinner_driver_error", [ready], WebDriverException("token=fixture-secret"), None, 0),
            ("query_spinner_driver_error", [ready],
             MaxRetryError(None, "https://fixture.invalid/?token=fixture-secret"), None, 0),
            ("query_source_changed_before_read", [ready, changed], None, None, 0),
            ("query_source_changed_after_read", [ready, ready, changed], None, None, 1),
            ("query_state_read_failed_before_rows", [ready, JavascriptException("token=fixture-secret")], None, None, 0),
            ("query_state_read_failed_after_rows", [ready, ready, JavascriptException("token=fixture-secret")], None, None, 1),
            ("query_state_read_failed_before_rows", [ready, MaxRetryError(None, "https://fixture.invalid/?token=fixture-secret")], None, None, 0),
            ("query_state_read_failed_after_rows", [ready, ready, MaxRetryError(None, "https://fixture.invalid/?token=fixture-secret")], None, None, 1),
            ("query_rows_read_failed", [ready, ready], None,
             WebDriverException("cookie=private-cookie https://fixture.invalid/?token=fixture-secret"), 1),
            ("query_rows_read_failed", [ready, ready], None,
             MaxRetryError(None, "https://fixture.invalid/?token=fixture-secret"), 1),
        )
        for context, operation in (("maintain", "maintenance"), ("equip", "equipment")):
            for code, states, spinner_error, read_error, read_count in failures:
                with self.subTest(operation=operation, code=code):
                    driver, wait, _old, _new, actions = self._query_fixture(states)
                    if read_error is not None:
                        driver.find_elements.side_effect = read_error
                    spinner = mock.patch.object(module, "wait_for_spinner", side_effect=[None, spinner_error]) if spinner_error else mock.patch.object(module, "wait_for_spinner")
                    report = module.new_vehicle_report("2026-10-19")
                    with spinner, mock.patch("builtins.print") as output:
                        with self.assertRaises(module.VehicleVerificationError) as failure:
                            module.verified_grid_rows(driver, wait, operation, report)
                    lines = [str(call.args[0]) for call in output.call_args_list if call.args]
                    diagnostics = [line for line in lines if line.startswith("[sinposmart-vehicle-diagnostic] ")]
                    self.assertEqual(len(diagnostics), 1, lines)
                    self.assertIn(code, diagnostics[0])
                    self.assertEqual(report["reason_code"], "state_unconfirmed")
                    self.assertEqual(driver.find_elements.call_count, read_count)
                    self.assertIn(f"[sinposmart-stage] {operation}_query_verify", lines)
                    for private in ("fixture-secret", "private-cookie", "https://fixture.invalid"):
                        self.assertNotIn(private, "\n".join(lines) + str(failure.exception))
                    self.assertEqual(actions[-1], "cleanup")

    def test_query_stages_identify_last_fixed_phase_without_raw_details(self) -> None:
        module = self.module
        for operation, context in (("maintenance", "maintain"), ("equipment", "equip")):
            with self.subTest(operation=operation):
                driver, wait, _old, _new, _actions = self._query_fixture([{"ready": True, "failed": False}])
                fixture_driver, _wait, _state, rows = self._fixture()
                fixture_driver.mode = context
                original = driver.execute_script.side_effect
                driver.execute_script.side_effect = lambda script, *args: (
                    {"total": 1, "page": 1, "page_size": 20, "filtered": False}
                    if script == module.GRID_STATE_SCRIPT else fixture_driver.execute_script(script, *args)
                    if "grid.dataItem" in script else original(script, *args))
                driver.find_elements.side_effect = rows
                with mock.patch("builtins.print") as output:
                    module.verified_grid_rows(driver, wait, operation, module.new_vehicle_report("2026-10-17"))
                stages = [call.args[0] for call in output.call_args_list
                          if call.args and str(call.args[0]).startswith("[sinposmart-stage] ")]
                self.assertEqual(stages, [f"[sinposmart-stage] {operation}_{step}" for step in
                                         ("page_ready", "query_start", "query_wait", "query_verify",
                                          "inventory_verify", "check")])
                driver, wait, _old, _new, _actions = self._query_fixture([{"ready": False, "failed": True}])
                with mock.patch("builtins.print") as output, self.assertRaises(module.VehicleVerificationError):
                    module.query_grid_rows(driver, wait, context)
                self.assertEqual([call.args[0] for call in output.call_args_list],
                                 [f"[sinposmart-stage] {operation}_{step}"
                                  for step in ("page_ready", "query_start", "query_wait")])

    def test_same_path_navigation_waits_for_old_document_to_become_stale_before_query(self) -> None:
        from selenium.common.exceptions import StaleElementReferenceException

        module = self.module
        for operation in ("maintenance", "equipment"):
            with self.subTest(operation=operation):
                driver, _wait, state, query = self._fixture(maintenance_done=True, equipment_done=True)
                state["committed"], state["stale_polls"] = False, 0
                original_find = driver.find_element
                original_get = driver.get

                class OldDocument:
                    def is_enabled(self):
                        state["stale_polls"] += 1
                        if state["stale_polls"] >= 2:
                            state["committed"] = True
                            raise StaleElementReferenceException("fixture new document committed")
                        return True

                driver.find_element = lambda by, selector: OldDocument() if selector == "html" else original_find(by, selector)
                driver.get = lambda url: original_get(url)
                driver.current_url = module.MAINTAIN_LIST_URL if operation == "maintenance" else module.EQUIP_CHECK_LIST_URL

                def current_rows(*args):
                    self.assertTrue(state["committed"], "query used the old complete document at the same pathname")
                    return query(*args)

                with mock.patch.object(module, "query_grid_rows", side_effect=current_rows), \
                        mock.patch.object(module, "ensure_requested_date"):
                    wait = module.WebDriverWait(driver, 0.2, poll_frequency=0.01)
                    if operation == "maintenance":
                        today = datetime(2026, 10, 17)
                        module.process_maintain_checks(driver, wait, module.build_today_strings(today), today)
                    else:
                        module.process_equip_checks(driver, wait, module.new_vehicle_report("2026-10-17"))
                self.assertGreaterEqual(state["stale_polls"], 2)

    def test_refresh_waits_for_new_calendar_after_server_save_before_readback(self) -> None:
        from selenium.common.exceptions import StaleElementReferenceException

        module = self.module
        driver, _wait, state, query = self._fixture(persist=False)
        original_find, original_get, original_script = driver.find_element, driver.get, driver.execute_script
        state.update(document=0, refresh_pending=False, refresh_polls=0, server_saved=False)

        class Document:
            def __init__(self): self.generation = state["document"]
            def is_enabled(self):
                if state["refresh_pending"]:
                    state["refresh_polls"] += 1
                    if state["refresh_polls"] >= 2:
                        state["refresh_pending"] = False
                        state["document"] += 1
                        state["maintenance_done"] = state["server_saved"]
                if self.generation != state["document"]:
                    raise StaleElementReferenceException("fixture new calendar committed")
                return True

        def navigate(url):
            state["document"] += 1
            original_get(url)

        def execute(script, *args):
            result = original_script(script, *args)
            if args and getattr(args[0], "role", "") == "save":
                state["server_saved"] = True
            return result

        driver.find_element = lambda by, selector: Document() if selector == "html" else original_find(by, selector)
        driver.get = navigate
        driver.refresh = lambda: state.update(refresh_pending=True)
        driver.execute_script = execute
        with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                mock.patch.object(module, "ensure_requested_date"):
            today = datetime(2026, 10, 17)
            module.process_maintain_checks(driver, module.WebDriverWait(driver, 0.2, poll_frequency=0.01),
                                          module.build_today_strings(today), today)
        self.assertTrue(state["server_saved"])
        self.assertGreaterEqual(state["refresh_polls"], 2)
        self.assertEqual(state["saved"], 1)

    def test_query_observer_ignores_old_binding_and_removes_only_its_handlers(self) -> None:
        javascript = """
function observable() {
    return {handlers: {}, bind: function(name, callback) {
        (this.handlers[name] || (this.handlers[name] = [])).push(callback);
    }, unbind: function(name, callback) {
        this.handlers[name] = (this.handlers[name] || []).filter(fn => fn !== callback);
    }, fire: function(name, event) {
        (this.handlers[name] || []).slice().forEach(fn => fn(event));
    }};
}
var grid = observable(), source = observable();
grid.dataSource = source;
source.view = () => [{uid: 'current-row'}];
grid.tbody = {find: () => [{getAttribute: () => 'current-row'}]};
var window = {jQuery: function() {return {data: function() {return grid;}};}};
window.jQuery.active = 0;
function query(action, token, button) { return (function() {
""" + self.module.GRID_QUERY_OBSERVER_SCRIPT + """
}).apply(null, [action, token, button]); }
var externalCalls = 0;
var external = function() {externalCalls++;};
grid.bind('dataBound', external);
query('begin', 'first');
var oldStart = source.handlers.requestStart[0], oldBound = grid.handlers.dataBound[1];
grid.fire('dataBound', {});
var beforeClick = query('state', 'first').ready;
query('click', 'first', {click: function() {}});
grid.fire('dataBound', {});
var oldResponse = query('state', 'first').ready;
source.fire('requestStart', {type: 'read'});
grid.fire('dataBound', {});
var currentResponse = query('state', 'first').ready;
query('begin', 'second');
query('cleanup', 'first');
query('click', 'second', {click: function() {}});
oldStart({type: 'read'}); oldBound();
var staleCallback = query('state', 'second').ready;
source.fire('error', {});
var failed = query('state', 'second').failed;
query('cleanup', 'second');
JSON.stringify({beforeClick: beforeClick, oldResponse: oldResponse,
    currentResponse: currentResponse, staleCallback: staleCallback, failed: failed,
    gridHandlers: grid.handlers.dataBound.length,
    startHandlers: source.handlers.requestStart.length, errorHandlers: source.handlers.error.length,
    externalRetained: grid.handlers.dataBound[0] === external});
"""
        runner = (
            "import json\nfrom PySide6.QtCore import QCoreApplication\n"
            "from PySide6.QtQml import QJSEngine\napp=QCoreApplication([])\nengine=QJSEngine()\n"
            "result=engine.evaluate(" + repr(javascript) + ")\n"
            "assert not result.isError(), result.toString()\nprint(result.toString())\n"
        )
        result = subprocess.run([sys.executable, "-B", "-c", runner], cwd=PROJECT_ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=15,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {
            "beforeClick": False, "oldResponse": False, "currentResponse": True,
            "staleCallback": False, "failed": True, "gridHandlers": 1,
            "startHandlers": 0, "errorHandlers": 0, "externalRetained": True})

    def test_query_observer_waits_for_all_pending_responses_and_keeps_latched_changes_failed(self) -> None:
        javascript = """
function emitter() {return {handlers: {}, bind: function(n,f) {(this.handlers[n]||(this.handlers[n]=[])).push(f);},
    unbind: function(n,f) {this.handlers[n]=(this.handlers[n]||[]).filter(x=>x!==f);},
    fire: function(n,e) {(this.handlers[n]||[]).slice().forEach(f=>f(e));}};}
function source(ids) {const s=emitter(); s.ids=ids; s.view=()=>s.ids.map(uid=>({uid})); return s;}
function makeGrid(ids) {const g=emitter(); g.dataSource=source(ids); g.dom=ids;
    g.tbody={find:()=>g.dom.map(uid=>({getAttribute:()=>uid}))}; return g;}
var current=makeGrid(['old']);
var window={jQuery:()=>({data:()=>current})}; window.jQuery.active=0;
function query(a,t,b) {return (function() {
""" + self.module.GRID_QUERY_OBSERVER_SCRIPT + """
}).apply(null,[a,t,b]);}
function snapshot(t) {const s=query('state',t); return {ready:s.ready,failed:s.failed};}
query('begin','overlap');
query('click','overlap',{click:()=>{window.jQuery.active=2;}});
current.dataSource=source(['auto-A']); current.dom=['auto-A']; window.jQuery.active=1;
var firstPending=snapshot('overlap');
current.dataSource=source(['manual-B']); current.dom=['manual-B']; window.jQuery.active=0;
var lastCompleted=snapshot('overlap'), stableRead=snapshot('overlap');
window.jQuery.active=1; var newPending=snapshot('overlap');
const latchedSource=current.dataSource;
current.dataSource=source(['late-C']); current.dom=['late-C'];
var sourceChanged=snapshot('overlap');
current.dataSource=latchedSource; current.dom=['manual-B']; window.jQuery.active=0;
var sourceRestored=snapshot('overlap'); query('cleanup','overlap');
query('begin','uids');
query('click','uids',{click:()=>{current=makeGrid(['uid-A']);}});
var uidLatched=snapshot('uids');
current.dataSource.ids=['uid-B']; current.dom=['uid-B'];
var uidChanged=snapshot('uids');
current.dataSource.ids=['uid-A']; current.dom=['uid-A'];
var uidRestored=snapshot('uids'); query('cleanup','uids');
JSON.stringify({firstPending,lastCompleted,stableRead,newPending,sourceChanged,sourceRestored,
    uidLatched,uidChanged,uidRestored});
"""
        runner = ("from PySide6.QtCore import QCoreApplication\nfrom PySide6.QtQml import QJSEngine\n"
                  "app=QCoreApplication([])\nengine=QJSEngine()\nresult=engine.evaluate(" + repr(javascript) + ")\n"
                  "assert not result.isError(),result.toString()\nprint(result.toString())\n")
        result = subprocess.run([sys.executable, "-B", "-c", runner], cwd=PROJECT_ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=15,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(result.returncode, 0, result.stderr)
        evidence = json.loads(result.stdout)
        for state, expected in {
            "firstPending": {"ready": False, "failed": False},
            "lastCompleted": {"ready": True, "failed": False},
            "stableRead": {"ready": True, "failed": False},
            "newPending": {"ready": False, "failed": False},
            "sourceChanged": {"ready": False, "failed": True},
            "sourceRestored": {"ready": False, "failed": True},
            "uidLatched": {"ready": True, "failed": False},
            "uidChanged": {"ready": False, "failed": True},
            "uidRestored": {"ready": False, "failed": True},
        }.items():
            with self.subTest(state=state):
                self.assertEqual(evidence[state], expected, evidence)

    def test_live_query_replacement_rebuilt_and_initial_grid_require_current_render(self) -> None:
        javascript = """
function emitter() {return {handlers: {}, bind: function(n,f) {(this.handlers[n]||(this.handlers[n]=[])).push(f);},
    unbind: function(n,f) {this.handlers[n]=(this.handlers[n]||[]).filter(x=>x!==f);},
    fire: function(n,e) {(this.handlers[n]||[]).slice().forEach(f=>f(e));}};}
function source(ids) {const s=emitter(); s.view=()=>ids.map(uid=>({uid})); return s;}
function makeGrid(ids) {const g=emitter(); g.dataSource=source(ids); g.dom=ids;
    g.tbody={find:()=>g.dom.map(uid=>({getAttribute:()=>uid}))}; return g;}
var current=makeGrid(['old']);
var window={jQuery:()=>({data:()=>current})};
window.jQuery.active = 0;
function query(a,t,b) {return (function() {
""" + self.module.GRID_QUERY_OBSERVER_SCRIPT + """
}).apply(null,[a,t,b]);}
function ready(t) {const s=query('state',t); return !!s && s.ready===true;}
query('begin','replacement');
query('click','replacement',{click:()=>{current.dataSource=source(['replacement']);
    current.dom=['replacement']; current.fire('dataBound',{});}});
var replacement=ready('replacement'); query('cleanup','replacement');
query('begin','rebuilt');
query('click','rebuilt',{click:()=>{current=null; current=makeGrid(['rebuilt']);}});
var rebuilt=ready('rebuilt'); query('cleanup','rebuilt');
current=null; var initial=query('begin','initial');
query('click','initial',{click:()=>{current=makeGrid(['initial']);}});
initial=initial && ready('initial'); query('cleanup','initial');
query('begin','render');
query('click','render',{click:()=>{current=makeGrid(['new-view']);current.dom=['old-dom'];}});
var staleRender=ready('render'); current.dom=['new-view']; var rendered=ready('render');
current=makeGrid(['late-replacement']); var changed=query('state','render');
query('cleanup','render');
JSON.stringify({replacement,rebuilt,initial,staleRender,rendered,changedRejected:!changed.ready});
"""
        runner = ("from PySide6.QtCore import QCoreApplication\nfrom PySide6.QtQml import QJSEngine\n"
                  "app=QCoreApplication([])\nengine=QJSEngine()\nresult=engine.evaluate(" + repr(javascript) + ")\n"
                  "assert not result.isError(),result.toString()\nprint(result.toString())\n")
        result = subprocess.run([sys.executable, "-B", "-c", runner], capture_output=True,
                                text=True, encoding="utf-8", timeout=15,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {"replacement": True, "rebuilt": True,
            "initial": True, "staleRender": False, "rendered": True, "changedRejected": True})

    def _fixture(self, **options):
        from selenium.common.exceptions import NoSuchElementException, StaleElementReferenceException

        module = self.module
        state = {"maintenance_done": False, "equipment_done": False, "equipment_started": False,
                 "persist": True, "event_class": "unchecked", "missing_today": False,
                 "event_classes": None, "reviewed": False,
                 "total": 1, "filtered": False, "placeholder": False,
                 "equipment_text": "", "equipment_button": True,
                 "checker": None, "start": None, "end": None,
                 "contradictory_button": False, "no_paging": False,
                 "row_names": None,
                 "button_only_evidence": False, "change_after_save": False,
                 "record_id": "101", "readback_id": None, "model_vehicle": None, "model_date": None,
                 "dialog_mode": None, "dialog_open": False, "dialog_reads": 0,
                 "selected": False, "selection_clicks": 0, "confirm_clicks": 0, "cancel_clicks": 0,
                 "selection_ready_after": 0, "selection_reads": 0, "selection_incomplete": False,
                 "refresh_header_on_select": False, "header_generation": 0,
                 "header_lookups": 0, "stale_header_uses": 0,
                 "continue_present": True, "continue_displayed": True, "continue_enabled": True,
                 "complete_on_view": False, "view_record_id": None,
                 "view_clicks": 0, "continue_clicks": 0, "view_ids": [], "direct_take_gets": 0,
                 "end_take_ready_after": 0, "end_take_reads": 0, "end_take_pending": False,
                 "end_take_interrupted": False, "terminal_view_urls": [],
                 "take_ids": [], "saved": 0}
        state.update(options)

        class Element:
            def __init__(self, role="button", text="", displayed=True, event_class=None, enabled=True):
                self.role, self.text, self.displayed = role, text, displayed
                self.event_class, self.enabled = event_class, enabled
                self.header_generation = state["header_generation"]

            def is_displayed(self):
                return self.displayed

            def is_enabled(self):
                return self.enabled

            def is_selected(self):
                return state["selected"]

            def get_attribute(self, name):
                if name == "textContent":
                    return self.text
                if name == "aria-label":
                    if self.role == "checkbox":
                        return "Deselect all rows" if state["selected"] else "Select all rows"
                    return "2026年10月16日" if state["missing_today"] else "2026年10月17日"
                if name == "class":
                    return "checked" if state["maintenance_done"] else self.event_class or state["event_class"]
                return ""

            def find_elements(self, by, selector):
                if self.role == "equipment-status" and state["button_only_evidence"]:
                    return [Element(text="完成清點")]
                if self.role == "event" and selector == ".reviewed" and state["reviewed"]:
                    return [Element()]
                return []

        class Row(Element):
            def __init__(self, vehicle=None):
                super().__init__()
                self.vehicle = vehicle

            def find_elements(self, by, selector):
                if selector == "td":
                    if state["placeholder"]:
                        return [Element(text="查無資料")]
                    vehicle = self.vehicle or ("CAR-02" if state["change_after_save"] and state["saved"] else "CAR-01")
                    if driver.mode == "maintain":
                        cells = [Element() for _ in range(13)]
                        cells[8] = Element(text=vehicle)
                        cells[9] = Element(text="115/10/17" if state["maintenance_done"] else "未完成")
                        return cells
                    text = state["equipment_text"] or "115/10/17"
                    started = state["equipment_done"] or state["equipment_started"]
                    checker = state["checker"] if state["checker"] is not None else "同仁" if started else "-"
                    start = state["start"] if state["start"] is not None else "08:40" if started else "-"
                    end = state["end"] if state["end"] is not None else "08:40" if state["equipment_done"] else "-"
                    return [Element(), Element(), Element(), Element(text=vehicle), Element("equipment-status", text),
                            Element(text="-"), Element(text=checker), Element(text=start), Element(text=end), Element()]
                if selector == ".k-grid-view":
                    return [Element("view")] if state["equipment_done"] or state["equipment_started"] else []
                if selector not in (".k-grid-check", self_module.EQUIPMENT_BUTTON_XPATH):
                    return []
                if driver.mode == "equip" and state["equipment_button"] and (
                        not (state["equipment_done"] or state["equipment_started"]) or state["contradictory_button"]):
                    return [Element("check")]
                return []

        class Document:
            def __init__(self, generation): self.generation = generation
            def is_enabled(self):
                if self.generation != driver.document_generation:
                    raise StaleElementReferenceException("fixture document was replaced")
                return True

        class Driver:
            mode = "maintain"
            current_url = "fixture://maintain"
            document_generation = 0
            control_navigation = False

            def get(self, url):
                if "/CarEquipCheck/Index" in url and state["end_take_pending"]:
                    state["end_take_interrupted"] = True
                    state["end_take_pending"] = False
                self.document_generation += 1
                self.mode = "equip" if "CarEquipCheck" in url else "maintain"
                self.current_url = url
                if "/CarEquipCheck/Take?id=" in url:
                    state["take_ids"].append(url.rsplit("=", 1)[-1])
                    if not self.control_navigation:
                        state["direct_take_gets"] += 1
                if "/CarEquipCheck/View?id=" in url:
                    state["view_ids"].append(url.rsplit("=", 1)[-1])
                    if state["complete_on_view"]:
                        state["equipment_done"] = True
                if "/CarEquipCheck/EndTake?id=" in url:
                    state["end_take_pending"] = True
                    if not state["end_take_ready_after"]:
                        self.commit_end_take()

            def commit_end_take(self):
                if state["persist"]:
                    state["equipment_done"] = True
                    if state["checker"] in {"", "-", "—", "－"}:
                        state["checker"] = "同仁"
                state["end_take_pending"] = False
                self.document_generation += 1
                self.current_url = module.BASE_URL + "/CarEquipCheck/View/" + state["record_id"]
                state["terminal_view_urls"].append(self.current_url)

            def navigate_control(self, url):
                self.control_navigation = True
                try:
                    self.get(url)
                finally:
                    self.control_navigation = False

            def continue_control(self):
                if ("/CarEquipCheck/View?id=" not in self.current_url or not state["equipment_started"]
                        or state["equipment_done"] or not state["continue_present"]):
                    return None
                return Element("continue", text="繼續清點", displayed=state["continue_displayed"],
                               enabled=state["continue_enabled"])

            def refresh(self):
                self.document_generation += 1
                self.mode = "calendar"

            def find_element(self, by, selector):
                if selector == "html":
                    return Document(self.document_generation)
                if "繼續清點" in selector:
                    control = self.continue_control()
                    if control is None:
                        raise NoSuchElementException("fixture has no continue control")
                    return control
                if selector == "spinner-container":
                    return Element(displayed=False)
                if "tbody/tr[" in selector:
                    return Row()
                if "SaveProc" in selector:
                    return Element("save")
                if "btnCancel" in selector:
                    return Element("return")
                if "完成清點" in selector:
                    return Element("finish")
                if "Select all rows" in selector or "Deselect all rows" in selector:
                    label = "Deselect all rows" if state["selected"] else "Select all rows"
                    if label not in selector:
                        raise NoSuchElementException("fixture header changed its selection label")
                    state["header_lookups"] += 1
                    return Element("checkbox")
                if "保養" in selector:
                    return Element("maintain")
                return Element()

            def find_elements(self, by, selector):
                if "繼續清點" in selector:
                    control = self.continue_control()
                    return [control] if control is not None else []
                if selector == "#scheduler .k-event":
                    return ([Element("event", event_class=value) for value in state["event_classes"]]
                            if state["event_classes"] is not None else [Element("event")])
                if "swal-button" in selector:
                    if not state["dialog_open"]:
                        return []
                    state["dialog_reads"] += 1
                    visible = state["dialog_mode"] != "delayed" or state["dialog_reads"] > 2
                    confirm = Element("confirm", displayed=visible)
                    return ([Element("cancel"), confirm] if "swal-button--confirm" not in selector else [confirm])
                return []

            def execute_script(self, script, *args):
                if "__sinposmartEquipmentDialog" in script:
                    return "saved notice" if state["dialog_mode"] == "two_step" and state["confirm_clicks"] else "confirm save"
                if "__sinposmartEquipmentSelection" in script:
                    if args[0].header_generation != state["header_generation"]:
                        state["stale_header_uses"] += 1
                        raise StaleElementReferenceException("fixture header was replaced by TreeList.refresh")
                    state["selection_reads"] += 1
                    return {"ready": state["selection_reads"] > state["selection_ready_after"],
                            "selected": state["selected"] and not state["selection_incomplete"]}
                if "grid.dataItem" in script:
                    return {"id": state["readback_id"] if state["saved"] and state["readback_id"] is not None
                            else state["record_id"], "vehicle": state["model_vehicle"] or "CAR-01",
                            "date": state["model_date"] or "115/10/17"}
                if "document.readyState" in script:
                    if state["end_take_pending"]:
                        state["end_take_reads"] += 1
                        if state["end_take_reads"] <= state["end_take_ready_after"]:
                            return False
                        self.commit_end_take()
                    return True
                if "dataSource" in script:
                    total = len(state["row_names"]) if state["row_names"] is not None else state["total"]
                    return {"total": total, "page": None if state["no_paging"] else 1,
                            "page_size": None if state["no_paging"] else 20,
                            "filtered": state["filtered"]}
                element = args[0]
                if element.role == "maintain":
                    self.mode = "calendar"
                elif element.role == "return":
                    self.mode = "maintain"
                elif element.role == "check":
                    state["equipment_started"] = True
                    self.navigate_control(module.BASE_URL + "/CarEquipCheck/Take?id=" + state["record_id"])
                elif element.role == "view":
                    state["view_clicks"] += 1
                    self.navigate_control(module.BASE_URL + "/CarEquipCheck/View?id=" +
                                          (state["view_record_id"] or state["record_id"]))
                elif element.role == "continue":
                    state["continue_clicks"] += 1
                    self.navigate_control(module.BASE_URL + "/CarEquipCheck/Take?id=" + state["record_id"])
                elif element.role == "checkbox":
                    state["selection_clicks"] += 1
                    state["selected"] = not state["selected"]
                    if state["refresh_header_on_select"]:
                        state["header_generation"] += 1
                elif element.role == "finish" and state["dialog_mode"]:
                    state["dialog_open"] = True
                elif element.role == "cancel":
                    state["cancel_clicks"] += 1
                    state["dialog_open"] = False
                elif element.role == "confirm" and state["dialog_mode"] == "stuck":
                    state["confirm_clicks"] += 1
                elif element.role == "confirm" and state["dialog_mode"] == "two_step" and state["confirm_clicks"] == 1:
                    state["confirm_clicks"] += 1
                    state["dialog_open"] = False
                    self.get(module.BASE_URL + "/CarEquipCheck/EndTake?id=" + state["record_id"])
                elif element.role in ("save", "finish", "confirm"):
                    if element.role == "confirm":
                        state["confirm_clicks"] += 1
                        state["dialog_open"] = False
                    state["saved"] += 1
                    if state["persist"] and element.role == "save":
                        state["maintenance_done"] = True
                    if element.role != "save":
                        if state["dialog_mode"] == "two_step" and element.role == "confirm":
                            state["dialog_open"] = True
                        else:
                            self.get(module.BASE_URL + "/CarEquipCheck/EndTake?id=" + state["record_id"])
                return None

        class Wait:
            def until(self, condition):
                for _ in range(6):
                    try:
                        result = condition(driver)
                    except NoSuchElementException:
                        continue
                    if result:
                        return result
                raise module.TimeoutException("fixture condition was not confirmed")

        self_module = module
        driver = Driver()
        return driver, Wait(), state, lambda *_args: (
            [Row(vehicle) for vehicle in state["row_names"]] if state["row_names"] is not None else [Row()])

    def _run_operation(self, operation, report=None, **options):
        module = self.module
        driver, wait, state, query = self._fixture(**options)
        report = report if report is not None else module.new_vehicle_report("2026-10-17")
        with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                mock.patch.object(module, "ensure_requested_date"):
            if operation == "maintenance":
                today = datetime(2026, 10, 17)
                module.process_maintain_checks(driver, wait, module.build_today_strings(today), today, report=report)
            else:
                module.process_equip_checks(driver, wait, report=report)
        return report, state

    def test_calendar_without_today_is_unconfirmed(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "today_not_found"):
            self._run_operation("maintenance", missing_today=True)

    def test_unchecked_is_submitted_and_verified_after_save(self) -> None:
        report, state = self._run_operation("maintenance")
        self.assertEqual(state["saved"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_all_disabled_today_is_explicitly_not_required(self) -> None:
        report, state = self._run_operation("maintenance", event_classes=["k-event disable", "k-event disabled"])
        self.assertEqual(state["saved"], 0)
        self.assertEqual(report["items"], [{"vehicle": "CAR-01", "operation": "maintenance",
                                          "status": "not_required", "reason_code": "not_required"}])

    def test_mixed_or_reviewed_disabled_events_never_become_not_required(self) -> None:
        for classes, reviewed in ((["disable", "checked"], False), (["disable", "unchecked"], False),
                                  (["disable unchecked"], False), (["disable reviewed"], False), (["disable"], True)):
            with self.subTest(classes=classes, reviewed=reviewed), self.assertRaises(RuntimeError):
                self._run_operation("maintenance", event_classes=classes, reviewed=reviewed)

    def test_live_equipment_fields_prove_completion_without_literal_complete(self) -> None:
        report, state = self._run_operation("equipment", equipment_done=True)
        self.assertEqual(state["saved"], 0)
        self.assertEqual(report["items"][0]["status"], "already_done")
        for changes in ({"checker": "-"}, {"start": "-"}, {"end": "-"}, {"start": "99:40"}):
            with self.subTest(changes=changes), self.assertRaises(RuntimeError):
                self._run_operation("equipment", equipment_done=True, **changes)

    def test_view_only_equipment_missing_end_time_reports_precise_diagnostic_without_saving(self) -> None:
        module = self.module
        for missing_end in ("", "-", "—", "－"):
            with self.subTest(missing_end=missing_end):
                driver, wait, state, query = self._fixture(equipment_done=True, end=missing_end, record_id="")
                driver.execute_script = mock.Mock(wraps=driver.execute_script)
                report = module.new_vehicle_report("2026-10-17")
                with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                        mock.patch.object(module, "ensure_requested_date"), \
                        mock.patch("builtins.print") as output:
                    with self.assertRaisesRegex(module.VehicleVerificationError, "state_unconfirmed"):
                        module.process_equip_checks(driver, wait, report=report)

                diagnostics = [call.args[0] for call in output.call_args_list if call.args
                               and str(call.args[0]).startswith("[sinposmart-vehicle-diagnostic] ")]
                self.assertEqual(diagnostics, ["[sinposmart-vehicle-diagnostic] equipment_end_time_missing"])
                self.assertEqual(report["reason_code"], "state_unconfirmed")
                self.assertEqual(report["items"], [{"vehicle": "CAR-01", "operation": "equipment",
                                                   "status": "unconfirmed", "reason_code": "state_unconfirmed"}])
                self.assertEqual(state["saved"], 0)
                self.assertFalse(any(call.args[0] == "arguments[0].click();"
                                     for call in driver.execute_script.call_args_list))

    def test_missing_equipment_end_time_diagnostic_does_not_mask_other_states(self) -> None:
        cases = (
            {"equipment_done": True, "end": "99:40"},
            {"equipment_done": False, "equipment_button": False, "end": "-"},
            {"equipment_done": True, "contradictory_button": True, "end": "-"},
        )
        for changes in cases:
            with self.subTest(changes=changes), mock.patch("builtins.print") as output:
                with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
                    self._run_operation("equipment", **changes)
                self.assertFalse(any("equipment_end_time_missing" in str(call.args[0])
                                     for call in output.call_args_list if call.args))

    def test_missing_equipment_end_time_diagnostic_has_actionable_safe_error(self) -> None:
        service = package_module("app_core.daily_vehicle_service").DailyVehicleService
        output = "[sinposmart-vehicle-diagnostic] equipment_end_time_missing\n"
        self.assertEqual(service._safe_error(output),
                         "器材清點尚未結束，無法核對原始紀錄，請至公務網站檢查。")
        for code, expected in {
            "equipment_continue_unavailable": "器材清點頁面未提供可用的「繼續清點」，已停止後續操作。",
            "equipment_detail_unconfirmed": "器材清點明細尚未完整載入或紀錄不符，已停止送出。",
            "equipment_selection_unconfirmed": "器材清點未能確認所有項目已勾選，已停止送出。",
            "equipment_confirmation_unconfirmed": "器材清點尚未確認儲存結果，請至公務網站檢查。",
        }.items():
            with self.subTest(code=code):
                self.assertEqual(service._safe_error(f"[sinposmart-vehicle-diagnostic] {code}\n"), expected)

    def test_not_required_is_maintenance_only_and_can_complete_full_report(self) -> None:
        module = self.module
        report, _state = self._run_operation("maintenance", event_class="disable")
        self._run_operation("equipment", report, equipment_done=True)
        self.assertTrue(module.verify_vehicle_report(report))
        report["reason_code"] = "not_required"
        self.assertFalse(module.verify_vehicle_report(report))
        report["reason_code"] = ""
        report["items"][1].update(status="not_required", reason_code="not_required")
        self.assertFalse(module.verify_vehicle_report(report))

    def test_save_without_persisted_completion_is_unconfirmed(self) -> None:
        for operation in ("maintenance", "equipment"):
            with self.subTest(operation=operation), self.assertRaisesRegex(RuntimeError, "save_unconfirmed"):
                self._run_operation(operation, persist=False)

    def test_placeholder_paging_filter_and_changed_inventory_cannot_complete(self) -> None:
        for options in ({"placeholder": True}, {"total": 2}, {"filtered": True}, {"change_after_save": True}):
            with self.subTest(options=options), self.assertRaises(RuntimeError):
                self._run_operation("maintenance", **options)

    def test_missing_equipment_button_needs_positive_today_evidence(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
            self._run_operation("equipment", equipment_button=False)
        report, state = self._run_operation("equipment", equipment_done=True)
        self.assertEqual(state["saved"], 0)
        self.assertEqual(report["items"][0]["status"], "already_done")

    def test_equipment_save_is_verified_from_fresh_list(self) -> None:
        report, state = self._run_operation("equipment")
        self.assertEqual(state["saved"], 1)
        self.assertEqual(state["terminal_view_urls"], [self.module.BASE_URL + "/CarEquipCheck/View/101"])
        self.assertEqual(report["items"][0]["reason_code"], "saved_verified")

    def test_started_equipment_continues_original_record_and_verifies_save(self) -> None:
        report, state = self._run_operation("equipment", equipment_started=True)
        self.assertEqual((state["view_clicks"], state["continue_clicks"], state["direct_take_gets"]), (1, 1, 0))
        self.assertEqual(state["view_ids"], ["101"])
        self.assertEqual(state["take_ids"], ["101"])
        self.assertEqual(state["saved"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_started_equipment_without_checker_can_continue_original_record(self) -> None:
        report, state = self._run_operation("equipment", equipment_started=True, checker="-")
        self.assertEqual((state["view_clicks"], state["continue_clicks"], state["direct_take_gets"]), (1, 1, 0))
        self.assertEqual(state["take_ids"], ["101"])
        self.assertEqual(state["saved"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_started_equipment_requires_visible_enabled_continue_before_take(self) -> None:
        module = self.module
        for changes in ({"continue_present": False}, {"continue_displayed": False}, {"continue_enabled": False}):
            with self.subTest(changes=changes):
                driver, wait, state, query = self._fixture(equipment_started=True, **changes)
                report = module.new_vehicle_report("2026-10-17")
                with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                        mock.patch.object(module, "ensure_requested_date"), mock.patch("builtins.print") as output:
                    with self.assertRaisesRegex(module.VehicleVerificationError, "state_unconfirmed"):
                        module.process_equip_checks(driver, wait, report)
                self.assertEqual((state["view_clicks"], state["continue_clicks"], state["saved"]), (1, 0, 0))
                self.assertEqual(state["take_ids"], [])
                self.assertEqual(state["direct_take_gets"], 0)
                self.assertIn("[sinposmart-vehicle-diagnostic] equipment_continue_unavailable",
                              [call.args[0] for call in output.call_args_list if call.args])

    def test_started_equipment_completed_during_view_never_reenters_take(self) -> None:
        module = self.module
        driver, wait, state, query = self._fixture(equipment_started=True, complete_on_view=True)
        with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                mock.patch.object(module, "ensure_requested_date"):
            with self.assertRaisesRegex(module.VehicleVerificationError, "state_unconfirmed"):
                module.process_equip_checks(driver, wait, module.new_vehicle_report("2026-10-17"))
        self.assertTrue(state["equipment_done"])
        self.assertEqual((state["view_clicks"], state["continue_clicks"], state["saved"]), (1, 0, 0))
        self.assertEqual(state["take_ids"], [])
        self.assertEqual(state["direct_take_gets"], 0)

    def test_started_equipment_view_requires_original_record_before_continue(self) -> None:
        module = self.module
        driver, wait, state, query = self._fixture(equipment_started=True, view_record_id="102")
        with mock.patch.object(module, "query_grid_rows", side_effect=query), \
                mock.patch.object(module, "ensure_requested_date"):
            with self.assertRaisesRegex(module.VehicleVerificationError, "state_unconfirmed"):
                module.process_equip_checks(driver, wait, module.new_vehicle_report("2026-10-17"))
        self.assertEqual(state["view_ids"], ["102"])
        self.assertEqual((state["continue_clicks"], state["saved"]), (0, 0))
        self.assertEqual(state["take_ids"], [])
        self.assertEqual(state["direct_take_gets"], 0)

    def test_first_equipment_confirmation_never_clicks_cancel(self) -> None:
        report, state = self._run_operation("equipment", dialog_mode="cancel_first")
        self.assertEqual((state["cancel_clicks"], state["confirm_clicks"], state["saved"]), (0, 1, 1))
        self.assertEqual(report["items"][0]["reason_code"], "saved_verified")

    def test_first_equipment_waits_for_visible_confirmation(self) -> None:
        report, state = self._run_operation("equipment", dialog_mode="delayed")
        self.assertEqual(state["confirm_clicks"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_equipment_confirmation_then_saved_notice_is_acknowledged_once_each(self) -> None:
        report, state = self._run_operation("equipment", dialog_mode="two_step")
        self.assertEqual((state["saved"], state["confirm_clicks"], state["cancel_clicks"]), (1, 2, 0))
        self.assertEqual(report["items"][0]["reason_code"], "saved_verified")

    def test_equipment_same_visible_confirmation_is_never_clicked_twice(self) -> None:
        driver, wait, state, query = self._fixture(dialog_mode="stuck")
        with mock.patch.object(self.module, "query_grid_rows", side_effect=query), \
                mock.patch.object(self.module, "ensure_requested_date"):
            with self.assertRaisesRegex(RuntimeError, "save_unconfirmed"):
                self.module.process_equip_checks(driver, wait, self.module.new_vehicle_report("2026-10-17"))
        self.assertEqual((state["confirm_clicks"], state["cancel_clicks"], state["saved"]), (1, 0, 0))

    def test_equipment_already_selected_header_is_not_toggled_off(self) -> None:
        _report, state = self._run_operation("equipment", selected=True)
        self.assertEqual(state["selection_clicks"], 0)
        self.assertTrue(state["selected"])

    def test_equipment_waits_for_loaded_details_before_selection(self) -> None:
        _report, state = self._run_operation("equipment", selection_ready_after=2)
        self.assertGreaterEqual(state["selection_reads"], 3)
        self.assertEqual(state["saved"], 1)

    def test_equipment_reacquires_header_after_selection_refresh(self) -> None:
        report, state = self._run_operation("equipment", refresh_header_on_select=True)
        self.assertGreaterEqual(state["header_lookups"], 2)
        self.assertEqual(state["stale_header_uses"], 0)
        self.assertEqual(state["selection_clicks"], 1)
        self.assertEqual(state["saved"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_equipment_waits_for_end_take_commit_before_returning_to_index(self) -> None:
        report, state = self._run_operation("equipment", end_take_ready_after=2)
        self.assertGreaterEqual(state["end_take_reads"], 3)
        self.assertFalse(state["end_take_interrupted"])
        self.assertEqual(state["saved"], 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_equipment_selection_requires_full_bound_rows_and_checked_details(self) -> None:
        javascript = """
let records = [], rows = [], total = 0;
const box = {checked:false, indeterminate:false};
const checks = [{checked:false, indeterminate:false}, {checked:false, indeterminate:false}];
const grid = {dataSource:{data:()=>records,total:()=>total},tbody:{find:()=>rows}};
const host = {data:name=>name==='kendoTreeList'?grid:null};
var window = {jQuery:()=>({closest:()=>host})}; window.jQuery.active=0;
function snapshot() { return (function() {
""" + self.module.EQUIPMENT_SELECTION_SCRIPT + """
}).apply(null,[box]); }
const empty = snapshot();
records = [{uid:'a',Checked:'N'},{uid:'b',Checked:'N'}]; total=2;
rows = ['a','b'].map((uid,index)=>({getAttribute:()=>uid,querySelectorAll:()=>[checks[index]]}));
const loaded = snapshot(); box.checked=true; checks[0].checked=true;
const partial = snapshot(); checks[1].checked=true; const domOnly = snapshot();
records.forEach(record=>record.Checked='Y'); const selected = snapshot();
records[1].Checked=true; const unknownModel = snapshot(); records[1].Checked='Y';
box.indeterminate=true; const indeterminate = snapshot(); box.indeterminate=false;
rows[1]={getAttribute:()=>'stale',querySelectorAll:()=>[checks[1]]}; const stale=snapshot();
rows[1]={getAttribute:()=>'b',querySelectorAll:()=>[checks[1]]}; total=3; const paged=snapshot();
JSON.stringify({empty,loaded,partial,domOnly,selected,unknownModel,indeterminate,stale,paged});
"""
        runner = ("from PySide6.QtCore import QCoreApplication\nfrom PySide6.QtQml import QJSEngine\n"
                  "app=QCoreApplication([])\nengine=QJSEngine()\nresult=engine.evaluate(" + repr(javascript) + ")\n"
                  "assert not result.isError(),result.toString()\nprint(result.toString())\n")
        result = subprocess.run([sys.executable, "-B", "-c", runner], capture_output=True,
                                text=True, encoding="utf-8", timeout=15,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {
            "empty": None, "loaded": {"ready": True, "selected": False},
            "domOnly": {"ready": True, "selected": False}, "unknownModel": {"ready": True, "selected": False},
            "partial": {"ready": True, "selected": False}, "selected": {"ready": True, "selected": True},
            "indeterminate": {"ready": True, "selected": False}, "stale": None, "paged": None,
        })

    def test_equipment_incomplete_selection_never_submits(self) -> None:
        driver, wait, state, query = self._fixture(selection_incomplete=True)
        with mock.patch.object(self.module, "query_grid_rows", side_effect=query), \
                mock.patch.object(self.module, "ensure_requested_date"):
            with self.assertRaises(RuntimeError):
                self.module.process_equip_checks(driver, wait, self.module.new_vehicle_report("2026-10-17"))
        self.assertEqual(state["saved"], 0)

    def test_equipment_continuation_rejects_wrong_date_vehicle_or_record(self) -> None:
        for options in ({"equipment_text": "115/10/16"}, {"model_date": "115/10/16"},
                        {"model_vehicle": "CAR-02"}, {"record_id": "../Duplicate/101"},
                        {"start": "99:40"}):
            with self.subTest(options=options), self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
                self._run_operation("equipment", equipment_started=True, **options)

    def test_equipment_save_requires_original_record_id_on_readback(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "save_unconfirmed"):
            self._run_operation("equipment", readback_id="102")

    def _equipment_inventory(self, records):
        module = self.module
        rows = []
        for record in records:
            vehicle, started = record.get("vehicle", "CAR-01"), record.get("start", "08:40")
            completed = record.get("completed", True)
            texts = ["", "", "", vehicle, record.get("date", "115/10/17"), "-",
                     record.get("checker", "同仁" if completed else "-"), started,
                     record.get("end", "08:41" if completed else "-"), ""]
            cells = []
            for text in texts:
                cell = mock.Mock()
                cell.get_attribute.side_effect = lambda name, value=text: value if name == "textContent" else ""
                cell.find_elements.return_value = []
                cells.append(cell)
            view = mock.Mock()
            view.is_displayed.return_value = view.is_enabled.return_value = True
            row = mock.Mock()
            row.record = {"id": record["id"], "vehicle": vehicle, "date": texts[4]}
            row.find_elements.side_effect = lambda by, selector, values=cells, button=view: (
                values if selector == "td" else [button] if selector == ".k-grid-view" else [])
            rows.append(row)
        driver = mock.Mock()
        driver.execute_script.side_effect = lambda script, *args: (
            {"total": len(rows), "page": 1, "page_size": 20, "filtered": False}
            if script == module.GRID_STATE_SCRIPT else dict(args[0].record))
        report = module.new_vehicle_report("2026-10-17")
        with mock.patch.object(module, "query_grid_rows", return_value=rows):
            inventory = module.verified_grid_rows(driver, mock.Mock(), "equipment", report)
        return inventory, rows, report

    def test_equipment_inventory_selects_only_unfinished_record_with_completed_history(self) -> None:
        inventory, rows, report = self._equipment_inventory([
            {"id": "101"}, {"id": "102", "completed": False},
        ])
        self.assertEqual(set(inventory), {"CAR-01"})
        self.assertIs(inventory["CAR-01"][1], rows[1])
        self.assertEqual(report["items"], [])

    def test_equipment_inventory_all_completed_uses_latest_numeric_record(self) -> None:
        inventory, rows, report = self._equipment_inventory([{"id": "101"}, {"id": "9"}])
        self.assertEqual(set(inventory), {"CAR-01"})
        self.assertIs(inventory["CAR-01"][1], rows[0])
        self.assertEqual(report["items"], [])

    def test_equipment_inventory_rejects_ambiguous_or_mismatched_duplicates(self) -> None:
        cases = (
            [{"id": "101", "completed": False}, {"id": "102", "completed": False}],
            [{"id": "101", "completed": False}, {"id": "101", "vehicle": "CAR-02", "completed": False}],
            [{"id": "../Duplicate/101"}],
            [{"id": "101", "completed": False}, {"id": "102"}],
            [{"id": "101"}, {"id": "101"}],
            [{"id": "101", "date": "115/10/16"}, {"id": "102", "completed": False}],
            [{"id": "101"}, {"id": "102", "vehicle": "car-01"}],
            [{"id": "101", "checker": "-"}, {"id": "102", "completed": False}],
            [{"id": "../Duplicate/101"}, {"id": "102", "completed": False}],
        )
        for records in cases:
            with self.subTest(records=records), self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
                self._equipment_inventory(records)

    def test_today_completed_text_with_pending_button_is_unconfirmed(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
            self._run_operation("equipment", equipment_done=True, contradictory_button=True)

    def test_loaded_grid_without_paging_is_accepted(self) -> None:
        report, _state = self._run_operation("maintenance", maintenance_done=True, no_paging=True)
        self.assertEqual(report["items"][0]["status"], "already_done")

    def test_duplicate_vehicle_case_cannot_complete(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "state_unconfirmed"):
            self._run_operation("maintenance", maintenance_done=True, row_names=["CAR-01", "car-01"])

    def test_vehicle_names_match_protocol_and_casefold_upsert(self) -> None:
        module = self.module
        self.assertEqual(module.safe_vehicle_name("トヨタ"), "")
        self.assertEqual(module.safe_vehicle_name("  新坡  91  "), "新坡 91")
        report = module.new_vehicle_report("2026-10-17")
        module.record_vehicle_result(report, "CAR-01", "maintenance", "unconfirmed", "save_unconfirmed")
        module.record_vehicle_result(report, "car-01", "maintenance", "completed", "saved_verified")
        self.assertEqual(len(report["items"]), 1)
        self.assertEqual(report["items"][0]["status"], "completed")

    def test_negative_date_status_and_button_text_cannot_be_completion(self) -> None:
        module = self.module
        today_strings = module.build_today_strings(datetime(2026, 10, 17))
        for status in ("未完成", "尚未完成", "待完成", "未清點", "待清點", "未知狀態"):
            with self.subTest(status=status):
                self.assertFalse(module.is_task_done("115/10/17 " + status, today_strings))
                with self.assertRaises(RuntimeError):
                    self._run_operation("equipment", equipment_button=False,
                                        equipment_text="115/10/17 " + status)
        with self.assertRaises(RuntimeError):
            self._run_operation("equipment", equipment_button=False,
                                equipment_text="115/10/17 完成清點", button_only_evidence=True)

    def test_positive_report_requires_both_operations_and_safe_unique_rows(self) -> None:
        module = self.module
        report, _ = self._run_operation("maintenance", maintenance_done=True)
        self.assertFalse(module.verify_vehicle_report(report))
        self._run_operation("equipment", report, equipment_done=True)
        self.assertTrue(module.verify_vehicle_report(report))
        self.assertEqual(len(report["items"]), 2)
        module.record_vehicle_result(report, "CAR-01", "maintenance", "already_done", "verified_today")
        self.assertEqual(len(report["items"]), 2)
        for name in ("account-secret", "帳號測試", "CAR-01\nTOKEN", "CAR-01\n", "https://fixture.invalid", "A" * 33):
            with self.subTest(name=name):
                self.assertEqual(module.safe_vehicle_name(name), "")
        report = module.new_vehicle_report("2026-10-17")
        for index in range(200):
            module.record_vehicle_result(report, f"CAR-{index}", "maintenance", "completed", "saved_verified")
        with self.assertRaisesRegex(RuntimeError, "row_limit"):
            module.record_vehicle_result(report, "OVER-LIMIT", "equipment", "completed", "saved_verified")
        self.assertEqual(len(report["items"]), 200)

    def _run_main(self, *, keep_browser=False, configuration_error=False, **options):
        from contextlib import ExitStack, redirect_stdout
        from io import StringIO
        module = self.module
        driver, wait, state, query = self._fixture(**options)

        class FixedDateTime(datetime):
            @classmethod
            def now(cls):
                return cls(2026, 10, 17, 9)

        configuration = {"username": "fixture-user", "password": "fixture-password",
                         "headless": False, "keep_browser_open": keep_browser,
                         "timeout_seconds": 1, "selenium_remote_url": ""}
        output = StringIO()
        error = None
        with ExitStack() as stack:
            stack.enter_context(mock.patch.object(module.dt, "datetime", FixedDateTime))
            stack.enter_context(mock.patch.dict(module.os.environ, {"PPE_TARGET_DATE": "2026-10-17"}))
            stack.enter_context(mock.patch.object(module, "query_grid_rows", side_effect=query))
            stack.enter_context(mock.patch.object(module, "ensure_requested_date"))
            stack.enter_context(mock.patch.object(module, "load_config", side_effect=RuntimeError("token=fixture-secret"))
                                if configuration_error else mock.patch.object(module, "load_config", return_value=configuration))
            stack.enter_context(mock.patch.object(module.webdriver, "ChromeOptions"))
            stack.enter_context(mock.patch.object(module, "retry_duty_browser_session_open", return_value=driver))
            stack.enter_context(mock.patch.object(module, "WebDriverWait", return_value=wait))
            cleanup = stack.enter_context(mock.patch.object(module, "quit_driver"))
            sleep = stack.enter_context(mock.patch.object(module.time, "sleep"))
            stack.enter_context(redirect_stdout(output))
            try:
                module.main([])
            except RuntimeError as exception:
                error = exception
        lines = output.getvalue().splitlines()
        reports = [json.loads(line[len(module.REPORT_PREFIX):]) for line in lines if line.startswith(module.REPORT_PREFIX)]
        return reports, lines, error, cleanup, sleep

    def test_main_emits_one_verified_report_before_done_and_retains_browser(self) -> None:
        reports, lines, error, cleanup, sleep = self._run_main(
            keep_browser=True, maintenance_done=True, equipment_done=True)
        self.assertIsNone(error)
        self.assertEqual(len(reports), 1)
        self.assertTrue(reports[0]["verified"])
        self.assertEqual(reports[0]["target_date"], "2026-10-17")
        report_line = next(index for index, line in enumerate(lines) if line.startswith(self.module.REPORT_PREFIX))
        self.assertLess(report_line, lines.index("[done] automation finished"))
        sleep.assert_called_once_with(self.module.DEFAULT_BROWSER_RETENTION_SECONDS)
        cleanup.assert_called_once()

    def test_main_failure_preserves_verified_items_without_done(self) -> None:
        reports, lines, error, cleanup, _sleep = self._run_main(maintenance_done=True, persist=False)
        self.assertIsNotNone(error)
        self.assertEqual(len(reports), 1)
        self.assertFalse(reports[0]["verified"])
        self.assertEqual([item["status"] for item in reports[0]["items"]], ["already_done", "unconfirmed"])
        self.assertNotIn("[done] automation finished", lines)
        cleanup.assert_called_once()

    def test_configuration_failure_emits_safe_unconfirmed_report(self) -> None:
        reports, lines, error, cleanup, _sleep = self._run_main(configuration_error=True)
        self.assertIsNotNone(error)
        self.assertEqual(len(reports), 1)
        self.assertFalse(reports[0]["verified"])
        self.assertEqual(reports[0]["reason_code"], "unexpected_error")
        self.assertNotIn("fixture-secret", "\n".join(lines) + str(error))
        self.assertNotIn("[done] automation finished", lines)
        cleanup.assert_not_called()


class PackageSmokeTests(unittest.TestCase):
    def test_retained_browser_is_scheduled_to_close_after_ten_minutes(self) -> None:
        module = duty_rehearsal_module()
        module.RETAINED_BROWSER_DRIVERS.clear()
        driver = mock.Mock()
        with mock.patch.object(module.threading, "Timer") as timer_factory:
            module.retain_browser_for_seconds(driver)

        timer_factory.assert_called_once()
        self.assertEqual(timer_factory.call_args.args[0], 600)
        timer = timer_factory.return_value
        self.assertTrue(timer.daemon)
        timer.start.assert_called_once()
        timer_factory.call_args.args[1](*timer_factory.call_args.kwargs["args"])
        driver.quit.assert_called_once_with()
        self.assertNotIn(driver, module.RETAINED_BROWSER_DRIVERS)

    def test_duty_driver_closes_unless_verified_success_is_retained(self) -> None:
        module = legacy_duty_sheet_module()

        self.assertTrue(module.should_close_duty_driver(True, True, False))
        self.assertFalse(module.should_close_duty_driver(True, True, True))
        self.assertTrue(module.should_close_duty_driver(True, False, True))
        self.assertFalse(module.should_close_duty_driver(False, False, False))

    def test_updater_busy_handshake_wait_is_bounded_and_preserves_process_identity(self) -> None:
        source = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        self.assertIn("function Wait-UpdateLogoutEvent", source)
        helper = source[source.index("function Wait-UpdateLogoutEvent"):source.index("function Stop-RunningDutyGui")]
        for responses, identities, timeout, expected, calls in (
            (["busy", "busy", "ready"], [123, 123, 123], 300, "ready", 3),
            (["busy"], [123], 0, "busy_timeout", 1),
            (["failed"], [123], 300, "failed", 1),
            (["timeout"], [123], 300, "timeout", 1),
            (["busy", "ready"], [123, 456], 300, "failed", 1),
        ):
            with self.subTest(responses=responses, identities=identities):
                replies = ",".join(f"'{value}'" for value in responses)
                pids = ",".join(str(value) for value in identities)
                result = run_powershell_contract(f"""
$ErrorActionPreference = 'Stop'
{helper}
$script:responses = @({replies})
$script:identities = @({pids})
$script:calls = 0
$script:queries = 0
function Send-UpdateLogoutEvent {{ $script:calls++; return $script:responses[$script:calls - 1] }}
function Get-RunningDutyGuiProcesses {{
    $script:queries++
    return @([pscustomobject]@{{ ProcessId = $script:identities[$script:queries - 1] }})
}}
function Test-IsQtDutyGuiProcess {{ return $true }}
function Write-UpdateProgress {{ }}
function Start-Sleep {{ }}
$result = Wait-UpdateLogoutEvent -ExpectedProcessId 123 -TimeoutSeconds {timeout}
if ($result -ne '{expected}' -or $script:calls -ne {calls}) {{ throw "Unexpected result: $result, calls: $script:calls" }}
""" )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_hidden_process_streams_both_pipes_and_times_out(self):
        import base64
        source = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        hidden = source[source.index("function Invoke-HiddenProcess"):source.index("function Start-DutyGui")]
        for command, timeout, expected in (("[Console]::Out.WriteLine('first'); [Console]::Error.WriteLine('second')", 10, "ok"), ("Start-Sleep -Seconds 30", 1, "timed out")):
            encoded = base64.b64encode(command.encode("utf-16-le")).decode()
            result = run_powershell_contract(f"""
$ErrorActionPreference='Stop'
$packageDir=$env:TEMP
{hidden}
function Write-UpdateProgress {{ }}
try {{
$r=Invoke-HiddenProcess -FileName powershell -Arguments '-NoProfile -EncodedCommand {encoded}' -TimeoutSeconds {timeout} -ProgressPhase setup
if ($r.Output -notmatch 'first' -or $r.Error -notmatch 'second') {{ throw 'missing output' }}
'ok'
}} catch {{ $_.Exception.Message }}
""")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn(expected, result.stdout)

    def test_hidden_python_discovery_preserves_chinese_paths(self) -> None:
        source = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        hidden = source[source.index("function Invoke-HiddenProcess"):source.index("function Start-DutyGui")]
        finder = source[source.index("function Get-WinPythonExe"):source.index("function Invoke-SetupAfterUpdate")]
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            expected = r"C:\更新測試 WinPython\python.exe"
            (root / "find_winpython.ps1").write_text(f"Write-Output '{expected}'\n", encoding="utf-8-sig")
            result = run_powershell_contract(
                f"$ErrorActionPreference='Stop'\n$packageDir='{root}'\n{hidden}\n{finder}\n"
                f"$found=Get-WinPythonExe\nif ($found -ne '{expected}') {{ throw 'Unicode path was corrupted.' }}\n'ok'\n"
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_update_progress_full_install_uses_only_local_fixtures(self) -> None:
        import hashlib
        import http.server
        import io
        import re
        import zipfile

        source = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        required = source.split("$requiredQtPackageFiles = @(", 1)[1].split("\n    )", 1)[0]
        paths = re.findall(r'"([^"\n]+)"', required)
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_STORED) as archive:
            for path in paths:
                archive.writestr(path.replace("\\", "/"), "# isolated fixture\n")
            archive.writestr("VERSION.txt", "2026.09.13.2359")
            archive.writestr("fixture-marker.txt", "installed fixture")
            archive.writestr("work_log_defaults.json", '"must not replace"')
        zip_bytes = zip_buffer.getvalue()
        routes = {
            "/version": b"2026.09.13.2359",
            "/sha": hashlib.sha256(zip_bytes).hexdigest().encode("ascii"),
            "/package": zip_bytes,
        }

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                payload = routes[self.path]
                self.send_response(200)
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            def log_message(self, *_args):
                pass

        server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                root = Path(temp_dir)
                installed = root / "installed"
                installed.mkdir()
                (installed / "VERSION.txt").write_text("2026.09.12.0001", encoding="utf-8")
                (installed / "work_log_defaults.json").write_text('"preserved fixture"', encoding="utf-8")
                url = f"http://127.0.0.1:{server.server_port}"
                overrides = f"""
$packageDir = '{installed}'
$localVersionPath = Join-Path $packageDir 'VERSION.txt'
$tempDir = '{root / 'stage'}'
$zipPath = Join-Path $tempDir 'package.zip'
$extractDir = Join-Path $tempDir 'extract'
$backupDir = '{root / 'backups'}'
$ProgressPath = '{root / 'progress.json'}'
$remoteVersionUrl = '{url}/version'
$remoteSha256Url = '{url}/sha'
$remoteZipUrl = '{url}/package'
$AssumeYes = $true
$RestartAfterUpdate = $true
function Get-RunningDutyGuiProcesses {{ return @([pscustomobject]@{{ ProcessId = 123 }}) }}
function Test-IsQtDutyGuiProcess {{ return $true }}
$script:prepareCalls = 0
function Send-UpdateLogoutEvent {{
    $script:prepareCalls++
    if (Test-Path -LiteralPath (Join-Path $packageDir 'fixture-marker.txt')) {{ throw 'Copied before ready.' }}
    if ($script:prepareCalls -lt 3) {{ return 'busy' }}
    return 'ready'
}}
function Start-Sleep {{ }}
function Stop-RunningDutyGui {{
    if ($script:prepareCalls -ne 3) {{ throw 'Stopped without waiting for ready.' }}
    return $true
}}
function Invoke-SetupAfterUpdate {{ Write-UpdateProgress -Phase 'setup' -Percent 85 }}
function Start-DutyGui {{ $script:restartedProcessId = 456; return $true }}
"""
                marker = 'Write-UpdateProgress -Phase "checking" -Percent 2\nif'
                prefix, body = source.split(marker, 1)
                contract = prefix + overrides + marker + body
                result = run_powershell_contract(contract)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                self.assertEqual((installed / "fixture-marker.txt").read_text(), "installed fixture")
                self.assertEqual((installed / "VERSION.txt").read_text(encoding="utf-8-sig").strip(), "2026.09.13.2359")
                self.assertEqual((installed / "work_log_defaults.json").read_text(), '"preserved fixture"')
                progress = json.loads((root / "progress.json").read_text(encoding="utf-8-sig"))
                self.assertEqual(progress, {"phase": "installed", "percent": 98, "pid": 456})
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)

    def duty_board_schedule_payload(self) -> dict[str, object]:
        return {
            "target_date": "1150716",
            "today": {
                "roc_date": "1150716",
                "rows": [
                    {"slot": "8-9", "columns": {"值班": ["1", "2"]}},
                    {"slot": "9-10", "columns": {"值班": ["3"]}},
                    {"slot": "23-0", "columns": {"值班": ["4"]}},
                    {"slot": "0-1", "columns": {"值班": ["5"]}},
                ],
                "staff": {
                    "1": {"name": "王小明", "role": "隊員"},
                    "2": {"name": "李小華", "role": "隊員"},
                    "3": {"name": "陳小美", "role": "隊員"},
                    "4": {"name": "林小強", "role": "隊員"},
                    "5": {"name": "周小安", "role": "隊員"},
                },
            },
            "tomorrow": {
                "roc_date": "1150717",
                "rows": [{"slot": "8-9", "columns": {"值班": ["6"]}}],
                "staff": {"6": {"name": "張小雲", "role": "隊員"}},
            },
        }

    def test_package_entry_files_exist(self) -> None:
        root = package_dir()

        for relative in (
            "requirements.txt",
            "duty_gui.pyw",
            "duty_gui.py",
            "RUN_DUTY_GUI_WINPYTHON.bat",
            "update_package.ps1",
        ):
            with self.subTest(relative=relative):
                self.assertTrue((root / relative).is_file())

        windowed_entry = (root / "duty_gui.pyw").read_text(encoding="utf-8-sig")
        self.assertIn("from qt_app.main import main", windowed_entry)
        self.assertNotIn("from duty_gui import main", windowed_entry)

        compatibility_entry = (root / "duty_gui.py").read_text(encoding="utf-8-sig")
        self.assertIn("from qt_app.main import main", compatibility_entry)
        self.assertNotIn("tkinter", compatibility_entry)
        self.assertNotIn("from legacy_tk", compatibility_entry)
        self.assertTrue(legacy_tk_path("duty_gui.py").is_file())

        launcher = (root / "RUN_DUTY_GUI_WINPYTHON.bat").read_text(encoding="utf-8-sig")
        self.assertIn('set "PYTHONW_EXE=%%F"', launcher)
        self.assertIn('start "" /b "%PYTHONW_EXE%" "%~dp0duty_gui.pyw"', launcher)
        self.assertIn("-Windowed", launcher)
        self.assertIn("exit /b 0", launcher)
        self.assertNotIn("System.Diagnostics.ProcessStartInfo", launcher)
        self.assertNotIn("PYTHON_EXE", launcher)

    def test_environment_check_targets_qt_runtime_not_legacy_tk(self) -> None:
        source = (package_dir() / "check_environment.py").read_text(encoding="utf-8-sig")

        self.assertNotIn("import tkinter", source)
        self.assertNotIn('"customtkinter"', source)
        self.assertNotIn('"tkcalendar"', source)
        self.assertNotIn('"pystray"', source)
        self.assertIn("from PySide6.QtQml import QQmlApplicationEngine", source)
        self.assertIn("PySide6/QML runtime imports succeeded", source)

    def test_legacy_rescue_video_tool_launches_without_delete_flags(self) -> None:
        root = package_dir()
        gui_source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertTrue(legacy_tk_path("rescue_video_gui.py").is_file())
        self.assertTrue((root / "rescue_video" / "classify_rescue_video.py").is_file())
        self.assertIn('text="行車紀錄器（BETA）"', gui_source)
        update_source = (root / "update_package.ps1").read_text(encoding="utf-8-sig")
        self.assertIn('"rescue_video\\classify_rescue_video.py"', update_source)
        self.assertIn('"rescue_video\\rescue_video_core.py"', update_source)

        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        messages: list[tuple[str, str]] = []
        gui.notify_user = lambda title, message, **_kwargs: messages.append((title, message))

        with mock.patch.object(module.subprocess, "Popen") as popen:
            gui.open_rescue_video_tool()

        command = popen.call_args.args[0]
        self.assertEqual(Path(command[-1]).name, "rescue_video_gui.py")
        self.assertNotIn("--delete-source", command)
        self.assertNotIn("--apply", command)
        self.assertEqual(messages, [(module.APP_DISPLAY_NAME, "已開啟行車紀錄器（BETA）。")])

    def test_python_sources_compile(self) -> None:
        root = package_dir()
        sources = sorted(list(root.rglob("*.py")) + list(root.rglob("*.pyw")))

        self.assertGreaterEqual(len(sources), 8)
        for path in sources:
            if "__pycache__" in path.parts:
                continue
            with self.subTest(path=str(path.relative_to(PROJECT_ROOT))):
                source = path.read_text(encoding="utf-8-sig")
                compile(source, str(path), "exec")

    def test_background_chrome_options_have_offscreen_fallback(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")
        tree = ast.parse(source)
        helper = next(
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.FunctionDef) and node.name == "background_chrome_options"
        )
        helper_args = [
            node.args[0].value
            for node in ast.walk(helper)
            if isinstance(node, ast.Call)
            and isinstance(node.func, ast.Attribute)
            and node.func.attr == "add_argument"
            and node.args
            and isinstance(node.args[0], ast.Constant)
        ]

        self.assertIn("def background_chrome_options() -> Options:", source)
        self.assertIn("--headless=new", helper_args)
        self.assertIn("--window-size=1280,900", helper_args)
        self.assertIn("--window-position=-32000,-32000", helper_args)
        self.assertNotIn("--disable-popup-blocking", helper_args)
        self.assertGreaterEqual(source.count("background_chrome_options()"), 5)

    def test_readonly_background_queries_suppress_window_open(self) -> None:
        source = (package_dir() / "duty_rehearsal.py").read_text(encoding="utf-8-sig")

        self.assertIn("def suppress_window_open_for_background_query", source)
        self.assertGreaterEqual(source.count("suppress_window_open_for_background_query(driver)"), 3)

    def test_update_package_excludes_sensitive_and_runtime_files(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")

        for expected in (
            '".env"',
            '".jsonl"',
            '"logs"',
            '"runtime_outputs"',
            '".zip"',
            '".token"',
        ):
            with self.subTest(expected=expected):
                self.assertIn(expected, script)

    def test_update_package_migrates_legacy_tk_fallback(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        skip_dirs = script[
            script.index("$skipDirs = @("):
            script.index("$alwaysSkipFiles = @(")
        ]

        self.assertNotIn('"legacy_tk"', skip_dirs)
        self.assertIn("function Move-LegacyTkFallbacks", script)
        self.assertIn("Move-LegacyTkFallbacks -PackageDir $packageDir", script)
        for old_path in (
            '"duty_gui.py"',
            '"duty_sheet_automation.py"',
            '"daily_vehicle_automation.py"',
            '"rescue_video\\救護影片分類GUI.py"',
        ):
            with self.subTest(old_path=old_path):
                self.assertIn(old_path, script)

    def test_update_package_moves_legacy_tk_files_without_overwriting_fallback(self) -> None:
        contract = r'''
$updatePath = Join-Path (Get-Location) "update_package.ps1"
$updateScript = Get-Content -LiteralPath $updatePath -Raw -Encoding UTF8
$functionBlock = [regex]::Match(
    $updateScript,
    '(?s)function Test-LegacyTkScript\s*\{.*?(?=function New-PackageBackup)'
)
if (-not $functionBlock.Success) {
    throw "Could not load legacy Tk migration functions."
}

$stamp = "migration-test"
Invoke-Expression $functionBlock.Value
$packageDir = Join-Path $env:TEMP ("SinpoSmartLegacyTkTest-" + [guid]::NewGuid().ToString("N"))

try {
    New-Item -ItemType Directory -Path $packageDir -Force | Out-Null
    $rescueDir = Join-Path $packageDir "rescue_video"
    $legacyDir = Join-Path $packageDir "legacy_tk"
    New-Item -ItemType Directory -Path $rescueDir -Force | Out-Null
    New-Item -ItemType Directory -Path $legacyDir -Force | Out-Null

    Set-Content -LiteralPath (Join-Path $packageDir "duty_gui.py") -Value "import tkinter as tk" -Encoding UTF8
    Set-Content -LiteralPath (Join-Path $packageDir "duty_sheet_automation.py") -Value "from tkinter import messagebox" -Encoding UTF8
    Set-Content -LiteralPath (Join-Path $packageDir "daily_vehicle_automation.py") -Value "# QML-compatible source" -Encoding UTF8
    Set-Content -LiteralPath (Join-Path $rescueDir "救護影片分類GUI.py") -Value "import customtkinter" -Encoding UTF8
    Set-Content -LiteralPath (Join-Path $legacyDir "duty_sheet_automation.py") -Value "existing fallback" -Encoding UTF8

    Move-LegacyTkFallbacks -PackageDir $packageDir 6>$null

    if (Test-Path -LiteralPath (Join-Path $packageDir "duty_gui.py")) {
        throw "Tk duty GUI was not moved."
    }
    if (Test-Path -LiteralPath (Join-Path $packageDir "duty_sheet_automation.py")) {
        throw "Tk duty sheet automation was not moved."
    }
    if (-not (Test-Path -LiteralPath (Join-Path $packageDir "daily_vehicle_automation.py"))) {
        throw "Non-Tk source should remain at its root path."
    }
    if (Test-Path -LiteralPath (Join-Path $rescueDir "救護影片分類GUI.py")) {
        throw "Tk rescue video GUI was not moved."
    }
    if (-not (Test-Path -LiteralPath (Join-Path $legacyDir "duty_gui.py"))) {
        throw "Tk duty GUI fallback is missing."
    }
    if (-not (Test-Path -LiteralPath (Join-Path $legacyDir "rescue_video_gui.py"))) {
        throw "Tk rescue video fallback is missing."
    }
    if ((Get-Content -LiteralPath (Join-Path $legacyDir "duty_sheet_automation.py") -Raw -Encoding UTF8).Trim() -ne "existing fallback") {
        throw "Existing Tk fallback was overwritten."
    }
    $migrationBackup = Join-Path $legacyDir "migration_backup\\migration-test\\duty_sheet_automation.py"
    if (-not (Test-Path -LiteralPath $migrationBackup)) {
        throw "Existing Tk fallback was not preserved in migration backup."
    }
    if ((Get-Content -LiteralPath $migrationBackup -Raw -Encoding UTF8) -notmatch "from tkinter") {
        throw "Migration backup content is incorrect."
    }
} finally {
    if (Test-Path -LiteralPath $packageDir) {
        Remove-Item -LiteralPath $packageDir -Recurse -Force
    }
}
'''

        result = run_powershell_contract(contract)

        self.assertEqual(result.returncode, 0, msg=(result.stdout or "") + (result.stderr or ""))

    def test_update_package_preserves_existing_work_log_settings(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        preserve_section = script[
            script.index("$preserveIfExistsFiles = @("):
            script.index("$skipExtensions = @(")
        ]

        self.assertIn('"work_log_defaults.json"', preserve_section)

    def test_sinposmart_events_include_installed_version(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn("def current_app_version(", source)
        self.assertIn("snapshot_data = sanitize_frontend_json(dict(snapshot or {}))", source)
        self.assertIn('snapshot_data.setdefault("app_version", current_app_version())', source)
        self.assertIn('"snapshot": snapshot_data', source)

    def test_sinposmart_backend_event_call_keywords_match_signature(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")
        tree = ast.parse(source)
        signature_keywords: set[str] = set()
        call_keywords: set[str] = set()

        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef) and node.name == "send_sinposmart_backend_event":
                signature_keywords = {arg.arg for arg in node.args.args[2:]}
                signature_keywords.update(arg.arg for arg in node.args.kwonlyargs)
            if not isinstance(node, ast.Call):
                continue
            func = node.func
            if isinstance(func, ast.Attribute) and func.attr == "send_sinposmart_backend_event":
                call_keywords.update(keyword.arg for keyword in node.keywords if keyword.arg)

        self.assertIn("send_sinposmart_backend_event", source)
        self.assertFalse(call_keywords - signature_keywords)

    def test_frontend_error_payload_sanitizes_selenium_and_sensitive_details(self) -> None:
        module = duty_gui_module()
        from selenium.common.exceptions import NoSuchElementException, TimeoutException

        cases = (
            (
                TimeoutException("Timed out receiving message from renderer\nStacktrace:\nChromeDriver"),
                "timeout",
                "網頁等待逾時：勤務系統可能登入失敗、網頁變慢，或頁面結構已變更。",
            ),
            (
                NoSuchElementException("no such element: Unable to locate element\nStacktrace:\nChromeDriver"),
                "no_such_element",
                "找不到網頁元素：可能勤務系統頁面改版，或尚未成功登入。",
            ),
            (
                RuntimeError("仍停留在 login119，找不到 _txtUsername 以外的登入後元素"),
                "login_failed",
                "登入失敗：帳號或密碼可能已變更，請登出後重新登入系統。",
            ),
            (
                RuntimeError("帳號密碼有誤或尚未申請帳號權限,請確認後再重新登入"),
                "login_failed",
                "登入失敗：帳號或密碼可能已變更，請登出後重新登入系統。",
            ),
            (
                RuntimeError("勤務表檢查未通過，已停止登打。"),
                "duty_sheet_preflight_failed",
                "勤務表檢查未通過，已停止登打。",
            ),
            (
                RuntimeError("unexpected failure\ntraceback\nsession token abc\ncookie xyz\npassword secret\nChromeDriver"),
                "unknown_error",
                "執行失敗：系統發生未預期錯誤，請查看後端日誌。",
            ),
        )

        for exc, error_code, message in cases:
            with self.subTest(error_code=error_code):
                payload = module.frontend_error_payload(exc)
                self.assertEqual(payload["error_code"], error_code)
                self.assertEqual(payload["message"], message)
                self.assertRegex(payload["timestamp"], r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}")
                rendered = json.dumps(payload, ensure_ascii=False).lower()
                for forbidden in ("stacktrace", "traceback", "chromedriver", "cookie", "session token", "password secret"):
                    self.assertNotIn(forbidden, rendered)

    def test_automation_failure_result_uses_safe_error_fields(self) -> None:
        module = duty_gui_module()
        error = RuntimeError("raw traceback\nChromeDriver\ncookie=abc\nsession token=def\npassword secret")
        result = module.automation_failure_result(
            error,
            action_index=3,
            action={"kind": "work_log", "password": "secret", "headers": {"Cookie": "abc"}},
            save=True,
            visible=False,
        )

        self.assertEqual(result["stage"], "failed")
        self.assertEqual(result["error_code"], "unknown_error")
        self.assertEqual(result["message"], "執行失敗：系統發生未預期錯誤，請查看後端日誌。")
        self.assertIn("timestamp", result)
        rendered = json.dumps(result, ensure_ascii=False).lower()
        for forbidden in ("stacktrace", "traceback", "chromedriver", "cookie", "session token", "password secret"):
            self.assertNotIn(forbidden, rendered)

    def test_tool_dialog_error_formatters_hide_selenium_details(self) -> None:
        expected = {
            "TimeoutException: timed out receiving message from renderer\nStacktrace:\nChromeDriver": "網頁等待逾時：勤務系統可能登入失敗、網頁變慢，或頁面結構已變更。",
            "NoSuchElementException: no such element: Unable to locate element\nStacktrace:\nChromeDriver": "找不到網頁元素：可能勤務系統頁面改版，或尚未成功登入。",
            "仍停留在 login119，找不到 _txtUsername 以外的登入後元素": "登入失敗：帳號或密碼可能已變更，請登出後重新登入系統。",
            "帳號密碼有誤或尚未申請帳號權限,請確認後再重新登入": "登入失敗：帳號或密碼可能已變更，請登出後重新登入系統。",
            "Traceback\nChromeDriver\ncookie=abc\nsession token=def\npassword secret": "執行失敗：系統發生未預期錯誤，請查看後端日誌。",
        }

        for module_name, module in (
            ("duty_sheet_automation", legacy_tk_module("duty_sheet_automation")),
            ("rest_time_automation", package_module("rest_time_automation")),
            ("daily_vehicle_automation", legacy_tk_module("daily_vehicle_automation")),
        ):
            for raw, safe in expected.items():
                with self.subTest(module=module_name, safe=safe):
                    self.assertEqual(module.format_automation_error(RuntimeError(raw)), safe)

    def test_four_tool_entries_register_sinposmart_callbacks(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        for tool_name in ("duty_sheet", "rest_time", "monthly_base", "daily_vehicle"):
            with self.subTest(tool_name=tool_name):
                self.assertIn(f'self.sinposmart_tool_event_callbacks("{tool_name}"', source)
        for callback_name in ("on_start=on_start", "on_finish=on_finish", "on_error=on_error"):
            with self.subTest(callback_name=callback_name):
                self.assertGreaterEqual(source.count(callback_name), 4)

    def test_tool_finish_callbacks_show_date_or_month_in_completion_notification(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        notifications: list[tuple[str, str]] = []
        gui.send_tool_start_event = lambda *_args, **_kwargs: None
        gui.send_tool_finish_event = lambda *_args, **_kwargs: None
        gui.notify_user = lambda title, message, **_kwargs: notifications.append((title, message))
        gui.after = lambda _delay, callback: callback()

        cases = (
            ("duty_sheet", "勤務表登打", "勤務表登打完成：1150728", "已完成：1150728 勤務表登打"),
            ("daily_vehicle", "車輛保養清點", "車輛保養清點已完成。", "已完成：1150728 車輛保養清點"),
            ("rest_time", "休息時間登打", "115年7月 休息時間登打完成：完成。", "已完成：115年7月 休息時間登打"),
            ("monthly_base", "勤務基準表登打", "115年7月 勤務基準表登打完成：完成。", "已完成：115年7月 勤務基準表登打"),
        )

        with mock.patch.object(module, "duty_business_roc_date", return_value="1150728"), mock.patch.object(
            module.threading, "Timer"
        ):
            for tool_name, tool_label, result, expected in cases:
                with self.subTest(tool_name=tool_name):
                    start, finish, _fail = gui.sinposmart_tool_event_callbacks(tool_name, tool_label)
                    start()
                    finish(result)
                    finish(result)
                    self.assertEqual(notifications[-1], (module.APP_DISPLAY_NAME, expected))

        self.assertEqual(len(notifications), len(cases))
        rest_source = legacy_tk_path("rest_time_dialogs.py").read_text(encoding="utf-8-sig")
        self.assertIn('on_finish(f"{expected_roc_year}年{expected_month}月 休息時間登打完成：{result}")', rest_source)
        self.assertIn('on_finish(f"{expected_roc_year}年{expected_month}月 勤務基準表登打完成：{result}")', rest_source)

    def test_rest_and_monthly_base_dialogs_use_fixed_year_and_three_month_combo(self) -> None:
        source = legacy_tk_path("rest_time_dialogs.py").read_text(encoding="utf-8-sig")

        self.assertIn("month_var", source)
        self.assertIn("nearby_month_options", source)
        self.assertIn("CTK_COMBO_STYLE", source)
        self.assertIn("ctk.CTkComboBox", source)
        self.assertNotIn("ctk.CTkOptionMenu", source)
        self.assertNotIn("year_var", source)
        self.assertIn("selected_year_month", source)
        self.assertIn("expected_roc_year", source)
        self.assertIn("expected_month", source)

        module = rest_time_module()
        self.assertEqual(module.nearby_month_options(6), ["05", "06", "07"])
        self.assertEqual(module.nearby_month_options(12), ["11", "12", "01"])

    def test_rest_time_rejects_workbook_month_mismatch(self) -> None:
        module = rest_time_module()
        workbook = module.openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(row=2, column=4).value = 115
        sheet.cell(row=2, column=5).value = 7
        sheet.cell(row=2, column=7).value = 31

        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "rest.xlsx"
            workbook.save(path)

            with self.assertRaisesRegex(RuntimeError, "Excel.*115年07月"):
                module.validate_workbook_year_month(path, 115, 6)

    def test_rest_time_rejects_directory_and_non_excel_workbook_paths(self) -> None:
        module = rest_time_module()

        with tempfile.TemporaryDirectory() as tmp_dir:
            directory = Path(tmp_dir)
            non_excel = directory / "duty.csv"
            non_excel.write_text("not an Excel workbook", encoding="utf-8")

            for path in (directory, non_excel):
                with self.subTest(path=path):
                    with self.assertRaisesRegex(RuntimeError, "有效的勤務表 Excel 檔案"):
                        module.validate_rest_workbook_path(path)

    def test_rest_and_monthly_base_failures_use_safe_diagnostic_messages(self) -> None:
        module = rest_time_module()
        from selenium.common.exceptions import WebDriverException

        self.assertEqual(
            module.format_automation_error(WebDriverException("driver startup failed")),
            "瀏覽器啟動或連線失敗：請關閉卡住的 Chrome 後重試。",
        )
        session_error = duty_rehearsal_module().DutyBrowserSessionOpenError(attempts=2)
        self.assertEqual(
            module.format_automation_error(session_error),
            "瀏覽器在登入或開啟勤務頁面時中斷：已自動使用新的工作階段重試，仍失敗請重新登入後再試。",
        )
        self.assertEqual(
            module.format_automation_error(RuntimeError("讀取固定 Google 試算表失敗：HTTP 503")),
            "勤務基準表登打失敗：輪休基準表無法讀取，請確認網路與 Google 試算表後重試。",
        )

    def test_tool_failure_payload_classifies_rest_and_monthly_base_diagnostics(self) -> None:
        module = duty_gui_module()

        cases = (
            ("休息時間登打失敗：請選擇有效的勤務表 Excel 檔案（.xlsx 或 .xlsm）。", "rest_workbook_invalid"),
            ("瀏覽器啟動或連線失敗：請關閉卡住的 Chrome 後重試。", "browser_error"),
            ("勤務基準表登打失敗：輪休基準表無法讀取，請確認網路與 Google 試算表後重試。", "monthly_base_source_failed"),
        )
        for message, error_code in cases:
            with self.subTest(error_code=error_code):
                self.assertEqual(module.frontend_error_payload(message)["error_code"], error_code)

    def test_monthly_base_rejects_selected_month_mismatch(self) -> None:
        module = rest_time_module()

        with self.assertRaisesRegex(RuntimeError, "勤務基準表.*115年07月"):
            module.validate_selected_year_month("勤務基準表", 115, 6, 115, 7)

    def test_duty_sheet_filters_trainee_numbers_from_shift_workbook(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        roster = workbook.active
        roster.title = "輪休基準表"
        roster.cell(row=5, column=12).value = "番號"
        roster.cell(row=5, column=13).value = "姓名"
        roster.cell(row=5, column=14).value = "班表欄位"
        roster.cell(row=6, column=12).value = 28
        roster.cell(row=6, column=13).value = "正式人員"
        roster.cell(row=6, column=14).value = "B班"
        roster.cell(row=7, column=12).value = 29
        roster.cell(row=7, column=13).value = "實習人員"
        roster.cell(row=7, column=14).value = "實習生"

        excluded = module.trainee_numbers_from_workbook(workbook)

        self.assertEqual(excluded, {"29"})
        self.assertEqual(module.clean_to_list_excluding("18,29,23", excluded), ["18", "23"])
        self.assertEqual(module.clean_v_excluding("29,2", excluded), "2")
        self.assertEqual(
            module.clean_to_list("1,2 3，4.5．6、7·8。9"),
            ["1", "2", "3", "4", "5", "6", "7", "8", "9"],
        )

    def test_duty_sheet_captures_images_while_website_submission_runs(self) -> None:
        module = legacy_duty_sheet_module()
        daily_capture = {"image_path": "daily.png", "capture_range": "B3:AM36"}
        night_capture = {"image_path": "night.png", "capture_range": "B24:AM33"}

        with mock.patch.object(module, "preview_excel_capture", return_value=daily_capture), mock.patch.object(
            module, "preview_night_excel_capture", return_value=night_capture
        ):
            captures = module.capture_duty_sheet_images("duty.xlsm", "1150809")

        self.assertEqual(captures, (daily_capture, night_capture))
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(encoding="utf-8-sig")
        self.assertIn("capture_executor.submit(capture_duty_sheet_images, excel_path, target_date)", source)
        self.assertLess(
            source.index("capture_executor.submit(capture_duty_sheet_images, excel_path, target_date)"),
            source.index("driver = retry_duty_number_popup_preflight("),
        )

    def test_external_duty_followed_by_rest_returns_before_rest_starts(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150808",
            rows=[
                module.DutyRow("14~15", {"值班": ["19"], "備勤": ["5", "25"]}),
                module.DutyRow("15~16", {"值班": ["19"], "備勤": ["5", "25"]}),
                module.DutyRow("16~17", {"值班": ["27"], "休息": ["15"], "防溺車巡暨車輛駕訓": ["5", "25"]}),
                module.DutyRow("17~18", {"值班": ["27"], "休息": ["15"], "防溺車巡暨車輛駕訓": ["5", "25"]}),
                module.DutyRow("18~19", {"值班": ["5"], "休息": ["25"]}),
                module.DutyRow("19~20", {"值班": ["5"], "休息": ["25"]}),
            ],
            summary={"在勤": ["5", "15", "19", "25", "27"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150808"), [], None)
        person_actions = [
            (action.time, action.source, action.actor, action.fields.get("出或入"), action.fields.get("領用事由及地點"))
            for action in actions
            if action.target == "25" and action.source in ("外勤簽出", "外勤簽入", "休息簽出", "休息結束")
        ]

        self.assertEqual(
            person_actions,
            [
                ("16:00", "外勤簽出", "19", "出", "防溺車巡"),
                ("18:00", "外勤簽入", "27", "入", "防溺車巡返隊"),
                ("18:00", "休息簽出", "27", "出", "休息"),
                ("20:00", "休息結束", "5", "入", "休息返隊"),
            ],
        )

    def test_external_duty_after_rest_uses_schedule_time_for_partial_public_leave(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150812",
            rows=[
                module.DutyRow("09~17", {"值班": ["17"], "救助複訓": ["21"]}),
                module.DutyRow("17~18", {"值班": ["23"], "休息": ["21"]}),
                module.DutyRow("18~20", {"值班": ["5"], "防溺車巡": ["21"]}),
            ],
            summary={"公假": ["21"], "在勤": ["5", "17", "23"]},
        )

        self.assertEqual(
            [block for block in module.external_duty_blocks(today) if block[1] == "21"],
            [
                ("救助複訓", "21", 9, 17, 0),
                ("防溺車巡", "21", 18, 20, 0),
            ],
        )
        self.assertEqual(
            [block for block in module.rest_blocks(today) if block[0] == "21"],
            [("21", 17, 18)],
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150812"), [], None)
        person_actions = [
            (
                action.time,
                action.source,
                action.actor,
                action.target,
                action.fields.get("出或入"),
                action.fields.get("領用事由及地點"),
                action.date_offset,
            )
            for action in actions
            if action.target == "21"
            and action.source in ("外勤簽出", "外勤簽入", "休息簽出", "休息結束")
        ]

        self.assertEqual(
            person_actions,
            [
                ("09:00", "外勤簽出", "17", "21", "出", "救助複訓", 0),
                ("17:00", "外勤簽入", "17", "21", "入", "返隊", 0),
                ("17:00", "休息簽出", "17", "21", "出", "休息", 0),
                ("18:00", "休息結束", "23", "21", "入", "休息返隊", 0),
                ("18:00", "外勤簽出", "23", "21", "出", "防溺車巡", 0),
                ("20:00", "外勤簽入", "5", "21", "入", "防溺車巡返隊", 0),
            ],
        )

    def test_drowning_patrol_uses_car_patrol_fields_and_one_work_log(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150826",
            rows=[
                module.DutyRow("14~16", {"值班": ["19"], "備勤": ["5", "25"]}),
                module.DutyRow("16~17", {"值班": ["19"], "防溺車巡暨車輛駕訓": ["5", "25"]}),
                module.DutyRow("17~18", {"值班": ["19"], "防溺車巡暨車輛駕訓": ["5", "25"]}),
                module.DutyRow("18~20", {"值班": ["27"], "備勤": ["5", "25"]}),
            ],
            summary={"在勤": ["5", "19", "25", "27"]},
            staff={
                "5": {"role": "隊員", "name": "甲"},
                "19": {"role": "隊員", "name": "值班員"},
                "25": {"role": "小隊長", "name": "乙"},
                "27": {"role": "隊員", "name": "接班員"},
            },
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150826"), [], None)
        external_entries = [action for action in actions if action.source in ("外勤簽出", "外勤簽入")]
        work_actions = [action for action in actions if action.kind == "work_log" and action.source == "防溺車巡"]

        self.assertEqual(len(external_entries), 4)
        for action in external_entries:
            self.assertEqual(action.fields["勤務項目"], "車巡")
            self.assertEqual(action.fields["事由"], "防溺")
        self.assertEqual(
            {(action.target, action.fields["出或入"], action.fields["領用事由及地點"]) for action in external_entries},
            {
                ("5", "出", "防溺車巡"),
                ("25", "出", "防溺車巡"),
                ("5", "入", "防溺車巡返隊"),
                ("25", "入", "防溺車巡返隊"),
            },
        )
        self.assertEqual(len(work_actions), 1)
        work = work_actions[0]
        self.assertEqual(work.time, "18:00")
        self.assertEqual(work.actor, "19")
        self.assertEqual(work.target, "5,25")
        self.assertEqual(
            work.fields,
            {
                "工作時間": "18:00",
                "勤務項目": "車巡",
                "事由": "防溺",
                "工作概述": (
                    "一、時間：1600-1800\n"
                    "二、地點：\n"
                    "至轄內16處危險水域沿路巡視，巡視學生埤、死囝仔埤、坡內埤埤塘、8-14號埤塘、銅鑼埤塘、崙坪埤塘、鉤漕埤塘、桃園大圳、富源里富源62-25號對面埤塘、泉州埤塘、公埤塘、大埤塘、草埤塘、豬屠埤塘、桃園大圳第八支線、龜墓埤塘。\n"
                    "三、人車:隊員 甲，小隊長 乙，新坡15車。\n"
                    "四、處理情形：勸導民眾勿戲水人數0人。"
                ),
                "處理情形": "未發現有民眾戲水之情形，發放宣導單0份。 16處危險水域易發生溺水案件處所車巡期間無事故。",
                "服勤人員": ["5", "25"],
            },
        )

    def test_drowning_patrol_work_duplicate_requires_drowning_reason(self) -> None:
        find_work_matches = package_module("compare_rehearsal_records").find_work_matches

        action = {
            "kind": "work_log",
            "time": "18:00",
            "source": "防溺車巡",
            "fields": {
                "工作時間": "18:00",
                "勤務項目": "車巡",
                "事由": "防溺",
                "工作概述": "一、時間：1600-1800",
            },
        }
        rows = [
            "115/08/26 18:00 | 車巡 | 防颱",
            "115/08/26 18:00 | 車巡 | 防溺 | 一、時間：1600-1800",
        ]

        self.assertEqual(find_work_matches(rows, "1150826", {}, action), [rows[1]])

    def test_partial_public_leave_rest_at_0800_uses_previous_day_worker_state(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150812",
            rows=[
                module.DutyRow("08~09", {"值班": ["17"], "休息": ["21"]}),
                module.DutyRow("09~17", {"值班": ["17"], "救助複訓": ["21"]}),
            ],
            summary={"公假": ["21"], "在勤": ["17"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150811",
            rows=[module.DutyRow("06~08", {"值班": ["10"], "備勤": ["21"]})],
            summary={"在勤": ["10", "21"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150812"), [], None)
        person_actions = [
            (
                action.time,
                action.source,
                action.actor,
                action.fields.get("出或入"),
                action.fields.get("領用事由及地點"),
                action.date_offset,
            )
            for action in actions
            if action.target == "21"
        ]

        self.assertEqual(
            person_actions,
            [
                ("08:00", "休息簽出", "10", "出", "休息", 0),
                ("09:00", "休息結束", "17", "入", "休息返隊", 0),
                ("09:00", "外勤簽出", "17", "出", "救助複訓", 0),
                ("17:00", "外勤簽入", "17", "入", "返隊", 0),
            ],
        )

    def test_partial_public_leave_rest_at_0800_delays_arrival_after_previous_day_off(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150812",
            rows=[
                module.DutyRow("08~09", {"值班": ["17"], "休息": ["21"]}),
                module.DutyRow("09~17", {"值班": ["17"], "救助複訓": ["21"]}),
            ],
            summary={"公假": ["21"], "在勤": ["17"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150811",
            rows=[module.DutyRow("06~08", {"值班": ["10"], "備勤": []})],
            summary={"輪休": ["21"], "在勤": ["10"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150812"), [], None)
        person_actions = [
            (
                action.time,
                action.source,
                action.actor,
                action.fields.get("出或入"),
                action.fields.get("領用事由及地點"),
                action.date_offset,
            )
            for action in actions
            if action.target == "21"
        ]

        self.assertEqual(
            person_actions,
            [
                ("09:00", "今日在勤且昨日未在勤", "17", "入", "到勤", 0),
                ("09:00", "外勤簽出", "17", "出", "救助複訓", 0),
                ("17:00", "外勤簽入", "17", "入", "返隊", 0),
            ],
        )

    def test_next_morning_rest_uses_actual_rows_before_public_leave_summary(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150812",
            rows=[
                module.DutyRow("06~08", {"值班": ["5"], "休息": ["21"]}),
                module.DutyRow("08~10", {"值班": ["17"], "休息": []}),
            ],
            summary={"在勤": ["5", "17", "21"]},
        )
        tomorrow = module.DutySheet(
            roc_date="1150813",
            rows=[module.DutyRow("08~10", {"值班": ["23"], "防溺車巡": ["21"]})],
            summary={"公假": ["21"], "在勤": ["23"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150812"), [], tomorrow)
        rest_actions = [
            (
                action.time,
                action.fields.get("出或入"),
                action.fields.get("領用事由及地點"),
                action.source,
                action.date_offset,
            )
            for action in actions
            if action.target == "21"
            and action.source in ("休息簽出", "休息結束", "休息後退勤")
        ]

        self.assertEqual(
            rest_actions,
            [
                ("06:00", "出", "休息", "休息簽出", 1),
                ("08:00", "入", "休息返隊", "休息結束", 1),
            ],
        )

    def test_next_morning_external_return_uses_actual_rows_before_public_leave_summary(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150812",
            rows=[
                module.DutyRow("06~08", {"值班": ["5"], "防溺車巡": ["21"]}),
                module.DutyRow("08~10", {"值班": ["17"], "防溺車巡": []}),
            ],
            summary={"在勤": ["5", "17", "21"]},
        )
        tomorrow = module.DutySheet(
            roc_date="1150813",
            rows=[module.DutyRow("08~10", {"值班": ["23"], "備勤": ["21"]})],
            summary={"公假": ["21"], "在勤": ["23"]},
        )

        self.assertEqual(
            [block for block in module.external_duty_blocks(today, tomorrow) if block[1] == "21"],
            [("防溺車巡", "21", 6, 8, 0)],
        )

    def test_duty_sheet_truncates_external_duty_names_to_24_display_units(self) -> None:
        module = legacy_duty_sheet_module()

        self.assertEqual(module.truncate_external_duty_name("測" * 13), "測" * 12)
        self.assertEqual(
            module.truncate_external_duty_name("救護" + "A" * 21),
            "救護" + "A" * 20,
        )
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        self.assertIn("task_name = truncate_external_duty_name(raw_name)", source)

    def test_fire_mission_does_not_reuse_people_across_disaster_vehicles(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission([], ["8", "8", "9", "10"], [], "")

        self.assertIsNotNone(mission)
        assigned = {
            vehicle: set(members.split(",")) - {""}
            for vehicle, members in mission.items()
        }
        self.assertFalse(assigned["attack"] & assigned["relay"])
        self.assertEqual(assigned["attack"] | assigned["relay"], {"8", "9", "10"})

    def test_fire_mission_fills_the_unique_pool_after_duplicate_removal(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission(["12"], ["8", "8", "9", "10", "11"], [], "")

        self.assertIsNotNone(mission)
        assigned = {
            vehicle: set(members.split(",")) - {""}
            for vehicle, members in mission.items()
        }
        self.assertFalse(assigned["attack"] & assigned["relay"])
        self.assertEqual(assigned["attack"] | assigned["relay"], {"8", "9", "10", "11", "12"})

    def test_fire_mission_does_not_reuse_officer_drivers_across_disaster_vehicles(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission([], ["1", "2"], [], "")

        self.assertIsNotNone(mission)
        assigned = {
            vehicle: set(members.split(",")) - {""}
            for vehicle, members in mission.items()
        }
        self.assertFalse(assigned["attack"] & assigned["relay"])
        self.assertEqual(assigned["attack"] | assigned["relay"], {"1", "2"})

    def test_ambulance2_uses_standby_then_on_duty_external_staff_without_ambulance1_overlap(self) -> None:
        module = legacy_duty_sheet_module()

        self.assertEqual(module.select_ambulance2_members(["8", "9"], ["10"], ["6", "7"]), ["8", "9"])
        self.assertEqual(module.select_ambulance2_members(["8"], ["6", "9"], ["6", "7"]), ["8", "9"])
        self.assertEqual(module.select_ambulance2_members([], ["6", "7", "8", "9"], ["6", "7"]), ["8", "9"])
        self.assertEqual(module.select_ambulance2_members(["6", "7"], ["8", "9"], ["6", "7"]), ["8", "9"])

    def test_fire_mission_uses_one_ambulance1_member_when_standby_has_four_people(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission(["6", "7"], ["8", "9", "10", "11"], [], "10")

        self.assertEqual(mission, {"attack": "9,10", "relay": "8,11,6"})

    def test_fire_mission_uses_both_ambulance1_members_when_standby_has_three_people(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission(["6", "7"], ["8", "9", "10"], [], "10")

        self.assertEqual(mission, {"attack": "9,10", "relay": "8,6,7"})

    def test_fire_mission_prefers_an_officer_for_attack_second_when_commander_is_blank(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission([], ["8", "9", "4", "10", "5"], [], "")

        self.assertEqual(mission, {"attack": "9,4", "relay": "8,10,5"})

    def test_fire_mission_uses_ambulance1_and_one_on_duty_external_member_when_standby_has_two_people(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission(["6", "7"], ["8", "9"], ["10"], "10")

        self.assertEqual(mission, {"attack": "9,10", "relay": "8,6,7"})

    def test_fire_mission_uses_the_specified_external_staffing_when_standby_has_one_or_zero_people(self) -> None:
        module = legacy_duty_sheet_module()

        self.assertEqual(
            module.calculate_fire_mission(["6", "7"], ["8"], ["9", "10", "11", "12"], "10"),
            {"attack": "9,10", "relay": "8,11,12"},
        )
        self.assertEqual(
            module.calculate_fire_mission(["6", "7"], [], ["8", "9", "10", "11", "12"], "9"),
            {"attack": "8,9", "relay": "10,11,12"},
        )

    def test_fire_mission_caps_standby_candidates_at_ten_and_replaces_a_conflicting_driver_for_the_commander(self) -> None:
        module = legacy_duty_sheet_module()

        mission = module.calculate_fire_mission(
            [], ["1", "2", "8", "9", "10", "11", "12", "13", "14", "15", "16"], [], "1"
        )

        self.assertEqual(mission, {"attack": "2,1,13,14,15", "relay": "8,9,10,11,12"})

    def test_daily_sheet_preflight_reports_duplicate_and_missing_assignments(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        roster = workbook.active
        roster.title = "輪休基準表"
        params = workbook.create_sheet("班別參數")
        sheet = workbook.create_sheet("7號")

        params.cell(row=4, column=1).value = None
        params.cell(row=4, column=2).value = "上班"
        roster.cell(row=4, column=5).value = 1
        roster.cell(row=4, column=6).value = 2
        roster.cell(row=4, column=7).value = 29
        roster.cell(row=5, column=3).value = "日期"
        roster.cell(row=5, column=5).value = "甲"
        roster.cell(row=5, column=6).value = "乙"
        roster.cell(row=5, column=7).value = "實習"
        roster.cell(row=6, column=3).value = 7

        sheet.cell(row=5, column=3).value = "值班"
        sheet.cell(row=6, column=8).value = "指揮官"
        for row in range(10, 34):
            sheet.cell(row=row, column=2).value = f"{row - 2:02d}-{row - 1:02d}"
            sheet.cell(row=row, column=3).value = "1"
            sheet.cell(row=row, column=8).value = "2"
        sheet.cell(row=10, column=8).value = "1"
        sheet.cell(row=10, column=5).value = "29"

        issues = module.validate_daily_sheet_assignments(workbook, sheet, 7, {"29"})

        self.assertTrue(any("08-09" in issue and "重複" in issue and "1" in issue for issue in issues))
        self.assertTrue(any("08-09" in issue and "漏排" in issue and "2" in issue for issue in issues))
        self.assertFalse(any("29" in issue for issue in issues))

    def test_daily_sheet_preflight_accepts_complete_assignments(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        roster = workbook.active
        roster.title = "輪休基準表"
        params = workbook.create_sheet("班別參數")
        sheet = workbook.create_sheet("7號")

        params.cell(row=4, column=1).value = None
        params.cell(row=4, column=2).value = "上班"
        roster.cell(row=4, column=5).value = 1
        roster.cell(row=4, column=6).value = 2
        roster.cell(row=5, column=3).value = "日期"
        roster.cell(row=5, column=5).value = "甲"
        roster.cell(row=5, column=6).value = "乙"
        roster.cell(row=6, column=3).value = 7

        sheet.cell(row=5, column=3).value = "值班"
        sheet.cell(row=6, column=8).value = "指揮官"
        for row in range(10, 34):
            sheet.cell(row=row, column=2).value = f"{row - 2:02d}-{row - 1:02d}"
            sheet.cell(row=row, column=3).value = "1"
            sheet.cell(row=row, column=8).value = "2"

        self.assertEqual(module.validate_daily_sheet_assignments(workbook, sheet, 7, set()), [])

    def test_daily_sheet_preflight_requires_commander_in_each_merged_time_slot(self) -> None:
        module = legacy_duty_sheet_module()
        for commander in (None, " \t\n\u3000", 0, "0"):
            with self.subTest(commander=commander):
                workbook = module.openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "24號"
                sheet.cell(5, 3, "值班")
                sheet.cell(6, 8, "指揮官")
                for row in range(10, 34):
                    sheet.cell(row, 2, f"{row - 2:02d}-{row - 1:02d}")
                    sheet.cell(row, 3, "16")
                    sheet.cell(row, 4, "28,1")
                    sheet.cell(row, 8, "2")
                sheet.cell(12, 8).value = commander
                sheet.merge_cells("H12:H13")
                with mock.patch.object(module, "expected_on_duty_numbers", return_value=["16", "28", "1", "2"]):
                    issues = module.validate_daily_sheet_assignments(workbook, sheet, 24, set())
                self.assertEqual(len(issues), 2)
                self.assertTrue(any("10-11" in issue and "【指揮官未排】" in issue for issue in issues))
                self.assertTrue(any("11-12" in issue and "【指揮官未排】" in issue for issue in issues))

    def test_daily_sheet_preflight_counts_merged_assignments_for_each_time_slot(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        roster = workbook.active
        roster.title = "輪休基準表"
        params = workbook.create_sheet("班別參數")
        sheet = workbook.create_sheet("7號")

        params.cell(row=4, column=1).value = None
        params.cell(row=4, column=2).value = "上班"
        roster.cell(row=4, column=5).value = 1
        roster.cell(row=4, column=6).value = 2
        roster.cell(row=5, column=3).value = "日期"
        roster.cell(row=5, column=5).value = "甲"
        roster.cell(row=5, column=6).value = "乙"
        roster.cell(row=6, column=3).value = 7

        sheet.cell(row=5, column=3).value = "值班"
        sheet.cell(row=6, column=8).value = "指揮官"
        for row in range(10, 34):
            sheet.cell(row=row, column=2).value = f"{row - 2:02d}-{row - 1:02d}"
            sheet.cell(row=row, column=3).value = "1"
            sheet.cell(row=row, column=8).value = "2"
        sheet.merge_cells(start_row=10, start_column=3, end_row=11, end_column=3)
        sheet.merge_cells(start_row=10, start_column=8, end_row=11, end_column=8)

        self.assertEqual(module.validate_daily_sheet_assignments(workbook, sheet, 7, set()), [])

    def test_daily_standby_members_require_blank_duty_number_leave_type(self) -> None:
        module = legacy_duty_sheet_module()
        website_rows = [
            {"staff_no": "10", "name": "甲", "leave_type": ""},
            {"staff_no": "11", "name": "乙", "leave_type": "○"},
            {"staff_no": "12", "name": "丙", "leave_type": ""},
        ]

        self.assertEqual(
            module.validate_daily_standby_against_duty_number_leave_types(
                ["10"], website_rows, set()
            ),
            [],
        )
        issues = module.validate_daily_standby_against_duty_number_leave_types(
            ["10", "11", "12"], website_rows, set()
        )

        self.assertTrue(any("11" in issue and "休假別" in issue for issue in issues))
        self.assertFalse(any("12" in issue for issue in issues))

    def test_duty_number_leave_type_check_reports_missing_staff(self) -> None:
        module = legacy_duty_sheet_module()

        issues = module.validate_daily_standby_against_duty_number_leave_types(
            ["10", "13"],
            [{"staff_no": "10", "name": "甲", "leave_type": ""}],
            set(),
        )

        self.assertTrue(any("13" in issue and "找不到" in issue for issue in issues))

    def test_duty_number_reader_anchors_each_row_from_designation_input(self) -> None:
        module = legacy_duty_sheet_module()
        captured = {}

        class FakeDriver:
            def execute_script(self, script: str):
                captured["script"] = script
                return [{"staff_no": "1", "name": "甲", "leave_type": ""}]

        self.assertEqual(
            module.read_duty_number_leave_rows(FakeDriver()),
            [{"staff_no": "1", "name": "甲", "leave_type": ""}],
        )
        self.assertIn("document.querySelectorAll('input[name^=\"_DESIGNATION\"]')", captured["script"])
        self.assertIn("numberInput.closest('tr')", captured["script"])
        self.assertNotIn("document.querySelectorAll('tr')", captured["script"])

    def test_daily_standby_members_are_read_from_daily_sheet_standby_summary(self) -> None:
        module = legacy_duty_sheet_module()
        sheet = module.openpyxl.Workbook().active
        sheet.cell(row=22, column=3).value = "備勤"
        sheet.cell(row=22, column=4).value = "10, 11"

        self.assertEqual(module.expected_on_duty_numbers_from_daily_sheet(sheet, set()), ["10", "11"])

    def test_mission_fill_returns_browser_result_and_readback_ignores_non_blocking_alerts(self) -> None:
        module = legacy_duty_sheet_module()

        class FakeDriver:
            def execute_script(self, *_args):
                return {"status": "saved", "found_ids": ["_pln_8_1"], "save_found": True}

        result = module.step_fill_mission_cells(FakeDriver(), {"_pln_8_1": "1,2"})

        self.assertEqual(result["status"], "saved")
        self.assertEqual(
            module.validate_mission_cell_values(
                {"_pln_8_1": "1,2"},
                {"_pln_8_1": "1,2"},
                ["提示：部分資料已存在"],
            ),
            [],
        )
        issues = module.validate_mission_cell_values(
            {"_pln_8_1": "1,2"},
            {"_pln_8_1": "1"},
            [],
        )
        self.assertTrue(any("_pln_8_1" in issue for issue in issues))

    def test_base_month_text_can_detect_site_month_mismatch(self) -> None:
        module = rest_time_module()

        self.assertEqual(module.extract_base_month_from_text("目前編輯月份為: 115年07月"), (115, 7))
        with self.assertRaisesRegex(RuntimeError, "網站.*115年07月"):
            module.validate_selected_year_month("網站", 115, 6, 115, 7)

    def test_monthly_base_fill_clears_existing_values_before_rewrite(self) -> None:
        source = (package_dir() / "rest_time_automation.py").read_text(encoding="utf-8-sig")

        self.assertIn("function clearRowValues(row, days)", source)
        self.assertIn("function dayPlanControls(row, days)", source)
        self.assertIn("Number(match[2]) >= 1", source)
        self.assertIn("el.value = ''", source)
        self.assertLess(source.index("clearRowValues(row, Object.keys(data).length)"), source.index("return setRowValues(row, data)"))

    def test_update_logout_command_reports_logout_synchronously(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn('elif message == "update_logout":', source)
        self.assertIn("def report_update_logout(self) -> bool:", source)
        self.assertIn("def update_logout_identity(self)", source)
        self.assertIn("last_update_logout_identity", source)
        self.assertIn('"logout"', source)
        self.assertIn('trigger_type="update"', source)
        self.assertIn("immediate=True", source)

    def test_duty_control_buttons_keep_spacing(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn('self.manual_pause_button.pack(side=tk.RIGHT, padx=(0, 6))', source)
        self.assertIn('self.resume_schedule_button.pack(side=tk.RIGHT, padx=(0, 6))', source)
        self.assertTrue(
            'self.early_submit_button.pack(side=tk.RIGHT, padx=(0, 6))' in source,
            "手動登打按鈕右側需保留 6px 間距，避免貼住手動暫停按鈕。",
        )
        self.assertEqual(
            source.count("self.early_submit_button.pack(side=tk.RIGHT)\n"),
            0,
            "手動登打按鈕不得用無 padx 的 pack 呼叫，登入後重新顯示也要保留間距。",
        )

    def test_logged_out_simple_mode_uses_compact_height(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertTrue(
            'self.geometry("550x320")' in source,
            "未登入值班模式視窗高度應貼近登入卡片內容，避免下方大片留白。",
        )
        self.assertTrue(
            "self.minsize(530, 300)" in source,
            "未登入值班模式最小高度應維持精簡。",
        )

    def test_auto_logout_waits_until_submit_queues_are_idle(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class StatusText:
            value = ""

            def set(self, value: str) -> None:
                self.value = value

        scheduled: list[tuple[int, object]] = []
        gui.auto_logout_after_id = None
        gui.auto_logout_deadline = None
        gui.auto_logout_handoff_at = None
        gui.auto_logout_actor_no = ""
        gui.pending_auto_logout_actor_no = ""
        gui.pending_auto_logout_deadline = None
        gui.pending_auto_logout_handoff_at = None
        gui.submit_queues = {"entry": [("queued",)], "work": []}
        gui.submit_worker_running = {"entry": True, "work": False}
        gui.duty_status_text = StatusText()
        gui.after = lambda delay_ms, callback: scheduled.append((delay_ms, callback)) or "after-id"
        gui.after_cancel = lambda _after_id: None
        gui.action_datetime = lambda _action: datetime(2026, 6, 26, 12, 0)
        action = {"kind": "entry_log", "time": "12:00", "fields": {"出或入": "值退"}}

        gui.schedule_auto_logout("10", action)

        self.assertIsNone(gui.auto_logout_after_id)
        self.assertEqual(gui.pending_auto_logout_actor_no, "10")
        self.assertEqual(scheduled, [])

        gui.submit_queues = {"entry": [], "work": []}
        gui.submit_worker_running = {"entry": False, "work": False}
        gui.schedule_pending_auto_logout_if_idle()

        self.assertEqual(gui.auto_logout_after_id, "after-id")
        self.assertEqual(gui.auto_logout_actor_no, "10")
        self.assertEqual(gui.pending_auto_logout_actor_no, "")
        self.assertEqual(len(scheduled), 1)

    def test_auto_logout_waits_until_manual_pause_is_resumed(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class StatusText:
            value = ""

            def set(self, value: str) -> None:
                self.value = value

        scheduled: list[tuple[int, object]] = []
        gui.auto_logout_after_id = None
        gui.auto_logout_deadline = None
        gui.auto_logout_handoff_at = None
        gui.auto_logout_actor_no = ""
        gui.pending_auto_logout_actor_no = ""
        gui.pending_auto_logout_deadline = None
        gui.pending_auto_logout_handoff_at = None
        gui.submit_queues = {"entry": [], "work": []}
        gui.submit_worker_running = {"entry": False, "work": False}
        gui.manual_paused_due_indices = {0: "10"}
        gui.duty_status_text = StatusText()
        gui.after = lambda delay_ms, callback: scheduled.append((delay_ms, callback)) or "after-id"
        gui.after_cancel = lambda _after_id: None
        gui.action_datetime = lambda _action: datetime(2026, 7, 2, 12, 0)
        action = {"kind": "entry_log", "time": "12:00", "fields": {"出或入": "值退"}}

        gui.schedule_auto_logout("10", action)

        self.assertIsNone(gui.auto_logout_after_id)
        self.assertEqual(gui.pending_auto_logout_actor_no, "10")
        self.assertEqual(scheduled, [])
        self.assertIn("人員手動暫停", gui.duty_status_text.value)

        gui.schedule_pending_auto_logout_if_idle()

        self.assertEqual(scheduled, [])

        gui.manual_paused_due_indices.clear()
        gui.schedule_pending_auto_logout_if_idle()

        self.assertEqual(gui.auto_logout_after_id, "after-id")
        self.assertEqual(gui.auto_logout_actor_no, "10")
        self.assertEqual(len(scheduled), 1)

    def test_manual_handoff_checkout_schedules_auto_logout_from_planned_time(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class StatusText:
            value = ""

            def set(self, value: str) -> None:
                self.value = value

        scheduled: list[tuple[int, object]] = []
        gui.auto_logout_after_id = None
        gui.auto_logout_deadline = None
        gui.auto_logout_handoff_at = None
        gui.auto_logout_actor_no = ""
        gui.pending_auto_logout_deadline = None
        gui.pending_auto_logout_handoff_at = None
        gui.pending_auto_logout_actor_no = ""
        gui.submit_queues = {"entry": [], "work": []}
        gui.submit_worker_running = {"entry": False, "work": False}
        gui.manual_paused_due_indices = {}
        gui.duty_status_text = StatusText()
        gui.after = lambda delay, callback: scheduled.append((delay, callback)) or "after-id"
        gui.after_cancel = lambda _after_id: None
        gui.action_datetime = lambda _action: datetime(2026, 7, 10, 18, 0)
        action = {
            "kind": "entry_log",
            "source": "值班交接",
            "fields": {"出或入": "值退"},
        }

        self.assertTrue(gui.should_schedule_auto_logout(action, "manual"))
        gui.schedule_auto_logout("28", action)

        self.assertEqual(gui.auto_logout_handoff_at, datetime(2026, 7, 10, 18, 0))
        self.assertEqual(gui.auto_logout_deadline, datetime(2026, 7, 10, 18, 10))
        self.assertEqual(gui.auto_logout_actor_no, "28")
        self.assertEqual(len(scheduled), 1)

    def test_auto_logout_rechecks_until_handoff_group_is_complete(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class StatusText:
            value = ""

            def set(self, value: str) -> None:
                self.value = value

        real_datetime = module.datetime

        class FixedDatetime(real_datetime):
            current = real_datetime(2026, 7, 10, 18, 10)

            @classmethod
            def now(cls, tz=None):
                return cls.current

        handoff_at = datetime(2026, 7, 10, 18, 0)
        first_check = datetime(2026, 7, 10, 18, 10)
        scheduled: list[tuple[int, object]] = []
        cleared: list[str] = []
        gui.session = module.LoginSession(actor_no="28", user_id="user28", password="secret", verified=True)
        gui.auto_logout_after_id = "after-id"
        gui.auto_logout_deadline = first_check
        gui.auto_logout_handoff_at = handoff_at
        gui.auto_logout_actor_no = "28"
        gui.pending_auto_logout_deadline = None
        gui.pending_auto_logout_handoff_at = None
        gui.pending_auto_logout_actor_no = ""
        gui.submit_queues = {"entry": [], "work": []}
        gui.submit_worker_running = {"entry": False, "work": False}
        gui.manual_paused_due_indices = {}
        gui.duty_status_text = StatusText()
        gui.login_status = StatusText()
        gui.after = lambda delay, callback: scheduled.append((delay, callback)) or f"after-{len(scheduled)}"
        gui.after_cancel = lambda _after_id: None
        gui.sync_duty_compare_from_audit = lambda: None
        gui.clear_login = lambda trigger_type="manual": cleared.append(trigger_type)
        gui.logged_in_identity_label = lambda _actor: "28番"
        gui.set_duty_status = lambda message, **_kwargs: gui.duty_status_text.set(message)
        gui.notify_user = lambda *_args, **_kwargs: None
        gui.action_completion_key = lambda action: action["key"]
        gui.action_datetime = lambda action: action["at"]
        gui.duty_actions = [
            {"key": "checkout", "at": handoff_at, "kind": "entry_log", "actor": "28", "source": "值班交接", "fields": {"出或入": "值退"}},
            {"key": "checkin", "at": handoff_at, "kind": "entry_log", "actor": "28", "source": "值班交接", "fields": {"出或入": "值班"}},
            {"key": "work", "at": handoff_at, "kind": "work_log", "actor": "28", "source": "值班交接", "fields": {}},
            {"key": "rest-return", "at": handoff_at, "kind": "entry_log", "actor": "28", "source": "休息結束", "fields": {"出或入": "入", "領用事由及地點": "休息返隊"}},
            {"key": "next-shift", "at": datetime(2026, 7, 10, 20, 0), "kind": "work_log", "actor": "12", "source": "值班交接", "fields": {}},
        ]
        gui.duty_action_compare = {
            0: {"group": "done"},
            1: {"group": "done"},
            2: {"group": "todo"},
            3: {"group": "todo"},
            4: {"group": "todo"},
        }
        gui.executed_due = {0, 1}
        gui.manual_completed_keys = set()

        try:
            module.datetime = FixedDatetime
            gui.run_auto_logout("28", first_check, handoff_at)

            self.assertEqual(cleared, [])
            self.assertEqual(gui.auto_logout_deadline, datetime(2026, 7, 10, 18, 20))
            self.assertEqual(gui.auto_logout_handoff_at, handoff_at)
            self.assertEqual(scheduled[0][0], 10 * 60 * 1000)
            self.assertIn("1 筆未完成", gui.duty_status_text.value)

            gui.duty_action_compare[2] = {"group": "done"}
            FixedDatetime.current = real_datetime(2026, 7, 10, 18, 20)
            scheduled[0][1]()
        finally:
            module.datetime = real_datetime

        self.assertEqual(cleared, ["system"])

    def test_completed_handoff_checkout_restores_auto_logout_timer(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        action = {
            "kind": "entry_log",
            "time": "18:00",
            "actor": "28",
            "source": "值班交接",
            "fields": {"出或入": "值退"},
        }
        scheduled: list[tuple[str, dict[str, object]]] = []
        submitted: list[int] = []
        gui.session = module.LoginSession(actor_no="28", user_id="user28", password="secret", verified=True)
        gui.duty_actions = [action]
        gui.duty_action_compare = {0: {"group": "done"}}
        gui.executed_due = {0}
        gui.submitting_indices = set()
        gui.failed_due_retry_after = {}
        gui.paused_due_indices = {}
        gui.manual_paused_due_indices = {}
        gui.duty_task_indices = lambda: [0]
        gui.sync_duty_compare_from_audit = lambda: None
        gui.action_datetime = lambda _action: datetime(2026, 7, 10, 18, 0)
        gui.ensure_auto_logout_scheduled = lambda actor, item: scheduled.append((actor, item))
        gui.submit_duty_action = lambda index, *_args, **_kwargs: submitted.append(index)

        gui.trigger_due_tasks(datetime(2026, 7, 10, 18, 0))

        self.assertEqual(scheduled, [("28", action)])
        self.assertEqual(submitted, [])

    def test_relogin_after_automatic_logout_does_not_reschedule_completed_handoff(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        handoff_at = datetime(2026, 7, 10, 18, 0)
        action = {
            "kind": "entry_log",
            "time": "18:00",
            "actor": "28",
            "source": "值班交接",
            "fields": {"出或入": "值退"},
        }
        scheduled: list[tuple[str, dict[str, object]]] = []
        gui.session = module.LoginSession(actor_no="28", user_id="user28", password="secret", verified=True)
        gui.auto_logout_login_started_at = datetime(2026, 7, 10, 18, 20)
        gui.duty_actions = [action]
        gui.duty_action_compare = {0: {"group": "done"}}
        gui.executed_due = {0}
        gui.submitting_indices = set()
        gui.duty_task_indices = lambda: [0]
        gui.sync_duty_compare_from_audit = lambda: None
        gui.action_datetime = lambda _action: handoff_at
        gui.ensure_auto_logout_scheduled = lambda actor, item: scheduled.append((actor, item))

        gui.trigger_due_tasks(datetime(2026, 7, 10, 18, 20))

        self.assertEqual(scheduled, [])

    def test_submit_login_failure_expires_session_and_stops_queues(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class Var:
            def __init__(self, value: str = "") -> None:
                self.value = value

            def set(self, value: str) -> None:
                self.value = value

            def get(self) -> str:
                return self.value

        sent_events: list[tuple[tuple[object, ...], dict[str, object]]] = []
        status_messages: list[str] = []
        notifications: list[tuple[str, str]] = []
        gui.session = module.LoginSession(actor_no="10", user_id="tyfd01000", password="old-pass", verified=True)
        gui.login_status = Var()
        gui.duty_actions = [{"kind": "entry_log", "time": "20:00", "fields": {}}]
        gui.submitting_indices = {0, 1}
        gui.submit_queues = {"entry": [("queued-entry",)], "work": [("queued-work",)]}
        gui.submit_worker_running = {"entry": True, "work": True}
        gui.work_submit_parallel_enabled = False
        gui.submit_needs_comparison_refresh = True
        gui.submit_comparison_refresh_dates = {"1150626"}
        gui.submit_comparison_refresh_scheduled = True
        gui.failed_due_retry_after = {}
        gui.auto_logout_after_id = None
        gui.auto_logout_deadline = None
        gui.auto_logout_actor_no = ""
        gui.pending_auto_logout_actor_no = ""
        gui.pending_auto_logout_deadline = None
        gui.action_completion_key = lambda _action: "completion-key"
        gui.log_trigger = lambda *_args, **_kwargs: None
        gui.send_sinposmart_backend_event = lambda *args, **kwargs: sent_events.append((args, kwargs))
        gui.export_issue_package = lambda **_kwargs: Path("issue.zip")
        gui.set_duty_status = lambda message, **_kwargs: status_messages.append(message)
        gui.notify_user = lambda title, message, **_kwargs: notifications.append((title, message))
        gui.update_login_panel = lambda: None
        gui.refresh_tasks = lambda: None
        gui.refresh_duty_tasks = lambda: None
        gui.after_cancel = lambda _after_id: None
        result_path = Path("runtime_outputs/form_tests/login_failed.json")

        gui._save_work_log_item_failed(
            0,
            "登入失敗：帳號或密碼可能已變更，請重新登入系統。",
            result_path,
            False,
            "due",
            "login_failed",
            "2026-06-26T20:00:00",
        )

        self.assertIsNone(gui.session)
        self.assertEqual(gui.submit_queues, {"entry": [], "work": []})
        self.assertEqual(gui.submit_worker_running, {"entry": False, "work": False})
        self.assertEqual(gui.submitting_indices, set())
        self.assertNotIn(0, gui.failed_due_retry_after)
        self.assertIn("登入狀態失效", gui.login_status.get())
        self.assertTrue(any(args and args[0] == "login_expired" for args, _kwargs in sent_events))
        self.assertTrue(any("登入失效" in message for message in status_messages))
        self.assertTrue(any("重新登入" in message for _title, message in notifications))

    def test_typed_saved_account_does_not_replace_manual_password(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class Var:
            def __init__(self, value: str = "") -> None:
                self.value = value

            def get(self) -> str:
                return self.value

            def set(self, value: str) -> None:
                self.value = value

        gui.saved_accounts = [{"actor_no": "10", "user_id": "tyfd01000", "password": "old-pass"}]
        gui.actor_no = Var("")
        gui.user_id = Var("tyfd01000")
        gui.password = Var("new-pass")
        gui.saved_account_choice = Var("")

        gui.sync_typed_account_choice()

        self.assertEqual(gui.actor_no.get(), "10")
        self.assertEqual(gui.user_id.get(), "tyfd01000")
        self.assertEqual(gui.password.get(), "new-pass")
        self.assertEqual(gui.saved_account_choice.get(), "tyfd01000 / 10番")

    def test_select_saved_account_still_fills_saved_password(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class Var:
            def __init__(self, value: str = "") -> None:
                self.value = value

            def get(self) -> str:
                return self.value

            def set(self, value: str) -> None:
                self.value = value

        gui.saved_accounts = [{"actor_no": "10", "user_id": "tyfd01000", "password": "old-pass"}]
        gui.actor_no = Var("")
        gui.user_id = Var("")
        gui.password = Var("")
        gui.saved_account_choice = Var("")

        gui.select_saved_account("tyfd01000", persist=False)

        self.assertEqual(gui.actor_no.get(), "10")
        self.assertEqual(gui.user_id.get(), "tyfd01000")
        self.assertEqual(gui.password.get(), "old-pass")

    def test_successful_existing_account_login_is_saved_even_when_remember_is_off(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        class BoolVar:
            def __init__(self, value: bool = False) -> None:
                self.value = value

            def get(self) -> bool:
                return self.value

        gui.saved_accounts = [{"actor_no": "10", "user_id": "tyfd01000", "password": "old-pass"}]
        gui.remember_login = BoolVar(False)

        self.assertTrue(gui.should_save_successful_login("10", "tyfd01000"))
        self.assertFalse(gui.should_save_successful_login("11", "tyfd01100"))

    def test_login_rejects_mismatched_detected_actor_no(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.actor_no_from_user_id = lambda _user_id: ""

        with self.assertRaises(module.LoginFailedError) as context:
            gui.resolve_verified_actor_no("11", "tyfd01510", "8")

        self.assertIn("8 番", str(context.exception))
        self.assertIn("11 番", str(context.exception))

    def test_sanitize_rebuilds_missing_schedule_actions(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        duty = "\u503c\u73ed"
        standby = "\u5099\u52e4"
        on_duty = "\u5728\u52e4"
        data = {
            "target_date": "1150628",
            "today": {
                "roc_date": "1150628",
                "rows": [{"slot": "08-12", "columns": {duty: ["11"], standby: ["3"]}}],
                "summary": {on_duty: ["3", "11"]},
                "staff": {"3": {"name": "三號"}, "11": {"name": "十一號"}},
            },
            "yesterday": {
                "roc_date": "1150627",
                "rows": [{"slot": "22-24", "columns": {duty: ["8"], standby: ["4"]}}],
                "summary": {on_duty: ["4", "8"]},
                "staff": {"4": {"name": "四號"}, "8": {"name": "八號"}},
            },
            "tomorrow": {"roc_date": "1150629", "rows": [], "summary": {}, "staff": {}},
            "cases": [],
            "yesterday_cases": [],
            "actions": [],
        }

        gui.sanitize_schedule_data(data)

        self.assertTrue(data["actions"])
        self.assertTrue(any(str(action.get("actor")) == "11" for action in data["actions"]))

    def test_sanitize_merges_missing_schedule_actions_when_actions_partial(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        duty = "\u503c\u73ed"
        standby = "\u5099\u52e4"
        on_duty = "\u5728\u52e4"
        handoff = "\u503c\u73ed\u4ea4\u63a5"
        duty_end = "\u503c\u9000"
        data = {
            "target_date": "1150628",
            "today": {
                "roc_date": "1150628",
                "rows": [{"slot": "08-12", "columns": {duty: ["11"], standby: ["3"]}}],
                "summary": {on_duty: ["3", "11"]},
                "staff": {"3": {"name": "三號"}, "11": {"name": "十一號"}},
            },
            "yesterday": {
                "roc_date": "1150627",
                "rows": [{"slot": "22-24", "columns": {duty: ["8"], standby: ["4"]}}],
                "summary": {on_duty: ["4", "8"]},
                "staff": {"4": {"name": "四號"}, "8": {"name": "八號"}},
            },
            "tomorrow": {"roc_date": "1150629", "rows": [], "summary": {}, "staff": {}},
            "cases": [],
            "yesterday_cases": [],
            "actions": [
                {
                    "kind": "entry_log",
                    "time": "08:00",
                    "actor": "8",
                    "target": "8",
                    "fields": {
                        "\u767b\u6253\u6642\u9593": "08:00",
                        "\u7cfb\u7d71\u5beb\u5165\u6642\u9593": "08:00",
                        "\u52e4\u52d9\u9805\u76ee": "\u503c\u73ed(\u5bbf)",
                        "\u51fa\u6216\u5165": duty_end,
                        "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede": duty_end,
                    },
                    "source": handoff,
                    "duplicate_key": f"entry:2026-06-28:8:{duty_end}:8",
                }
            ],
        }

        gui.sanitize_schedule_data(data)

        self.assertEqual(
            sum(1 for action in data["actions"] if action.get("duplicate_key") == f"entry:2026-06-28:8:{duty_end}:8"),
            1,
        )
        self.assertTrue(
            any(action.get("duplicate_key") == f"entry:2026-06-28:8:{duty}:11" for action in data["actions"])
        )
        self.assertTrue(
            any(action.get("duplicate_key") == f"work:2026-06-28:8:{handoff}:8" for action in data["actions"])
        )

    def test_duty_tasks_show_previous_handoff_work_for_incoming_actor(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        handoff = "\u503c\u73ed\u4ea4\u63a5"
        duty = "\u503c\u73ed"
        duty_end = "\u503c\u9000"

        class ActorNo:
            def get(self) -> str:
                return "11"

        gui.session = None
        gui.actor_no = ActorNo()
        gui.duty_data = {"target_date": "1150628"}
        gui.duty_action_compare = {}
        gui.submitting_indices = set()
        gui.paused_due_indices = {}
        gui.executed_due = set()
        gui.manual_completed_keys = set()
        gui.duty_actions = [
            {
                "kind": "entry_log",
                "time": "08:00",
                "actor": "8",
                "target": "8",
                "fields": {"\u51fa\u6216\u5165": duty_end, "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede": duty_end},
                "source": handoff,
            },
            {
                "kind": "entry_log",
                "time": "08:00",
                "actor": "8",
                "target": "11",
                "fields": {"\u51fa\u6216\u5165": duty, "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede": duty},
                "source": handoff,
            },
            {
                "kind": "work_log",
                "time": "08:00",
                "actor": "8",
                "target": "8",
                "fields": {"\u5de5\u4f5c\u6642\u9593": "08:00", "\u52e4\u52d9\u9805\u76ee": "\u503c\u73ed(\u5bbf)"},
                "source": handoff,
            },
        ]
        gui.duty_action_compare = {
            index: {"compare": "\u672a\u627e\u5230", "group": "todo", "matched": []}
            for index in range(len(gui.duty_actions))
        }

        indices = gui.duty_task_indices()
        status, tag, is_next_candidate = gui.resolve_duty_task_display(
            2,
            gui.duty_actions[2],
            gui.duty_action_compare[2],
            "11",
            datetime(2026, 6, 28, 8, 1),
        )

        self.assertIn(2, indices)
        self.assertEqual(status, "\u524d\u73ed\u624b\u52d5")
        self.assertEqual(tag, "waiting")
        self.assertFalse(is_next_candidate)

    def test_manual_pause_selected_marks_current_actor_task(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        status_messages: list[str] = []
        refreshes: list[str] = []
        gui.session = module.LoginSession(actor_no="10", user_id="tyfd01010", password="secret", verified=True)
        gui.duty_selected_iids = {"duty-0"}
        gui.duty_actions = [
            {"kind": "work_log", "time": "08:00", "actor": "10", "fields": {"勤務項目": "巡邏"}}
        ]
        gui.duty_action_compare = {0: {"compare": "未找到", "group": "todo", "matched": []}}
        gui.manual_paused_due_indices = {}
        gui.executed_due = set()
        gui.submitting_indices = set()
        gui.manual_completed_keys = set()
        gui.action_completion_key = lambda _action: "work-0"
        gui.action_datetime = lambda _action: datetime(2026, 7, 2, 8, 0)
        gui.is_auto_duty_action = lambda _action: True
        gui.compare_needs_manual_review = lambda _compare: False
        gui.set_duty_status = lambda message, **_kwargs: status_messages.append(message)
        gui.refresh_duty_tasks = lambda: refreshes.append("refresh")

        gui.manual_pause_selected()
        status, tag, is_next_candidate = gui.resolve_duty_task_display(
            0,
            gui.duty_actions[0],
            gui.duty_action_compare[0],
            "10",
            datetime(2026, 7, 2, 8, 1),
        )

        self.assertEqual(gui.manual_paused_due_indices, {0: "10"})
        self.assertEqual(status, "人員手動暫停")
        self.assertEqual(tag, "manual")
        self.assertFalse(is_next_candidate)
        self.assertTrue(any("人員手動暫停" in message for message in status_messages))
        self.assertEqual(refreshes, ["refresh"])

    def test_manual_pause_blocks_due_until_resume_schedule(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)

        submitted: list[tuple[int, str]] = []
        logged: list[tuple[int, str]] = []
        status_messages: list[str] = []
        gui.session = module.LoginSession(actor_no="10", user_id="tyfd01010", password="secret", verified=True)
        gui.duty_selected_iids = {"duty-0"}
        gui.duty_actions = [
            {"kind": "work_log", "time": "08:00", "actor": "10", "fields": {"勤務項目": "巡邏"}}
        ]
        gui.duty_action_compare = {0: {"compare": "未找到", "group": "todo", "matched": []}}
        gui.manual_paused_due_indices = {0: "10"}
        gui.paused_due_indices = {}
        gui.executed_due = set()
        gui.submitting_indices = set()
        gui.failed_due_retry_after = {}
        gui.sync_duty_compare_from_audit = lambda: None
        gui.duty_task_indices = lambda: [0]
        gui.action_datetime = lambda _action: datetime(2026, 7, 2, 8, 0)
        gui.action_target_roc_date = lambda _action: "1150702"
        gui.should_pause_due_action = lambda _action, _target_roc_date, now=None: ""
        gui.is_auto_duty_action = lambda _action: True
        gui.compare_needs_manual_review = lambda _compare: False
        gui.log_trigger = lambda index, _action, trigger_type, **_kwargs: logged.append((index, trigger_type))
        gui.submit_duty_action = lambda index, _action, **kwargs: submitted.append((index, kwargs["trigger_type"]))
        gui.set_duty_status = lambda message, **_kwargs: status_messages.append(message)
        gui.refresh_duty_tasks = lambda: None
        gui.schedule_pending_auto_logout_if_idle = lambda: None

        gui.trigger_due_tasks(datetime(2026, 7, 2, 8, 1))

        self.assertEqual(logged, [])
        self.assertEqual(submitted, [])

        gui.resume_selected_schedule()
        gui.trigger_due_tasks(datetime(2026, 7, 2, 8, 2))

        self.assertEqual(gui.manual_paused_due_indices, {})
        self.assertEqual(logged, [(0, "due")])
        self.assertEqual(submitted, [(0, "due")])
        self.assertTrue(any("繼續排程" in message for message in status_messages))

    def test_crossday_due_waits_for_comparison_before_submit(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        comparison_requests: list[tuple[str, list[str]]] = []
        submitted: list[int] = []
        action = {
            "kind": "entry_log",
            "time": "06:00",
            "actor": "10",
            "target": "25",
            "date_offset": 1,
            "fields": {"登打時間": "06:00", "系統寫入時間": "06:00", "出或入": "出", "領用事由及地點": "退勤"},
        }

        gui.session = module.LoginSession(actor_no="10", user_id="user10", password="secret", verified=True)
        gui.duty_data = {"target_date": "1150715"}
        gui.duty_actions = [action]
        gui.duty_action_compare = {0: {"compare": "未找到", "group": "todo", "matched": []}}
        gui.executed_due = set()
        gui.submitting_indices = set()
        gui.failed_due_retry_after = {}
        gui.paused_due_indices = {}
        gui.manual_paused_due_indices = {}
        gui.manual_resume_due_times = {}
        gui.comparison_waiting_due_indices = {}
        gui.comparison_running = False
        gui.duty_task_indices = lambda: [0]
        gui.sync_duty_compare_from_audit = lambda: None
        gui.action_datetime = lambda _action: datetime(2026, 7, 16, 6, 0)
        gui.action_target_roc_date = lambda _action: "1150716"
        gui.is_auto_duty_action = lambda _action: True
        gui.compare_needs_manual_review = lambda _compare: False
        gui.should_pause_due_action = lambda _action, _target_date, now=None: ""
        gui.comparison_data_available = lambda target_date: False
        gui.refresh_comparison_background = lambda target_date, _slot, comparison_dates=None, **_kwargs: comparison_requests.append((target_date, comparison_dates or []))
        gui.request_duty_task_refresh = lambda **_kwargs: None
        gui.submit_duty_action = lambda index, *_args, **_kwargs: submitted.append(index)

        gui.trigger_due_tasks(datetime(2026, 7, 16, 6, 0))

        self.assertEqual(submitted, [])
        self.assertIn(0, gui.comparison_waiting_due_indices)
        self.assertEqual(comparison_requests, [("1150715", ["1150714", "1150715", "1150716"])])

    def test_request_duty_task_refresh_coalesces_and_upgrades_to_full(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        scheduled: list[tuple[int, object]] = []
        refreshed: list[bool] = []

        gui.duty_task_refresh_after_id = None
        gui.duty_task_refresh_full_requested = False
        gui.after = lambda delay, callback: scheduled.append((delay, callback)) or "refresh-id"
        gui.refresh_duty_tasks = lambda full=False: refreshed.append(full)

        gui.request_duty_task_refresh()
        gui.request_duty_task_refresh(full=True)

        self.assertEqual(len(scheduled), 1)
        scheduled[0][1]()
        self.assertEqual(refreshed, [True])
        self.assertIsNone(gui.duty_task_refresh_after_id)

    def test_status_only_refresh_reuses_matching_task_cards(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        updates: list[tuple[object, ...]] = []
        rebuilds: list[str] = []

        class ActorNo:
            def get(self) -> str:
                return "10"

        class TextValue:
            def set(self, _value: str) -> None:
                pass

        class TaskList:
            def winfo_children(self) -> list[object]:
                return []

        action = {"kind": "work_log", "time": "08:00", "actor": "10", "target": "10", "fields": {"工作時間": "08:00", "勤務項目": "巡邏"}}
        gui.duty_task_list = TaskList()
        gui.session = None
        gui.actor_no = ActorNo()
        gui.logout_cleared = False
        gui.duty_selected_iids = set()
        gui.duty_actions = [action]
        gui.duty_action_compare = {0: {"compare": "未找到", "group": "todo", "matched": []}}
        gui.duty_visible_iids = ["duty-0"]
        gui.duty_card_rows = {"duty-0": object()}
        gui.duty_card_borders = {"duty-0": "#cbd5e1"}
        gui.duty_card_widgets = {"duty-0": {}}
        gui.sync_duty_compare_from_audit = lambda: None
        gui.duty_task_indices = lambda: [0]
        gui.resolve_duty_task_display = lambda *_args: ("正在登打", "running", False)
        gui.action_display_time = lambda _action: "08:00"
        gui.duty_task_card_time = lambda value: value
        gui.duty_task_columns = lambda _action: ("工作", "", "巡邏", "10番")
        gui.pending_previous_duty_count = lambda _actor: 0
        gui.update_duty_task_card = lambda *args, **kwargs: updates.append((args, kwargs))
        gui.create_duty_task_card = lambda **_kwargs: rebuilds.append("create")
        gui.reset_duty_task_scroll = lambda: rebuilds.append("scroll")
        gui.update_duty_card_selection = lambda: None
        gui.next_task_text = TextValue()
        gui.duty_status_text = TextValue()
        gui.active_duty_status_override = lambda: ""

        gui.refresh_duty_tasks(full=False)

        self.assertEqual(len(updates), 1)
        self.assertEqual(rebuilds, [])

    def test_manual_work_log_submit_uses_current_time(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        real_datetime = module.datetime

        class FixedDatetime(real_datetime):
            @classmethod
            def now(cls, tz=None):
                value = cls(2026, 7, 2, 9, 37)
                if tz is not None:
                    return value.replace(tzinfo=tz)
                return value

        action = {
            "kind": "work_log",
            "time": "08:00",
            "fields": {"工作時間": "08:00", "勤務項目": "巡邏"},
        }
        try:
            module.datetime = FixedDatetime
            updated = gui.action_for_manual_submit(action)
        finally:
            module.datetime = real_datetime

        self.assertIsNot(updated, action)
        self.assertEqual(updated["time"], "09:37")
        self.assertEqual(updated["submit_target_date"], "1150702")
        self.assertEqual(updated["fields"]["工作時間"], "09:37")
        self.assertEqual(action["fields"]["工作時間"], "08:00")

    def test_manual_entry_log_without_reason_whitelist_uses_current_time(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        real_datetime = module.datetime

        class FixedDatetime(real_datetime):
            @classmethod
            def now(cls, tz=None):
                value = cls(2026, 7, 2, 9, 41)
                if tz is not None:
                    return value.replace(tzinfo=tz)
                return value

        action = {
            "kind": "entry_log",
            "time": "08:05",
            "source": "昨日在勤且今日未在勤",
            "fields": {"出或入": "出", "領用事由及地點": "退勤", "登打時間": "08:05", "系統寫入時間": "08:05"},
        }
        try:
            module.datetime = FixedDatetime
            updated = gui.action_for_manual_submit(action)
        finally:
            module.datetime = real_datetime

        self.assertIsNot(updated, action)
        self.assertEqual(updated["time"], "09:41")
        self.assertEqual(updated["submit_target_date"], "1150702")
        self.assertEqual(updated["fields"]["登打時間"], "09:41")
        self.assertEqual(updated["fields"]["系統寫入時間"], "09:41")
        self.assertEqual(action["fields"]["系統寫入時間"], "08:05")

    def test_manual_submit_confirmation_mentions_current_time(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn("將使用按下「手動登打」時的當下時間登打。", source)
        self.assertEqual(source.count("將使用按下「手動登打」時的當下時間登打。"), 2)

    def test_resume_overdue_manual_pause_submits_resume_time(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        real_datetime = module.datetime

        class FixedDatetime(real_datetime):
            @classmethod
            def now(cls, tz=None):
                value = cls(2026, 7, 2, 8, 2)
                if tz is not None:
                    return value.replace(tzinfo=tz)
                return value

        submitted: list[dict[str, object]] = []
        gui.session = module.LoginSession(actor_no="10", user_id="tyfd01010", password="secret", verified=True)
        gui.duty_selected_iids = {"duty-0"}
        gui.duty_actions = [
            {"kind": "work_log", "time": "08:00", "actor": "10", "fields": {"工作時間": "08:00", "勤務項目": "巡邏"}}
        ]
        gui.duty_action_compare = {0: {"compare": "未找到", "group": "todo", "matched": []}}
        gui.manual_paused_due_indices = {0: "10"}
        gui.paused_due_indices = {}
        gui.executed_due = set()
        gui.submitting_indices = set()
        gui.failed_due_retry_after = {}
        gui.sync_duty_compare_from_audit = lambda: None
        gui.duty_task_indices = lambda: [0]
        gui.action_datetime = lambda _action: datetime(2026, 7, 2, 8, 0)
        gui.action_target_roc_date = lambda _action: "1150702"
        gui.should_pause_due_action = lambda _action, _target_roc_date, now=None: ""
        gui.is_auto_duty_action = lambda _action: True
        gui.compare_needs_manual_review = lambda _compare: False
        gui.log_trigger = lambda *_args, **_kwargs: None
        gui.submit_duty_action = lambda _index, action, **_kwargs: submitted.append(action)
        gui.set_duty_status = lambda *_args, **_kwargs: None
        gui.refresh_duty_tasks = lambda: None
        gui.schedule_pending_auto_logout_if_idle = lambda: None

        try:
            module.datetime = FixedDatetime
            gui.resume_selected_schedule()
            gui.trigger_due_tasks(datetime(2026, 7, 2, 8, 5))
        finally:
            module.datetime = real_datetime

        self.assertEqual(gui.manual_paused_due_indices, {})
        self.assertEqual(len(submitted), 1)
        self.assertEqual(submitted[0]["time"], "08:02")
        self.assertEqual(submitted[0]["submit_target_date"], "1150702")
        self.assertEqual(submitted[0]["fields"]["工作時間"], "08:02")
        self.assertEqual(gui.duty_actions[0]["fields"]["工作時間"], "08:00")

    def test_reset_duty_task_scroll_moves_hidden_canvas_to_top(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        calls: list[float] = []

        class Canvas:
            def yview_moveto(self, fraction: float) -> None:
                calls.append(fraction)

        class TaskList:
            _parent_canvas = Canvas()

        gui.duty_task_list = TaskList()

        gui.reset_duty_task_scroll()

        self.assertEqual(calls, [0])

    def test_post_submit_verification_requeries_without_cache_and_fails_when_missing(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        cache_args: list[object] = []

        def duplicate_matches(_driver: object, _action: dict[str, object], _target_date: str, duplicate_cache: object = None) -> list[str]:
            cache_args.append(duplicate_cache)
            return []

        gui.duplicate_matches_before_submit = duplicate_matches

        with self.assertRaisesRegex(RuntimeError, "登打後.*查到"):
            gui.verify_action_saved_after_submit(None, {"kind": "work_log"}, "1150628")

        self.assertEqual(cache_args, [None])

    def test_submit_worker_verifies_saved_record_before_marking_success(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn("def verify_action_saved_after_submit", source)
        self.assertLess(
            source.index("self.verify_action_saved_after_submit(driver, action, target_date)"),
            source.index('result["stage"] = "submitted" if save else "filled"'),
        )

    def test_sinposmart_event_worker_persists_pending_before_posting(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertLess(
            source.index("write_pending_sinposmart_backend_events(pending)"),
            source.index("response = post_sinposmart_backend_event(entry)"),
        )
        self.assertIn("SINPOSMART_BACKEND_EVENT_LOCK = threading.Lock()", source)
        self.assertLess(
            source.index("with SINPOSMART_BACKEND_EVENT_LOCK:"),
            source.index("pending = load_pending_sinposmart_backend_events()"),
        )

    def test_duty_sheet_preflight_false_result_reports_tool_error(self) -> None:
        source = legacy_tk_path("duty_sheet_automation.py").read_text(encoding="utf-8-sig")

        self.assertIn("automation_result = legacy.start_automation", source)
        self.assertIn("if automation_result is False:", source)
        self.assertLess(source.index("if automation_result is False:"), source.index("success = True"))
        self.assertLess(source.index("if automation_result is False:"), source.index("on_finish("))
        self.assertIn("on_error(error)", source)

    def test_login_success_auto_syncs_credentials_without_export_gate(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")
        tree = ast.parse(source)
        login_fn = next(
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.FunctionDef) and node.name == "_login_succeeded"
        )

        self.assertNotIn("CREDENTIAL_EXPORT_USER_ID", source)
        self.assertNotIn("self.credential_export_button = ctk.CTkButton", source)
        self.assertIn("sync_credentials_after_login", {node.func.attr for node in ast.walk(login_fn) if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)})

    def test_auto_credential_sync_sends_only_verified_login_person(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        captured: list[tuple[dict[str, object], int, bool]] = []

        class ImmediateThread:
            def __init__(self, target: object, args: tuple[object, ...] = (), daemon: bool | None = None) -> None:
                self.target = target
                self.args = args

            def start(self) -> None:
                self.target(*self.args)

        gui.saved_accounts = [
            {"actor_no": "8", "user_id": "tyfd01510", "password": "old-pass", "display_name": "8番 tyfd01510", "id_number": "B123017532"},
            {"actor_no": "9", "user_id": "tyfd00009", "password": "pass9", "display_name": "9番 王小明", "name": "王小明"},
        ]
        gui.session = module.LoginSession("8", "tyfd01510", "new-pass", verified=True)
        gui.duty_staff = {}
        gui.staff = {}
        gui.duty_data = {}
        gui.data = {}
        gui.work_log_rows_for_person = lambda _actor_no, _name: []
        gui._credential_sync_send_worker = lambda payload, count, notify_user=True: captured.append((payload, count, notify_user))
        gui._credential_sync_send_failed = lambda *_args, **_kwargs: None
        original_enabled = module.credential_sync_enabled
        original_thread = module.threading.Thread
        module.credential_sync_enabled = lambda: True
        module.threading.Thread = ImmediateThread
        try:
            gui.sync_credentials_after_login("8", "tyfd01510", "new-pass", name="曾彥綸")
        finally:
            module.credential_sync_enabled = original_enabled
            module.threading.Thread = original_thread

        self.assertEqual(len(captured), 1)
        payload, count, notify_user = captured[0]
        accounts = payload["accounts"]
        synced_8 = next(account for account in accounts if account["actor_no"] == "8")

        self.assertEqual(count, 1)
        self.assertEqual([account["user_id"] for account in accounts], ["tyfd01510"])
        self.assertFalse(notify_user)
        self.assertEqual(synced_8["user_id"], "tyfd01510")
        self.assertEqual(synced_8["password"], "new-pass")
        self.assertEqual(synced_8["display_name"], "8番 曾彥綸")
        self.assertEqual(synced_8["name"], "曾彥綸")
        self.assertEqual(synced_8["id_number"], "B123017532")

    def test_identify_logged_in_actor_keeps_greeting_name_without_staff_map(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.staff = {}
        gui.actor_no_from_name = lambda _name: ""

        class SwitchTo:
            def default_content(self) -> None:
                return None

            def frame(self, _frame: object) -> None:
                return None

        class Driver:
            switch_to = SwitchTo()

            def execute_script(self, _script: str) -> str:
                return "曾彥綸,您好"

            def find_elements(self, _by: str, _value: str) -> list[object]:
                return []

        self.assertEqual(gui.identify_logged_in_actor(Driver()), ("", "曾彥綸"))

    def test_auto_credential_sync_rejects_unverified_or_mismatched_login(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.saved_accounts = [{"actor_no": "9", "user_id": "other-user", "password": "other-password"}]
        gui.saved_accounts_for_credential_sync = mock.Mock(return_value=[{
            "actor_no": "9", "user_id": "other-user", "password": "other-password",
            "display_name": "其他人", "name": "其他人", "id_number": ""}])
        gui.account_for_credential_sync = mock.Mock(return_value={
            "actor_no": "8", "user_id": "current-user", "password": "current-password",
            "display_name": "當次登入者", "name": "測試員", "id_number": ""})
        gui._credential_sync_send_worker = mock.Mock()
        gui._credential_sync_send_failed = mock.Mock()
        for session in (None, module.LoginSession("8", "current-user", "current-password"),
                        module.LoginSession("9", "other-user", "other-password", verified=True)):
            with self.subTest(session=session is not None):
                gui.session = session
                with mock.patch.object(module, "credential_sync_enabled", return_value=True), \
                        mock.patch.object(module.threading, "Thread") as thread:
                    gui.sync_credentials_after_login("8", "current-user", "current-password")
                thread.assert_not_called()

    def test_manual_credential_sync_sends_verified_session_instead_of_saved_accounts(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.session = module.LoginSession("8", "current-user", "verified-password", verified=True)
        gui.saved_accounts = [{"actor_no": "8", "user_id": "current-user", "password": "stale-password"},
                              {"actor_no": "9", "user_id": "other-user", "password": "other-password"}]
        gui.saved_accounts_for_credential_sync = mock.Mock(return_value=[{
            **account, "display_name": "測試員", "name": "測試員", "id_number": ""}
            for account in gui.saved_accounts])
        gui.account_for_credential_sync = mock.Mock(side_effect=lambda actor, user, password, **_kwargs: {
            "actor_no": actor, "user_id": user, "password": password,
            "display_name": "當次登入者", "name": "測試員", "id_number": ""})
        gui.status_text = mock.Mock()
        with mock.patch.object(module, "credential_sync_enabled", return_value=True), \
                mock.patch.object(module.messagebox, "askyesno", return_value=True), \
                mock.patch.object(module.threading, "Thread") as thread:
            gui.sync_current_account_dialog()
        payload, count, notify_user = thread.call_args.kwargs["args"]
        self.assertEqual(count, 1)
        self.assertTrue(notify_user)
        self.assertEqual([row["user_id"] for row in payload["accounts"]], ["current-user"])
        self.assertEqual(payload["accounts"][0]["password"], "verified-password")

    def test_credential_sync_worker_supports_silent_background_mode(self) -> None:
        source = legacy_tk_path("duty_gui.py").read_text(encoding="utf-8-sig")

        self.assertIn("notify_user: bool = True", source)
        self.assertIn("notify_user=False", source)
        self.assertIn("if notify_user:", source)

    def test_manual_credential_sync_rechecks_session_after_each_dialog(self) -> None:
        module = duty_gui_module()
        for stage in ("relay_confirm", "export_confirm", "export_path"):
            for change in ("logout", "switch", "password", "unverified"):
                with self.subTest(stage=stage, change=change):
                    gui = object.__new__(module.DutyGui)
                    gui.session = module.LoginSession("8", "current-user", "verified-password", verified=True)
                    gui.account_for_credential_sync = mock.Mock(side_effect=lambda actor, user, password, **_kwargs: {
                        "actor_no": actor, "user_id": user, "password": password,
                        "display_name": "當次登入者", "name": "測試員", "id_number": ""})
                    gui.status_text = mock.Mock()

                    def change_session(*_args):
                        if change == "logout":
                            gui.session = None
                        elif change == "switch":
                            gui.session = module.LoginSession("9", "other-user", "other-password", verified=True)
                        elif change == "password":
                            gui.session.password = "changed-password"
                        else:
                            gui.session.verified = False
                        return True

                    def choose_path(**_kwargs):
                        if stage == "export_path":
                            change_session()
                        return "fixture-credential-export.json"

                    with mock.patch.object(module, "credential_sync_enabled", return_value=stage == "relay_confirm"), \
                            mock.patch.object(module.messagebox, "askyesno", side_effect=(
                                change_session if stage != "export_path" else lambda *_args: True)), \
                            mock.patch.object(module.messagebox, "showinfo"), \
                            mock.patch.object(module.filedialog, "asksaveasfilename", side_effect=choose_path) as choose, \
                            mock.patch.object(module.threading, "Thread") as thread, \
                            mock.patch.object(Path, "write_text") as write:
                        gui.sync_current_account_dialog()
                    thread.assert_not_called()
                    write.assert_not_called()
                    if stage == "export_confirm":
                        choose.assert_not_called()

    def test_manual_credential_export_keeps_unchanged_verified_session(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.session = module.LoginSession("8", "current-user", "verified-password", verified=True)
        gui.account_for_credential_sync = mock.Mock(side_effect=lambda actor, user, password, **_kwargs: {
            "actor_no": actor, "user_id": user, "password": password,
            "display_name": "當次登入者", "name": "測試員", "id_number": ""})
        gui.status_text = mock.Mock()
        with mock.patch.object(module, "credential_sync_enabled", return_value=False), \
                mock.patch.object(module.messagebox, "askyesno", return_value=True), \
                mock.patch.object(module.messagebox, "showinfo"), \
                mock.patch.object(module.filedialog, "asksaveasfilename", return_value="fixture-credential-export.json"), \
                mock.patch.object(Path, "write_text") as write:
            gui.sync_current_account_dialog()
        payload = json.loads(write.call_args.args[0])
        self.assertEqual([row["user_id"] for row in payload["accounts"]], ["current-user"])
        self.assertEqual(payload["accounts"][0]["password"], "verified-password")

    def test_build_driver_retries_only_with_its_own_temporary_profile(self) -> None:
        module = duty_rehearsal_module()
        profiles = [Path("C:/Temp/duty_gui_first"), Path("C:/Temp/duty_gui_second")]
        cleaned: list[Path] = []
        driver = mock.Mock()

        with mock.patch.object(module, "chrome_start_attempts", return_value=2), mock.patch.object(
            module, "duty_browser_profile_dir", side_effect=profiles
        ), mock.patch.object(
            module, "prune_stale_duty_browser_profiles", return_value=0
        ), mock.patch.object(
            module,
            "create_webdriver_chrome_with_timeout",
            side_effect=[module.WebDriverException("Chrome failed to start"), driver],
        ) as start_chrome, mock.patch.object(
            module, "cleanup_duty_browser_startup_failure", side_effect=cleaned.append
        ), mock.patch.object(module.time, "sleep"
        ):
            result = module.build_driver(
                headless=False,
                option_arguments=("--test-duty-option",),
                page_load_strategy="none",
            )

        self.assertIs(result, driver)
        self.assertEqual(cleaned, [profiles[0]])
        first_options = start_chrome.call_args_list[0].args[0]
        second_options = start_chrome.call_args_list[1].args[0]
        self.assertIn(f"--user-data-dir={profiles[0]}", first_options.arguments)
        self.assertIn(f"--user-data-dir={profiles[1]}", second_options.arguments)
        self.assertIn("--test-duty-option", first_options.arguments)
        self.assertIn("--test-duty-option", second_options.arguments)
        self.assertEqual(first_options.page_load_strategy, "none")
        self.assertEqual(second_options.page_load_strategy, "none")
        self.assertNotEqual(profiles[0], profiles[1])
        self.assertEqual(getattr(driver, "_sinposmart_duty_browser_profile"), str(profiles[1]))

    def test_build_driver_passes_only_explicit_preferences_as_a_copy(self) -> None:
        module = duty_rehearsal_module()
        preferences = {"credentials_enable_service": False, "profile.password_manager_enabled": False}
        for supplied in (False, True):
            with self.subTest(supplied=supplied):
                driver = mock.Mock()
                with mock.patch.object(module, "chrome_start_attempts", return_value=1), mock.patch.object(
                    module, "duty_browser_profile_dir", return_value=Path("C:/Temp/duty_gui_preferences")
                ), mock.patch.object(
                    module, "create_webdriver_chrome_with_timeout", return_value=driver
                ) as start_chrome:
                    result = module.build_driver(
                        headless=True,
                        prune_profiles=False,
                        write_diagnostics=False,
                        **({"preferences": preferences} if supplied else {}),
                    )

                self.assertIs(result, driver)
                options = start_chrome.call_args.args[0]
                if supplied:
                    self.assertEqual(options.experimental_options["prefs"], preferences)
                    self.assertIsNot(options.experimental_options["prefs"], preferences)
                else:
                    self.assertNotIn("prefs", options.experimental_options)

    def test_prewrite_session_open_retries_with_a_fresh_driver_only(self) -> None:
        module = duty_rehearsal_module()
        first_driver = object()
        second_driver = object()
        cleaned: list[object] = []
        diagnostics: list[tuple[str, str, int]] = []

        def initialize(driver: object) -> None:
            if driver is first_driver:
                raise module.WebDriverException("invalid session id")

        with mock.patch.object(
            module,
            "_write_duty_browser_startup_diagnostic",
            side_effect=lambda event, **payload: diagnostics.append(
                (event, payload["category"], payload["attempts"])
            ),
        ):
            result = module.retry_duty_browser_session_open(
                mock.Mock(side_effect=[first_driver, second_driver]),
                initialize,
                cleanup=cleaned.append,
            )

        self.assertIs(result, second_driver)
        self.assertEqual(cleaned, [first_driver])
        self.assertEqual(
            diagnostics,
            [
                ("session_open_retry", "browser_session_open", 1),
                ("session_open_recovered", "browser_session_open", 2),
            ],
        )

    def test_prewrite_session_open_does_not_retry_page_timeout(self) -> None:
        module = duty_rehearsal_module()
        from selenium.common.exceptions import TimeoutException

        driver = object()
        factory = mock.Mock(return_value=driver)
        cleanup = mock.Mock()

        with self.assertRaises(TimeoutException):
            module.retry_duty_browser_session_open(
                factory,
                lambda _driver: (_ for _ in ()).throw(TimeoutException("page timed out")),
                cleanup=cleanup,
            )

        factory.assert_called_once_with()
        cleanup.assert_called_once_with(driver)

    def test_duty_number_popup_preflight_rebuilds_once_before_writes(self) -> None:
        module = legacy_duty_sheet_module()
        first_driver = object()
        second_driver = object()
        opened: list[object] = []
        cleaned: list[object] = []
        attempts: list[object] = []

        def open_browser() -> object:
            driver = first_driver if not opened else second_driver
            opened.append(driver)
            return driver

        def preflight(driver: object) -> None:
            attempts.append(driver)
            if driver is first_driver:
                raise module.DutyNumberPopupOpenError(
                    {"button_found": True, "window_count_before": 1, "window_count_after": 1}
                )

        result = module.retry_duty_number_popup_preflight(
            open_browser,
            preflight,
            cleanup=cleaned.append,
        )

        self.assertIs(result, second_driver)
        self.assertEqual(opened, [first_driver, second_driver])
        self.assertEqual(attempts, [first_driver, second_driver])
        self.assertEqual(cleaned, [first_driver])

    def test_duty_number_popup_preflight_rebuilds_once_after_prewrite_query_timeout(self) -> None:
        module = legacy_duty_sheet_module()
        first_driver = object()
        second_driver = object()
        opened: list[object] = []
        cleaned: list[object] = []
        attempts: list[object] = []

        def open_browser() -> object:
            driver = first_driver if not opened else second_driver
            opened.append(driver)
            return driver

        def preflight(driver: object) -> None:
            attempts.append(driver)
            if driver is first_driver:
                raise module.DutyQueryPrewriteTimeout("勤務表查詢逾時")

        result = module.retry_duty_number_popup_preflight(
            open_browser,
            preflight,
            cleanup=cleaned.append,
        )

        self.assertIs(result, second_driver)
        self.assertEqual(opened, [first_driver, second_driver])
        self.assertEqual(attempts, [first_driver, second_driver])
        self.assertEqual(cleaned, [first_driver])

    def test_duty_number_popup_preflight_rebuilds_once_after_prewrite_delete_unavailable(self) -> None:
        module = legacy_duty_sheet_module()
        first_driver = object()
        second_driver = object()
        opened: list[object] = []
        cleaned: list[object] = []
        attempts: list[object] = []

        def open_browser() -> object:
            driver = first_driver if not opened else second_driver
            opened.append(driver)
            return driver

        def preflight(driver: object) -> None:
            attempts.append(driver)
            if driver is first_driver:
                raise module.DutyExistingDeletePrewriteUnavailable(
                    "既有勤務表刪除按鈕未就緒"
                )

        result = module.retry_duty_number_popup_preflight(
            open_browser,
            preflight,
            cleanup=cleaned.append,
        )

        self.assertIs(result, second_driver)
        self.assertEqual(opened, [first_driver, second_driver])
        self.assertEqual(attempts, [first_driver, second_driver])
        self.assertEqual(cleaned, [first_driver])

    def test_duty_number_popup_preflight_stops_after_second_failure(self) -> None:
        module = legacy_duty_sheet_module()
        first_driver = object()
        second_driver = object()
        opened: list[object] = []
        cleaned: list[object] = []

        def open_browser() -> object:
            driver = first_driver if not opened else second_driver
            opened.append(driver)
            return driver

        def preflight(_driver: object) -> None:
            raise module.DutyNumberPopupOpenError(
                {"button_found": False, "window_count_before": 1, "window_count_after": 1}
            )

        with self.assertRaises(module.DutyNumberPopupOpenError) as caught:
            module.retry_duty_number_popup_preflight(
                open_browser,
                preflight,
                cleanup=cleaned.append,
            )

        self.assertEqual(opened, [first_driver, second_driver])
        self.assertEqual(cleaned, [first_driver, second_driver])
        self.assertIn("預檢重試：2/2", str(caught.exception))
        self.assertIn("未開始勤務表刪除、填寫或儲存", str(caught.exception))

    def test_duty_sheet_popup_setup_never_deletes_after_configuration(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("def start_automation(")
        flow = source[start:]
        post_setup_flow = flow[flow.index("step_config_popups("):]

        self.assertLess(post_setup_flow.index("step_config_popups("), post_setup_flow.index("wait_for_duty_result_grid("))
        self.assertLess(post_setup_flow.index("wait_for_duty_result_grid("), post_setup_flow.index('report_stage("duty_fill")'))
        self.assertNotIn('"_btnDelete"', post_setup_flow)
        self.assertIn("retry_duty_number_popup_preflight(", flow)
        self.assertIn("duty_number_popup_preflight", flow)

    def test_duty_sheet_capture_and_line_notification_failures_are_separate(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index('report_stage("report")')
        end = source.index("# 計算總花費秒數", start)
        report_flow = source[start:end]

        self.assertIn(
            "capture_future = capture_executor.submit(capture_duty_sheet_images, excel_path, target_date)",
            source,
        )
        self.assertIn("except Exception as capture_error:", report_flow)
        self.assertIn("勤務表截圖保存失敗", report_flow)
        self.assertIn("except Exception as notify_error:", report_flow)
        self.assertIn("LINE 截圖通知失敗", report_flow)
        self.assertNotIn("勤務表截圖或 LINE 通知失敗", report_flow)

    def test_duty_query_click_waiter_retries_until_target_date_frame_is_ready(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class QueryButtonWait:
            def until(self, predicate):
                for _ in range(3):
                    if predicate(driver):
                        return True
                raise AssertionError("查詢按鈕未在測試期間就緒")

        driver.execute_script.side_effect = [False, False, True]

        self.assertTrue(
            module.click_duty_query_button_when_ready(
                driver,
                QueryButtonWait(),
                "1150827",
            )
        )

        self.assertEqual(driver.execute_script.call_count, 3)
        click_scripts = [call.args[0] for call in driver.execute_script.call_args_list]
        self.assertTrue(
            all(
                "_txtTaskDate" in script and "_btnQuery" in script
                for script in click_scripts
            )
        )
        self.assertTrue(
            all(call.args[1] == "1150827" for call in driver.execute_script.call_args_list)
        )

    def test_duty_query_completion_waiter_polls_until_blank_setup_form_is_ready(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class QueryCompletionWait:
            def until(self, predicate):
                for _ in range(3):
                    result = predicate(driver)
                    if result:
                        return result
                raise AssertionError("查詢後的空白勤務表未在測試期間就緒")

        driver.execute_script.side_effect = [False, False, "setup"]

        self.assertEqual(
            module.wait_for_duty_query_completion(driver, QueryCompletionWait()),
            "setup",
        )

        self.assertEqual(driver.execute_script.call_count, 3)
        readiness_scripts = [call.args[0] for call in driver.execute_script.call_args_list]
        self.assertTrue(all("設定勤務項目" in script for script in readiness_scripts))
        self.assertTrue(all("設定勤務番號" in script for script in readiness_scripts))

    def test_duty_query_completion_returns_existing_for_existing_grid(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class QueryCompletionWait:
            def until(self, predicate):
                result = predicate(driver)
                if result:
                    return result
                raise AssertionError("既有勤務表未在測試期間辨識")

        driver.execute_script.return_value = "existing"

        self.assertEqual(
            module.wait_for_duty_query_completion(driver, QueryCompletionWait()),
            "existing",
        )

    def test_duty_query_completion_waiter_stops_before_popup_preflight_when_timed_out(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class TimedOutQueryCompletionWait:
            def until(self, predicate):
                predicate(driver)
                raise module.TimeoutException("query document did not reload")

        driver.execute_script.return_value = False

        with self.assertRaisesRegex(
            RuntimeError,
            "勤務表查詢逾時.*設定頁.*既有勤務表",
        ):
            module.wait_for_duty_query_completion(driver, TimedOutQueryCompletionWait())

    def test_duty_query_completion_marks_prewrite_timeout_for_safe_retry(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class TimedOutQueryCompletionWait:
            def until(self, predicate):
                predicate(driver)
                raise module.TimeoutException("query document did not reload")

        driver.execute_script.return_value = False

        with self.assertRaises(module.DutyQueryPrewriteTimeout):
            module.wait_for_duty_query_completion(
                driver,
                TimedOutQueryCompletionWait(),
                prewrite=True,
            )

    def test_duty_query_click_marks_prewrite_timeout_for_safe_retry(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class TimedOutQueryClickWait:
            def until(self, predicate):
                predicate(driver)
                raise module.TimeoutException("query button was not ready")

        driver.execute_script.return_value = False

        with self.assertRaises(module.DutyQueryPrewriteTimeout):
            module.click_duty_query_button_when_ready(
                driver,
                TimedOutQueryClickWait(),
                "1150828",
                prewrite=True,
            )

    def test_duty_existing_delete_marks_unavailable_button_for_safe_retry(self) -> None:
        module = legacy_duty_sheet_module()

        with mock.patch.object(
            module,
            "click_duty_existing_delete_and_accept_alert",
            return_value=False,
        ):
            with self.assertRaises(module.DutyExistingDeletePrewriteUnavailable):
                module.ensure_existing_duty_delete(object())

    def test_duty_result_grid_waiter_polls_after_setup_is_complete(self) -> None:
        module = legacy_duty_sheet_module()
        driver = object()

        class DutyGridWait:
            def until(self, predicate):
                for _ in range(3):
                    if predicate(driver):
                        return True
                raise AssertionError("勤務設定完成後的表格未在測試期間就緒")

        with mock.patch.object(module, "super_js_execute", side_effect=[False, False, True]) as execute:
            self.assertTrue(module.wait_for_duty_result_grid(driver, DutyGridWait()))

        self.assertEqual(execute.call_count, 3)
        self.assertTrue(
            all(
                call.args == (driver, "_pln_8_1", "exists")
                for call in execute.call_args_list
            )
        )

    def test_duty_query_waits_for_current_date_frame_before_popup_preflight(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("def open_duty_browser_for_popup_preflight():")
        flow = source[start:source.index("def verify_duty_number_popup", start)]
        readiness_start = source.index("def wait_for_duty_query_completion(")
        readiness_end = source.index("def read_duty_number_leave_rows", readiness_start)
        readiness_source = source[readiness_start:readiness_end]

        self.assertLess(
            flow.index('super_js_execute(candidate, "_txtTaskDate", "set", target_date)'),
            flow.index("click_duty_query_button_when_ready("),
        )
        self.assertLess(
            flow.index("click_duty_query_button_when_ready("),
            flow.index("wait_for_duty_query_completion("),
        )
        self.assertNotIn("wait_for_duty_query_button_ready(", flow)
        self.assertNotIn('super_js_execute(candidate, "_btnQuery", "click")', flow)
        self.assertNotIn("mark_duty_query_document(", flow)
        self.assertIn("_btnOpenWinTaskCode", readiness_source)
        self.assertIn("_btnOpenWinUserNo", readiness_source)
        self.assertIn("設定勤務項目", readiness_source)
        self.assertIn("設定勤務番號", readiness_source)

    def test_duty_setup_reloads_grid_without_a_second_query_or_delete(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("if popup_main_window:")
        end = source.index('report_stage("vehicle_form_open")', start)
        duty_flow = source[start:end]

        self.assertNotIn('super_js_execute(driver, "_btnQuery", "click")', duty_flow)
        self.assertNotIn('"_btnDelete"', duty_flow)
        self.assertLess(
            duty_flow.index("step_config_popups("),
            duty_flow.index("wait_for_duty_result_grid("),
        )
        self.assertLess(
            duty_flow.index("wait_for_duty_result_grid("),
            duty_flow.index('report_stage("duty_fill")'),
        )

    def test_existing_duty_query_deletes_before_popup_setup(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("def open_duty_browser_for_popup_preflight():")
        query_flow = source[start:source.index("def verify_duty_number_popup", start)]

        self.assertLess(
            query_flow.index('query_mode = wait_for_duty_query_completion('),
            query_flow.index('if query_mode == "existing":'),
        )
        self.assertLess(
            query_flow.index('if query_mode == "existing":'),
            query_flow.index("ensure_existing_duty_delete(candidate)"),
        )
        self.assertLess(
            query_flow.index("ensure_existing_duty_delete(candidate)"),
            query_flow.index("return candidate"),
        )
        self.assertEqual(
            query_flow.count("query_mode = wait_for_duty_query_completion("),
            2,
        )
        self.assertLess(
            query_flow.index("ensure_existing_duty_delete(candidate)"),
            query_flow.rindex("query_mode = wait_for_duty_query_completion("),
        )
        self.assertNotIn('super_js_execute(candidate, "_btnDelete", "click")', query_flow)

    def test_existing_duty_delete_accepts_confirmation_after_null_click_result(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.execute_script.return_value = None
        with mock.patch.object(module, "accept_pending_alerts") as accept_alerts:
            module.ensure_existing_duty_delete(driver)
        driver.execute_script.assert_called_once()
        accept_alerts.assert_called_once_with(driver, timeout=3)

    def test_existing_duty_delete_never_retries_an_unconfirmed_null_result(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.execute_script.return_value = None
        with mock.patch.object(module, "WebDriverWait") as wait:
            wait.return_value.until.side_effect = module.TimeoutException()
            with self.assertRaisesRegex(RuntimeError, "刪除結果尚未確認") as caught:
                module.ensure_existing_duty_delete(driver)
        self.assertNotIsInstance(caught.exception, module.DutyExistingDeletePrewriteUnavailable)
        driver.execute_script.assert_called_once()

    def test_duty_delete_completion_waits_past_the_old_grid(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.execute_script.side_effect = ["existing", False, "setup"]
        wait = module.WebDriverWait(driver, 1, poll_frequency=0.01)
        self.assertEqual(
            module.wait_for_duty_query_completion(driver, wait, expected_mode="setup"),
            "setup",
        )
        self.assertEqual(driver.execute_script.call_count, 3)

    def test_existing_duty_delete_accepts_alert_that_interrupts_the_click(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.execute_script.side_effect = module.UnexpectedAlertPresentException(
            "delete confirmation opened"
        )

        with mock.patch.object(
            module,
            "accept_pending_alerts",
            return_value=["確定要刪除舊勤務表嗎？"],
        ) as accept_alerts:
            self.assertTrue(
                module.click_duty_existing_delete_and_accept_alert(driver)
            )

        accept_alerts.assert_called_once_with(driver, timeout=3)
        delete_script = driver.execute_script.call_args.args[0]
        self.assertIn("_btnDelete", delete_script)
        self.assertIn("deleteButton.click()", delete_script)
        self.assertNotIn("onclick()", delete_script)

    def test_duty_number_popup_waits_for_a_new_handle_without_double_clicking(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("def _click_duty_number_popup_button(")
        end = source.index("def preflight_duty_number_popup(", start)
        popup_source = source[start:end]

        self.assertIn("handle for handle in candidate.window_handles if handle not in before_handles", popup_source)
        self.assertIn("element.click();", popup_source)
        self.assertNotIn("element.onclick()", popup_source)

    def test_popup_save_recovers_when_current_popup_window_has_closed(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.window_handles = ["main-window"]
        driver.switch_to.window.side_effect = [
            module.NoSuchWindowException("popup already closed"),
            None,
        ]

        class PopupCloseWait:
            def until(self, predicate):
                for _ in range(2):
                    if predicate(driver):
                        return True
                raise AssertionError("主視窗未在測試期間恢復")

        self.assertTrue(
            module.wait_for_popup_close_and_return_to_main(
                driver,
                PopupCloseWait(),
                "main-window",
                "popup-window",
            )
        )

        self.assertEqual(
            driver.switch_to.window.call_args_list,
            [mock.call("main-window"), mock.call("main-window")],
        )

    def test_duty_number_and_external_popups_share_main_window_recovery(self) -> None:
        source = (package_dir() / "duty_sheet_legacy" / "sinposmart_1.py").read_text(
            encoding="utf-8-sig"
        )
        start = source.index("def save_duty_number_popup(")
        popup_flow = source[start:source.index("def step_select_vehicles_popup(", start)]

        self.assertIn(
            "save_duty_number_popup(driver, wait, daily_commander, main_window)",
            popup_flow,
        )
        self.assertEqual(
            popup_flow.count("wait_for_popup_close_and_return_to_main("),
            2,
        )

    def test_external_setup_button_waits_for_its_named_button_before_clicking(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        class ExternalButtonWait:
            def until(self, predicate):
                for _ in range(3):
                    if predicate(driver):
                        return True
                raise AssertionError("設定勤務項目按鈕未在測試期間就緒")

        driver.execute_script.side_effect = [False, False, True]

        self.assertTrue(
            module.click_duty_external_setup_button_when_ready(driver, ExternalButtonWait())
        )

        self.assertEqual(driver.execute_script.call_count, 3)
        button_scripts = [call.args[0] for call in driver.execute_script.call_args_list]
        self.assertTrue(
            all(
                "_btnOpenWinTaskCode" in script and "設定勤務項目" in script and
                "element.click()" in script
                for script in button_scripts
            )
        )

    def test_external_popup_failure_stops_before_external_settings_are_written(self) -> None:
        module = legacy_duty_sheet_module()
        driver = mock.Mock()

        with mock.patch.object(module, "save_duty_number_popup"), mock.patch.object(
            module, "step_prepare_content", return_value=True
        ), mock.patch.object(
            module,
            "click_duty_external_setup_button_when_ready",
            side_effect=RuntimeError("設定勤務項目按鈕逾時"),
        ):
            with self.assertRaisesRegex(RuntimeError, "設定勤務項目按鈕逾時"):
                module.step_config_popups(driver, mock.Mock(), ["防溺車巡"], "2", "main")

        driver.find_element.assert_not_called()

    def test_external_popup_sets_two_rescue_items_and_clears_unused_fields(self) -> None:
        module = legacy_duty_sheet_module()
        for names in (["勤務一", "勤務二", "勤務三", "勤務四", "勤務五"], ["勤務一"], []):
            with self.subTest(items=len(names)):
                driver = mock.Mock()
                driver.window_handles = ["main", "popup"]
                fields = {f"_txtNAME{i}": mock.Mock() for i in range(2, 9)}
                def find_element(_by, field_id):
                    if field_id == "_btnSave":
                        return mock.Mock()
                    return fields[field_id]
                driver.find_element.side_effect = find_element
                wait = mock.Mock()
                wait.until.side_effect = lambda predicate: predicate(driver)
                with mock.patch.object(module, "save_duty_number_popup"), mock.patch.object(
                    module, "step_prepare_content", return_value=True
                ), mock.patch.object(module, "click_duty_external_setup_button_when_ready"), mock.patch.object(
                    module, "wait_for_popup_close_and_return_to_main"
                ), mock.patch.object(module, "accept_pending_alerts"), mock.patch.object(module.time, "sleep"):
                    module.step_config_popups(driver, wait, names, "2", "main")
                for index, name in enumerate(names, start=2):
                    fields[f"_txtNAME{index}"].clear.assert_called_once()
                    fields[f"_txtNAME{index}"].send_keys.assert_called_once_with(name)
                for index in range(2 + len(names), 7):
                    fields[f"_txtNAME{index}"].clear.assert_called_once()
                    fields[f"_txtNAME{index}"].send_keys.assert_not_called()
                fields["_txtNAME7"].clear.assert_called_once()
                fields["_txtNAME7"].send_keys.assert_called_once_with("救護")
                fields["_txtNAME8"].clear.assert_called_once()
                fields["_txtNAME8"].send_keys.assert_called_once_with("救護2車")

    def test_external_popup_missing_eighth_field_stops_before_any_write(self) -> None:
        from selenium.common.exceptions import NoSuchElementException
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        driver.window_handles = ["main", "popup"]
        fields = {f"_txtNAME{i}": mock.Mock() for i in range(2, 8)}
        def find_element(_by, field_id):
            if field_id not in fields:
                raise NoSuchElementException("設定勤務項目缺少第八欄")
            return fields[field_id]
        driver.find_element.side_effect = find_element
        wait = mock.Mock()
        wait.until.side_effect = lambda predicate: predicate(driver)
        with mock.patch.object(module, "save_duty_number_popup"), mock.patch.object(
            module, "step_prepare_content", return_value=True
        ), mock.patch.object(module, "click_duty_external_setup_button_when_ready"), mock.patch.object(
            module, "wait_for_popup_close_and_return_to_main"
        ), mock.patch.object(module, "accept_pending_alerts"), mock.patch.object(module.time, "sleep"):
            with self.assertRaisesRegex(RuntimeError, "外勤設定儲存失敗"):
                module.step_config_popups(driver, wait, ["勤務一"] * 5, "2", "main")
        for field in fields.values():
            field.clear.assert_not_called()
            field.send_keys.assert_not_called()

    def test_duty_sheet_excess_external_items_stop_before_website_changes(self) -> None:
        import openpyxl
        module = legacy_duty_sheet_module()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "3號"
        sheet.cell(5, 3, "值班")
        for column in range(4, 10):
            sheet.cell(5, column, f"外勤{column - 3}")
        for column, label in ((11, "休息"), (12, "備勤緊急救護"), (14, "備勤救災"), (16, "指揮官")):
            sheet.cell(5, column, label)
        sheet.cell(10, 16, "2")
        error = mock.Mock()
        with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
            module, "trainee_numbers_from_workbook", return_value=set()
        ), mock.patch.object(module, "expected_on_duty_numbers_from_daily_sheet", return_value=["2"]), mock.patch.object(
            module, "validate_daily_sheet_assignments", return_value=[]
        ), mock.patch.object(module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")) as build_driver, mock.patch.object(
            module, "capture_duty_sheet_images", return_value=[]
        ):
            result = module.start_automation("test-user", "test-password", "1151003", "test.xlsm", {},
                error_callback=error, show_dialogs=False)
        self.assertFalse(result)
        self.assertIn("外勤項目超過 5 項", error.call_args.args[0])
        build_driver.assert_not_called()

    def test_duty_sheet_matches_main_rest_header_and_keeps_extra_staff_on_standby(self) -> None:
        import openpyxl
        module = legacy_duty_sheet_module()
        for rest_label, extra_label in (("休息", "停休休息班"), ("休\n息", "停休休息班"), ("休息", "休息")):
            with self.subTest(rest=rest_label, extra=extra_label):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "23號"
                headers = {3: "值班", 4: "領取氣瓶", 5: "外出採買", 6: "公所進駐", 7: rest_label,
                    8: "在隊訓練", 9: "備勤緊急救護", 13: "備勤救災"}
                for column, label in headers.items():
                    sheet.cell(5, column, label)
                sheet.cell(6, 29, extra_label)
                sheet.cell(6, 32, "指揮官")
                for column, value in {2: "08-09", 3: "8", 4: "17", 5: "18", 6: "19", 7: "30",
                    8: "31", 9: "6,7", 14: "10,11,12", 29: "99", 32: "3"}.items():
                    sheet.cell(10, column, value)
                driver = mock.Mock()
                driver.execute_script.return_value = False
                error = mock.Mock()
                def use_mock_browser(_open_browser, verify, **_kwargs):
                    verify(driver)
                    return driver
                with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
                    module, "trainee_numbers_from_workbook", return_value=set()
                ), mock.patch.object(module, "expected_on_duty_numbers_from_daily_sheet", return_value=["3", "6", "7", "8", "10", "11", "12", "17", "18", "19", "30", "31", "99"]), mock.patch.object(
                    module, "validate_daily_sheet_assignments", return_value=[]
                ), mock.patch.object(module, "ThreadPoolExecutor"), mock.patch.object(
                    module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")
                ) as build_driver, mock.patch.object(
                    module, "retry_duty_number_popup_preflight", side_effect=use_mock_browser
                ), mock.patch.object(module, "preflight_duty_number_popup", return_value="main"), mock.patch.object(
                    module, "step_config_popups"
                ) as configure, mock.patch.object(
                    module, "wait_for_duty_result_grid", return_value=True
                ), mock.patch.object(module, "step_batch_fill_duty") as fill_duty, mock.patch.object(
                    module, "step_navigate_to_task_table"
                ), mock.patch.object(module, "super_js_execute", return_value=True), mock.patch.object(
                    module, "step_fill_mission_cells", side_effect=RuntimeError("offline-stop-before-save")
                ) as fill_mission, mock.patch.object(module.time, "sleep"), mock.patch.object(module, "log_status"):
                    result = module.start_automation("test-user", "test-password", "1151023", "test.xlsm", {},
                        error_callback=error, show_dialogs=False)
                self.assertFalse(result)
                error.assert_called_once_with("offline-stop-before-save")
                self.assertEqual(configure.call_args.args[2], ["領取氣瓶", "外出採買", "公所進駐"])
                duty_map = fill_duty.call_args.args[1]
                self.assertEqual({key: duty_map[key] for key in ("_pln_8_2", "_pln_8_3", "_pln_8_4")},
                    {"_pln_8_2": "17", "_pln_8_3": "18", "_pln_8_4": "19"})
                self.assertEqual(duty_map["_pln_8_5"], "6,7")
                self.assertEqual(duty_map["_pln_8_6"], "10,11")
                self.assertEqual(duty_map["_pln_8_7"], "12,99,3")
                self.assertEqual(duty_map["_pln_8_8"], "30")
                mission_map = fill_mission.call_args.args[1]
                fire_members = set((mission_map["_pln_8_1"] + "," + mission_map["_pln_8_2"]).split(","))
                self.assertEqual(fire_members, {"10", "11", "12", "99", "3"})
                build_driver.assert_not_called()

    def test_duty_sheet_only_submits_excel_normal_on_duty_staff(self) -> None:
        module = legacy_duty_sheet_module()
        normal_staff = ["2", "8", "9", "10", "11", "12", "16", "20", "24", "28", "30", "31"]
        off_duty_staff = {"1", "3", "7", "18", "19", "21", "22"}
        for source in ("roster", "daily_summary"):
            with self.subTest(source=source):
                workbook = module.openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "23號"
                for column, label in {3: "值班", 4: "外勤", 5: "休息", 6: "備勤緊急救護", 10: "備勤救災"}.items():
                    sheet.cell(5, column, label)
                sheet.cell(6, 15, "停休休息班")
                sheet.cell(6, 16, "指揮官")
                sheet.cell(22, 20, "備勤")
                sheet.cell(22, 21, ",".join(normal_staff + (["1"] if source == "roster" else [])))
                for row in range(10, 34):
                    hour = (row - 2) % 24
                    for column, value in {2: f"{hour:02d}-{(hour + 1) % 24:02d}", 3: "20,18", 4: "31,19",
                        5: "30,21", 6: "16,7", 7: "28,24", 11: "8,9,10,11,12", 12: "29", 15: "3,22", 16: "1,2"}.items():
                        sheet.cell(row, column, value)
                if source == "roster":
                    roster = workbook.create_sheet("輪休基準表")
                    params = workbook.create_sheet("班別參數")
                    params.cell(1, 1, "停")
                    params.cell(1, 2, "停休")
                    roster.cell(5, 3, "日期")
                    roster.cell(6, 3, 23)
                    for column, number in enumerate(normal_staff + sorted(off_duty_staff) + ["29"], start=5):
                        roster.cell(4, column, number)
                        roster.cell(5, column, f"人員{number}")
                        if number in off_duty_staff:
                            roster.cell(6, column, "停")
                driver = mock.Mock()
                driver.execute_script.return_value = False
                error = mock.Mock()
                def use_mock_browser(_open_browser, verify, **_kwargs):
                    verify(driver)
                    return driver
                with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
                    module, "trainee_numbers_from_workbook", return_value={"29"}
                ), mock.patch.object(module, "ThreadPoolExecutor"), mock.patch.object(
                    module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")
                ) as build_driver, mock.patch.object(
                    module, "retry_duty_number_popup_preflight", side_effect=use_mock_browser
                ), mock.patch.object(module, "preflight_duty_number_popup", return_value="main") as preflight, mock.patch.object(
                    module, "step_config_popups"
                ) as configure, mock.patch.object(module, "wait_for_duty_result_grid", return_value=True), mock.patch.object(
                    module, "step_batch_fill_duty"
                ) as fill_duty, mock.patch.object(module, "step_navigate_to_task_table"), mock.patch.object(
                    module, "super_js_execute", return_value=True
                ), mock.patch.object(module, "step_fill_mission_cells", side_effect=RuntimeError("offline-stop-before-save")) as fill_mission, mock.patch.object(
                    module.time, "sleep"
                ), mock.patch.object(module, "log_status"):
                    result = module.start_automation("test-user", "test-password", "1151023", "test.xlsm", {},
                        error_callback=error, show_dialogs=False)
                self.assertFalse(result)
                error.assert_called_once_with("offline-stop-before-save")
                self.assertEqual(configure.call_args.args[3], 2)
                self.assertEqual(preflight.call_args.args[2], normal_staff)
                duty_map = fill_duty.call_args.args[1]
                mission_map = fill_mission.call_args.args[1]
                for hour in range(24):
                    self.assertEqual([duty_map[f"_pln_{hour}_{column}"] for column in range(1, 7)],
                        ["20", "31", "16,28,24", "8,9", "10,11,12,2", "30"])
                    self.assertEqual(mission_map[f"_pln_{hour}_6"], "16,28,24")
                for payload in (duty_map, mission_map):
                    assigned = {number for key, value in payload.items() if key.startswith("_pln_")
                        for number in value.split(",") if number}
                    self.assertTrue(assigned.issubset(set(normal_staff)))
                    self.assertTrue(assigned.isdisjoint(off_duty_staff | {"29"}))
                self.assertEqual(sheet.cell(10, 15).value, "3,22")
                build_driver.assert_not_called()

    def test_duty_sheet_off_duty_commander_stops_before_browser_or_capture(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "23號"
        for column, label in {3: "值班", 4: "休息", 5: "備勤緊急救護", 9: "備勤救災"}.items():
            sheet.cell(5, column, label)
        sheet.cell(6, 13, "指揮官")
        sheet.cell(22, 16, "備勤")
        sheet.cell(22, 17, "20")
        for row in range(10, 34):
            hour = (row - 2) % 24
            sheet.cell(row, 2, f"{hour:02d}-{(hour + 1) % 24:02d}")
            sheet.cell(row, 3, "20")
            sheet.cell(row, 13, "1")
        error = mock.Mock()
        with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
            module, "ThreadPoolExecutor", side_effect=AssertionError("測試不得啟動截圖")
        ) as executor, mock.patch.object(module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")) as build_driver, mock.patch.object(
            module, "log_status"
        ):
            result = module.start_automation("test-user", "test-password", "1151023", "test.xlsm", {},
                error_callback=error, show_dialogs=False)
        self.assertFalse(result)
        self.assertIn("【指揮官未排】", error.call_args.args[0])
        executor.assert_not_called()
        build_driver.assert_not_called()

    def test_duty_sheet_commander_moved_to_ambulance_stops_before_browser_or_capture(self) -> None:
        module = legacy_duty_sheet_module()
        for commander in (None, " \t\n\u3000", 0, "0"):
            with self.subTest(commander=commander):
                workbook = module.openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "24號"
                for column, label in {3: "值班", 4: "休息", 5: "備勤緊急救護", 9: "備勤救災"}.items():
                    sheet.cell(5, column, label)
                sheet.cell(6, 13, "指揮官")
                for row in range(10, 34):
                    hour = (row - 2) % 24
                    sheet.cell(row, 2, f"{hour:02d}-{(hour + 1) % 24:02d}")
                    sheet.cell(row, 3, "20")
                    sheet.cell(row, 5, "16")
                    sheet.cell(row, 6, "28")
                    sheet.cell(row, 13, "1")
                for row in (12, 13):
                    sheet.cell(row, 6, "28,1")
                    sheet.cell(row, 13).value = commander
                error = mock.Mock()
                with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
                    module, "trainee_numbers_from_workbook", return_value=set()
                ), mock.patch.object(module, "expected_on_duty_numbers_from_daily_sheet", return_value=["20", "16", "28", "1"]), mock.patch.object(
                    module, "ThreadPoolExecutor", side_effect=AssertionError("測試不得啟動截圖")
                ) as executor, mock.patch.object(module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")) as build_driver, mock.patch.object(
                    module, "capture_duty_sheet_images", side_effect=AssertionError("測試不得啟動截圖")
                ) as capture, mock.patch.object(module, "log_status"):
                    result = module.start_automation("test-user", "test-password", "1151024", "test.xlsm", {},
                        error_callback=error, show_dialogs=False)
                self.assertFalse(result)
                self.assertIn("10-11：【指揮官未排】", error.call_args.args[0])
                self.assertIn("11-12：【指揮官未排】", error.call_args.args[0])
                self.assertNotIn("【漏排】", error.call_args.args[0])
                executor.assert_not_called()
                capture.assert_not_called()
                build_driver.assert_not_called()

    def test_duty_sheet_keeps_three_ambulance_members_in_duty_and_mission_payloads(self) -> None:
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "24號"
        for column, label in {3: "值班", 4: "休息", 5: "備勤緊急救護", 9: "備勤救災"}.items():
            sheet.cell(5, column, label)
        sheet.cell(6, 13, "指揮官")
        for row in range(10, 34):
            hour = (row - 2) % 24
            for column, value in {2: f"{hour:02d}-{(hour + 1) % 24:02d}", 3: "20", 5: "16", 6: "28,1",
                10: "8,9,10,11,12", 13: "2"}.items():
                sheet.cell(row, column, value)
        driver = mock.Mock()
        driver.execute_script.return_value = False
        error = mock.Mock()
        def use_mock_browser(_open_browser, verify, **_kwargs):
            verify(driver)
            return driver
        with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
            module, "trainee_numbers_from_workbook", return_value=set()
        ), mock.patch.object(module, "expected_on_duty_numbers_from_daily_sheet", return_value=["20", "16", "28", "1", "8", "9", "10", "11", "12", "2"]), mock.patch.object(
            module, "ThreadPoolExecutor"
        ), mock.patch.object(module, "build_driver", side_effect=AssertionError("測試不得啟動瀏覽器")) as build_driver, mock.patch.object(
            module, "retry_duty_number_popup_preflight", side_effect=use_mock_browser
        ), mock.patch.object(module, "preflight_duty_number_popup", return_value="main"), mock.patch.object(
            module, "step_config_popups"
        ), mock.patch.object(module, "wait_for_duty_result_grid", return_value=True), mock.patch.object(
            module, "step_batch_fill_duty"
        ) as fill_duty, mock.patch.object(module, "step_navigate_to_task_table"), mock.patch.object(
            module, "super_js_execute", return_value=True
        ), mock.patch.object(module, "step_fill_mission_cells", side_effect=RuntimeError("offline-stop-before-save")) as fill_mission, mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(module, "log_status"):
            result = module.start_automation("test-user", "test-password", "1151024", "test.xlsm", {},
                error_callback=error, show_dialogs=False)
        self.assertFalse(result)
        error.assert_called_once_with("offline-stop-before-save")
        duty_map = fill_duty.call_args.args[1]
        mission_map = fill_mission.call_args.args[1]
        for hour in range(24):
            self.assertEqual(duty_map[f"_pln_{hour}_2"], "16,28,1")
            self.assertEqual(mission_map[f"_pln_{hour}_6"], "16,28,1")
            self.assertTrue(set(mission_map[f"_pln_{hour}_7"].split(",")).isdisjoint({"16", "28", "1"}))
        build_driver.assert_not_called()

    def test_duty_result_grid_missing_rescue_column_is_rejected(self) -> None:
        from selenium.common.exceptions import TimeoutException
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        class WaitOnce:
            def until(self, predicate):
                if not predicate(driver):
                    raise TimeoutException("勤務表少一欄")
                return True
        with mock.patch.object(module, "super_js_execute", side_effect=lambda _driver, element_id, action:
            element_id != "_pln_8_8"):
            with self.assertRaisesRegex(RuntimeError, "欄位不完整"):
                module.wait_for_duty_result_grid(driver, WaitOnce(), expected_columns=10)
        with mock.patch.object(module, "super_js_execute", return_value=True) as exists:
            self.assertTrue(module.wait_for_duty_result_grid(driver, WaitOnce(), expected_columns=10))
            self.assertEqual([call.args[1] for call in exists.call_args_list],
                [f"_pln_8_{index}" for index in range(1, 11)])

    def test_visible_driver_is_positioned_at_top_right_without_explicit_position(self) -> None:
        module = duty_rehearsal_module()
        driver = mock.Mock()

        with mock.patch.object(module, "chrome_start_attempts", return_value=1), mock.patch.object(
            module, "duty_browser_profile_dir", return_value=Path("C:/Temp/duty_gui_visible")
        ), mock.patch.object(
            module, "prune_stale_duty_browser_profiles", return_value=0
        ), mock.patch.object(
            module, "create_webdriver_chrome_with_timeout", return_value=driver
        ), mock.patch.object(module, "position_duty_browser_at_top_right") as position_browser:
            result = module.build_driver(headless=False)

        self.assertIs(result, driver)
        position_browser.assert_called_once_with(driver)

    def test_stale_profile_cleanup_is_bounded_and_skips_active_profiles(self) -> None:
        module = duty_rehearsal_module()
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            stale_one = root / "duty_gui_stale_one"
            stale_two = root / "duty_gui_stale_two"
            active = root / "duty_gui_active"
            unrelated = root / "other_profile"
            for profile in (stale_one, stale_two, active, unrelated):
                profile.mkdir()
                os.utime(profile, (1, 1))

            removed = module.prune_stale_duty_browser_profiles(
                root=root,
                active_profiles={active},
                now=1_000,
                minimum_age_seconds=60,
                maximum_profiles=1,
            )

            self.assertEqual(removed, 1)
            self.assertFalse(stale_one.exists())
            self.assertTrue(stale_two.exists())
            self.assertTrue(active.exists())
            self.assertTrue(unrelated.exists())

    def test_browser_profile_process_checks_never_open_a_console(self) -> None:
        module = duty_rehearsal_module()
        process_result = mock.Mock(returncode=0, stdout="")

        with mock.patch.object(module.subprocess, "run", return_value=process_result) as run_process:
            self.assertEqual(
                module._active_duty_browser_profiles(Path("C:/owned"), [Path("C:/owned/duty_gui_one")]),
                set(),
            )

        self.assertEqual(
            run_process.call_args.kwargs["creationflags"],
            getattr(module.subprocess, "CREATE_NO_WINDOW", 0),
        )

    def test_browser_profile_cleanup_never_opens_a_console(self) -> None:
        module = duty_rehearsal_module()
        process_result = mock.Mock(returncode=0, stdout="")

        with mock.patch.object(module, "_is_owned_duty_browser_profile", return_value=True), mock.patch.object(
            module.subprocess, "run", return_value=process_result
        ) as run_process, mock.patch.object(module.shutil, "rmtree"):
            module.cleanup_duty_browser_profile(Path("C:/owned/duty_gui_one"), terminate_processes=True)

        self.assertEqual(
            run_process.call_args.kwargs["creationflags"],
            getattr(module.subprocess, "CREATE_NO_WINDOW", 0),
        )

    def test_chromedriver_starts_without_a_console(self) -> None:
        module = duty_rehearsal_module()
        driver = mock.Mock()

        with mock.patch.object(module.webdriver, "Chrome", return_value=driver) as launch_chrome:
            self.assertIs(module.create_webdriver_chrome_with_timeout(module.Options()), driver)

        service = launch_chrome.call_args.kwargs["service"]
        self.assertEqual(
            service.creation_flags,
            getattr(module.subprocess, "CREATE_NO_WINDOW", 0),
        )
        self.assertEqual(service.log_output, module.subprocess.DEVNULL)

    def test_concurrent_chrome_start_requests_are_serialized(self) -> None:
        module = duty_rehearsal_module()
        first_started = threading.Event()
        second_thread_started = threading.Event()
        release_first = threading.Event()
        state_lock = threading.Lock()
        active = 0
        peak_active = 0
        call_count = 0

        def fake_chrome(**_kwargs: object) -> mock.Mock:
            nonlocal active, peak_active, call_count
            with state_lock:
                call_count += 1
                call_number = call_count
                active += 1
                peak_active = max(peak_active, active)
            if call_number == 1:
                first_started.set()
                release_first.wait(2)
            with state_lock:
                active -= 1
            return mock.Mock()

        first_result: list[object] = []
        second_result: list[object] = []

        def start_first() -> None:
            first_result.append(module.create_webdriver_chrome_with_timeout(module.Options()))

        def start_second() -> None:
            second_thread_started.set()
            second_result.append(module.create_webdriver_chrome_with_timeout(module.Options()))

        with mock.patch.object(module.webdriver, "Chrome", side_effect=fake_chrome):
            first_thread = threading.Thread(target=start_first)
            second_thread = threading.Thread(target=start_second)
            first_thread.start()
            self.assertTrue(first_started.wait(1))
            second_thread.start()
            self.assertTrue(second_thread_started.wait(1))
            self.assertFalse(release_first.wait(0.2))
            with state_lock:
                self.assertEqual(call_count, 1)
                self.assertEqual(peak_active, 1)
            release_first.set()
            first_thread.join(2)
            second_thread.join(2)

        self.assertFalse(first_thread.is_alive())
        self.assertFalse(second_thread.is_alive())
        self.assertEqual(len(first_result), 1)
        self.assertEqual(len(second_result), 1)
        self.assertEqual(call_count, 2)
        self.assertEqual(peak_active, 1)

    def test_pruner_preserves_stale_profile_owned_by_inflight_constructor(self) -> None:
        module = duty_rehearsal_module()
        release = threading.Event()
        driver = mock.Mock()

        def blocked_chrome(**_kwargs):
            release.wait(2)
            return driver

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            profile = module.duty_browser_profile_dir(root=root)
            options = module.Options()
            options.add_argument("--user-data-dir=" + str(profile))
            with mock.patch.object(module, "chrome_start_timeout_seconds", return_value=0.02), mock.patch.object(
                module.webdriver, "Chrome", side_effect=blocked_chrome
            ), mock.patch.object(module, "_write_duty_browser_startup_diagnostic"):
                try:
                    with self.assertRaises(TimeoutError):
                        module.create_webdriver_chrome_with_timeout(options)
                    removed = module.prune_stale_duty_browser_profiles(
                        root=root, now=module.time.time() + 601, minimum_age_seconds=600, active_profiles=set(),
                    )
                    self.assertEqual(removed, 0)
                    self.assertTrue(profile.exists())
                finally:
                    release.set()
                    self.assertTrue(module._CHROME_START_LOCK.acquire(timeout=1))
                    module._CHROME_START_LOCK.release()
                driver.quit.assert_called_once_with()
                self.assertFalse(profile.exists())

    def test_chrome_start_timeout_keeps_one_owner_and_queue_creates_no_thread(self) -> None:
        module = duty_rehearsal_module()
        release = threading.Event()
        completed = threading.Event()
        driver = mock.Mock()

        def blocked_chrome(**_kwargs):
            release.wait(2)
            return driver

        driver.quit.side_effect = completed.set
        with mock.patch.object(module, "chrome_start_timeout_seconds", return_value=0.04), mock.patch.object(
            module.webdriver, "Chrome", side_effect=blocked_chrome
        ) as constructor, mock.patch.object(module, "_write_duty_browser_startup_diagnostic") as diagnostic:
            try:
                with self.assertRaises(TimeoutError) as first:
                    module.create_webdriver_chrome_with_timeout(module.Options())
                with mock.patch.object(module.threading, "Thread", wraps=threading.Thread) as threads:
                    with self.assertRaises(TimeoutError) as queued:
                        module.create_webdriver_chrome_with_timeout(module.Options())
                self.assertEqual(threads.call_count, 0)
                self.assertEqual(constructor.call_count, 1)
                self.assertTrue(first.exception.constructor_started)
                self.assertEqual(queued.exception.startup_stage, "queue_wait")
                self.assertFalse(queued.exception.constructor_started)
                self.assertTrue(queued.exception.in_progress)
                self.assertEqual(diagnostic.call_args.kwargs["actual_attempts"], 0)
            finally:
                release.set()
                self.assertTrue(completed.wait(1))
            driver.quit.assert_called_once_with()
            self.assertEqual(driver.method_calls, [mock.call.quit()])
            with mock.patch.object(module.webdriver, "Chrome", return_value=mock.sentinel.next_driver):
                self.assertIs(module.create_webdriver_chrome_with_timeout(module.Options()), mock.sentinel.next_driver)

    def test_chrome_resolver_timeout_is_identified_before_driver_service_start(self) -> None:
        from selenium.webdriver.common.driver_finder import DriverFinder
        from selenium.webdriver.common.service import Service

        module = duty_rehearsal_module()
        release = threading.Event()
        finished = threading.Event()

        def blocked_resolution(_finder):
            try:
                release.wait(2)
                raise RuntimeError("private resolver detail must not escape")
            finally:
                finished.set()

        options = module.Options()
        options.add_argument("--user-data-dir=C:/private/notice_browser_profiles/duty_gui_test")
        with mock.patch.object(module, "chrome_start_timeout_seconds", return_value=0.04), mock.patch.object(
            DriverFinder, "_binary_paths", blocked_resolution
        ), mock.patch.object(Service, "start") as service_start, mock.patch.object(
            module, "_write_duty_browser_startup_diagnostic"
        ) as diagnostic, mock.patch.object(module, "cleanup_duty_browser_profile"):
            try:
                with self.assertRaises(TimeoutError) as raised:
                    module.create_webdriver_chrome_with_timeout(options)
                self.assertEqual(raised.exception.startup_stage, "driver_resolution")
                self.assertTrue(raised.exception.constructor_started)
                service_start.assert_not_called()
                self.assertEqual(diagnostic.call_args.kwargs["context"], "notice")
                self.assertNotIn("private", str(diagnostic.call_args))
            finally:
                release.set()
                self.assertTrue(finished.wait(1))
                self.assertTrue(module._CHROME_START_LOCK.acquire(timeout=1))
                module._CHROME_START_LOCK.release()

    def test_chrome_constructor_is_not_started_after_request_already_timed_out(self) -> None:
        module = duty_rehearsal_module()
        release = threading.Event()

        def delayed_service(**_kwargs):
            release.wait(2)
            return mock.Mock()

        with mock.patch.object(module, "chrome_start_timeout_seconds", return_value=0.04), mock.patch.object(
            module, "ChromeService", side_effect=delayed_service
        ), mock.patch.object(module.webdriver, "Chrome") as constructor, mock.patch.object(
            module, "_write_duty_browser_startup_diagnostic"
        ):
            try:
                with self.assertRaises(TimeoutError) as raised:
                    module.create_webdriver_chrome_with_timeout(module.Options())
                self.assertFalse(raised.exception.constructor_started)
            finally:
                release.set()
                self.assertTrue(module._CHROME_START_LOCK.acquire(timeout=1))
                module._CHROME_START_LOCK.release()
            constructor.assert_not_called()

    def test_build_driver_does_not_retry_or_remove_profile_while_constructor_runs(self) -> None:
        module = duty_rehearsal_module()
        error = module.ChromeStartupTimeout("driver_resolution", constructor_started=True, in_progress=True)
        with mock.patch.object(module, "chrome_start_attempts", return_value=2), mock.patch.object(
            module, "prune_stale_duty_browser_profiles", return_value=0
        ), mock.patch.object(module, "duty_browser_profile_dir", return_value=Path("C:/owned/duty_gui_one")), mock.patch.object(
            module, "create_webdriver_chrome_with_timeout", side_effect=error
        ) as constructor, mock.patch.object(module, "cleanup_duty_browser_startup_failure") as cleanup, mock.patch.object(
            module, "_write_duty_browser_startup_diagnostic"
        ):
            with self.assertRaises(module.DutyBrowserStartupError) as raised:
                module.build_driver(headless=True, write_diagnostics=False)
        self.assertEqual(constructor.call_count, 1)
        cleanup.assert_not_called()
        self.assertEqual(raised.exception.startup_stage, "driver_resolution")
        self.assertEqual(raised.exception.attempts, 1)
        self.assertEqual(raised.exception.startup_message, module.CHROME_STARTUP_STAGE_MESSAGES["driver_resolution"])

    def test_chrome_service_and_session_timeouts_keep_safe_distinct_stages(self) -> None:
        from selenium.webdriver.common.driver_finder import DriverFinder
        from selenium.webdriver.common.service import Service
        from selenium.webdriver.remote.webdriver import WebDriver

        module = duty_rehearsal_module()
        for stage in ("driver_service_start", "session_create"):
            with self.subTest(stage=stage), ExitStack() as stack:
                release = threading.Event()
                driver_closed = threading.Event()

                def block(*_args, **_kwargs):
                    release.wait(2)
                    if stage == "driver_service_start":
                        raise RuntimeError("private startup details")
                    return {"value": {"sessionId": "test-only", "capabilities": {}}}

                stack.enter_context(mock.patch.object(module, "chrome_start_timeout_seconds", return_value=0.04))
                stack.enter_context(mock.patch.object(DriverFinder, "_binary_paths", return_value={
                    "driver_path": "C:/test-only/driver", "browser_path": ""}))
                stack.enter_context(mock.patch.object(module, "_write_duty_browser_startup_diagnostic"))
                if stage == "driver_service_start":
                    stack.enter_context(mock.patch.object(Service, "_start_process", block))
                else:
                    stack.enter_context(mock.patch.object(Service, "start"))
                    stack.enter_context(mock.patch.object(WebDriver, "execute", block))
                    stack.enter_context(mock.patch.object(module.webdriver.Chrome, "quit", side_effect=driver_closed.set))
                try:
                    with self.assertRaises(TimeoutError) as raised:
                        module.create_webdriver_chrome_with_timeout(module.Options())
                    self.assertEqual(raised.exception.startup_stage, stage)
                    self.assertTrue(raised.exception.in_progress)
                finally:
                    release.set()
                    self.assertTrue(module._CHROME_START_LOCK.acquire(timeout=1))
                    module._CHROME_START_LOCK.release()
                if stage == "session_create":
                    self.assertTrue(driver_closed.is_set())

    def test_startup_diagnostic_contains_only_safe_stage_and_context_metadata(self) -> None:
        module = duty_rehearsal_module()
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(module, "__file__", str(Path(directory) / "source.py")):
            module._write_duty_browser_startup_diagnostic(
                "startup_timeout", category="startup_timeout", attempts=0, profiles_pruned=0,
                startup_stage="queue_wait", context="C:/private/token", constructor_started=False,
                actual_attempts=0, in_progress=True, blocking_stage="driver_resolution",
            )
            files = list((Path(directory) / "runtime_outputs/browser").glob("browser_startup_*.jsonl"))
            self.assertEqual(len(files), 1)
            payload = json.loads(files[0].read_text(encoding="utf-8"))
        self.assertEqual(payload["context"], "duty")
        self.assertEqual(payload["startup_stage"], "queue_wait")
        self.assertEqual(payload["blocking_stage"], "driver_resolution")
        self.assertEqual(payload["actual_attempts"], 0)
        self.assertNotIn("private", json.dumps(payload))

    def test_build_driver_queue_timeout_has_zero_attempts_and_only_cleans_unused_profile(self) -> None:
        module = duty_rehearsal_module()
        error = module.ChromeStartupTimeout("queue_wait", constructor_started=False, in_progress=True,
                                           blocking_stage="driver_resolution")
        with mock.patch.object(module, "prune_stale_duty_browser_profiles", return_value=0), mock.patch.object(
            module, "duty_browser_profile_dir", return_value=Path("C:/owned/duty_gui_one")
        ), mock.patch.object(module, "create_webdriver_chrome_with_timeout", side_effect=error) as constructor, mock.patch.object(
            module, "cleanup_duty_browser_profile"
        ) as cleanup:
            with self.assertRaises(module.DutyBrowserStartupError) as raised:
                module.build_driver(headless=True)
        self.assertEqual(constructor.call_count, 1)
        self.assertEqual(raised.exception.attempts, 0)
        self.assertEqual(raised.exception.blocking_stage, "driver_resolution")
        cleanup.assert_called_once_with(Path("C:/owned/duty_gui_one"), root=None, terminate_processes=False)

    def test_legacy_daily_vehicle_launcher_never_opens_a_console(self) -> None:
        source = legacy_tk_path("daily_vehicle_automation.py").read_text(encoding="utf-8-sig")

        self.assertIn("process = subprocess.Popen(", source)
        self.assertIn('creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)', source)

    def test_quit_driver_cleans_only_the_attached_private_profile(self) -> None:
        module = duty_rehearsal_module()
        driver = mock.Mock()
        driver.service = mock.Mock()
        setattr(driver, "_sinposmart_duty_browser_profile", "C:/owned/duty_gui_profile")

        with mock.patch.object(module, "cleanup_duty_browser_profile") as cleanup:
            module.quit_driver(driver)

        driver.quit.assert_called_once_with()
        driver.service.stop.assert_called_once_with()
        cleanup.assert_called_once_with(Path("C:/owned/duty_gui_profile"), terminate_processes=False)

    def test_browser_startup_failure_writes_safe_diagnostic(self) -> None:
        module = duty_rehearsal_module()
        profiles = [Path("C:/Temp/duty_gui_first"), Path("C:/Temp/duty_gui_second")]
        diagnostics: list[dict[str, object]] = []

        def record(event: str, **fields: object) -> None:
            diagnostics.append({"event": event, **fields})

        with mock.patch.object(module, "chrome_start_attempts", return_value=2), mock.patch.object(
            module, "prune_stale_duty_browser_profiles", return_value=3
        ), mock.patch.object(module, "duty_browser_profile_dir", side_effect=profiles), mock.patch.object(
            module,
            "create_webdriver_chrome_with_timeout",
            side_effect=TimeoutError("startup timed out"),
        ), mock.patch.object(module, "cleanup_duty_browser_startup_failure"), mock.patch.object(
            module, "_write_duty_browser_startup_diagnostic", side_effect=record
        ), mock.patch.object(module.time, "sleep"):
            with self.assertRaises(module.DutyBrowserStartupError) as raised:
                module.build_driver(headless=True)

        self.assertEqual(raised.exception.diagnostic_category, "startup_timeout")
        self.assertEqual(
            diagnostics,
            [
                {
                    "event": "startup_failed",
                    "category": "startup_timeout",
                    "attempts": 2,
                    "profiles_pruned": 3,
                }
            ],
        )

    def test_all_local_duty_tools_use_the_safe_chrome_startup(self) -> None:
        root = package_dir()
        duty_sheet_source = (root / "duty_sheet_legacy" / "sinposmart_1.py").read_text(encoding="utf-8-sig")
        vehicle_source = (root / "daily_vehicle_legacy" / "automation" / "ppe_selenium_daily.py").read_text(encoding="utf-8-sig")
        rest_source = (root / "rest_time_automation.py").read_text(encoding="utf-8-sig")
        schedule_source = (root / "app_core" / "schedule_capture_service.py").read_text(encoding="utf-8-sig")
        submission_source = (root / "app_core" / "duty_submission_service.py").read_text(encoding="utf-8-sig")
        login_source = (root / "app_core" / "login_verifier.py").read_text(encoding="utf-8-sig")

        self.assertIn("from duty_rehearsal import (", duty_sheet_source)
        self.assertIn("build_driver,", duty_sheet_source)
        self.assertIn("retain_browser_for_seconds,", duty_sheet_source)
        self.assertIn("quit_driver(driver)", duty_sheet_source)
        self.assertIn("retry_duty_browser_session_open", duty_sheet_source)
        self.assertIn("from duty_rehearsal import (", vehicle_source)
        self.assertIn("build_driver,", vehicle_source)
        self.assertIn("quit_driver(driver)", vehicle_source)
        self.assertIn("retry_duty_browser_session_open", vehicle_source)
        self.assertIn("return build_driver(", vehicle_source)
        self.assertIn('page_load_strategy="none"', vehicle_source)
        self.assertIn("return webdriver.Remote(command_executor=remote_url, options=options)", vehicle_source)
        self.assertIn("build_initialized_driver", rest_source)
        self.assertIn("build_initialized_driver", schedule_source)
        self.assertIn("build_initialized_driver", submission_source)
        self.assertIn("retry_duty_browser_session_open", login_source)

    def test_next_morning_0600_rest_includes_manual_return_at_0800(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150621",
            rows=[
                module.DutyRow("06-08", {"值班": ["28"], "休息": ["16"]}),
                module.DutyRow("08-10", {"值班": ["12"], "休息": []}),
                module.DutyRow("22-24", {"值班": ["28"], "休息": []}),
            ],
            summary={"在勤": ["16", "28"]},
        )
        yesterday = module.DutySheet(roc_date="1150620", summary={"在勤": ["28"]})
        tomorrow = module.DutySheet(roc_date="1150622", summary={"在勤": ["16", "28"]})

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150621"), [], tomorrow)
        rest_actions = [
            action
            for action in actions
            if action.kind == "entry_log" and action.target == "16" and action.fields.get("領用事由及地點") in ("休息", "休息返隊")
        ]

        self.assertEqual(
            [(action.time, action.actor, action.fields["出或入"], action.fields["領用事由及地點"], action.date_offset) for action in rest_actions],
            [("06:00", "28", "出", "休息", 1), ("08:00", "28", "入", "休息返隊", 1)],
        )

    def test_next_morning_0002_rest_creates_return_even_when_off_duty_tomorrow(self) -> None:
        module = duty_rehearsal_module()
        duty = "\u503c\u73ed"
        rest = "\u4f11\u606f"
        on_duty = "\u5728\u52e4"
        direction = "\u51fa\u6216\u5165"
        reason_key = "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede"
        today = module.DutySheet(
            roc_date="1150627",
            rows=[
                module.DutyRow("0-1", {duty: ["8"], rest: ["2"]}),
                module.DutyRow("1-2", {duty: ["8"], rest: ["2"]}),
                module.DutyRow("2-3", {duty: ["8"], rest: []}),
                module.DutyRow("22-24", {duty: ["8"], rest: []}),
            ],
            summary={on_duty: ["2", "8"]},
        )
        tomorrow = module.DutySheet(roc_date="1150628", rows=[], summary={})

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150627"), [], tomorrow)
        rest_actions = [
            action
            for action in actions
            if action.kind == "entry_log" and action.target == "2" and action.fields.get(reason_key) in ("休息", "休息返隊")
        ]

        self.assertEqual(
            [(action.time, action.actor, action.fields[direction], action.fields[reason_key], action.date_offset) for action in rest_actions],
            [("00:00", "8", "出", "休息", 1), ("02:00", "8", "入", "休息返隊", 1)],
        )

    def test_overnight_2201_rest_is_one_block_not_two(self) -> None:
        module = duty_rehearsal_module()
        duty = "\u503c\u73ed"
        rest = "\u4f11\u606f"
        on_duty = "\u5728\u52e4"
        direction = "\u51fa\u6216\u5165"
        reason_key = "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede"
        today = module.DutySheet(
            roc_date="1150715",
            rows=[
                module.DutyRow("00-01", {duty: ["8"], rest: ["23"]}),
                module.DutyRow("01-02", {duty: ["8"], rest: []}),
                module.DutyRow("22-24", {duty: ["8"], rest: ["23"]}),
            ],
            summary={on_duty: ["8", "23"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150715"), [], None)
        rest_actions = [
            action
            for action in actions
            if action.kind == "entry_log" and action.target == "23" and action.fields.get(reason_key) in (rest, "\u4f11\u606f\u8fd4\u968a")
        ]

        self.assertEqual(
            [(action.time, action.fields[direction], action.fields[reason_key], action.date_offset, action.source) for action in rest_actions],
            [
                ("22:00", "\u51fa", rest, 0, "\u4f11\u606f\u7c3d\u51fa"),
                ("01:00", "\u5165", "\u4f11\u606f\u8fd4\u968a", 1, "\u4f11\u606f\u7d50\u675f"),
            ],
        )

    def test_schedule_load_replaces_stale_overnight_rest_actions(self) -> None:
        module = duty_gui_module()
        duty = "\u503c\u73ed"
        rest = "\u4f11\u606f"
        on_duty = "\u5728\u52e4"
        direction = "\u51fa\u6216\u5165"
        reason_key = "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede"
        data = {
            "target_date": "1150715",
            "today": {
                "rows": [
                    {"slot": "00-01", "columns": {duty: ["8"], rest: ["23"]}},
                    {"slot": "01-02", "columns": {duty: ["8"], rest: []}},
                    {"slot": "22-24", "columns": {duty: ["8"], rest: ["23"]}},
                ],
                "summary": {on_duty: ["8", "23"]},
            },
            "yesterday": {},
            "cases": [],
            "yesterday_cases": [],
            "actions": [
                {
                    "kind": "entry_log",
                    "time": "00:00",
                    "actor": "8",
                    "target": "23",
                    "fields": {direction: "\u5165", reason_key: "\u4f11\u606f\u8fd4\u968a"},
                    "source": "\u4f11\u606f\u7d50\u675f",
                    "duplicate_key": "entry:1150715:0:in:23:\u4f11\u606f\u8fd4\u968a",
                    "date_offset": 1,
                },
                {
                    "kind": "entry_log",
                    "time": "00:00",
                    "actor": "8",
                    "target": "23",
                    "fields": {direction: "\u51fa", reason_key: rest},
                    "source": "\u4f11\u606f\u7c3d\u51fa",
                    "duplicate_key": "entry:1150715:0:out:23:\u4f11\u606f",
                    "date_offset": 1,
                },
            ],
        }

        module.DutyGui.ensure_schedule_actions(object(), data)
        rest_actions = [
            action
            for action in data["actions"]
            if action.get("kind") == "entry_log"
            and action.get("target") == "23"
            and action.get("fields", {}).get(reason_key) in (rest, "\u4f11\u606f\u8fd4\u968a")
        ]

        self.assertEqual(
            [(action["time"], action["fields"][direction], action["fields"][reason_key], action["date_offset"]) for action in rest_actions],
            [("22:00", "\u51fa", rest, 0), ("01:00", "\u5165", "\u4f11\u606f\u8fd4\u968a", 1)],
        )

    def test_next_morning_0408_rest_for_off_duty_tomorrow_creates_rest_checkout(self) -> None:
        module = duty_rehearsal_module()
        duty = "\u503c\u73ed"
        rest = "\u4f11\u606f"
        on_duty = "\u5728\u52e4"
        direction = "\u51fa\u6216\u5165"
        reason_key = "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede"
        today = module.DutySheet(
            roc_date="1150627",
            rows=[
                module.DutyRow("04-06", {duty: ["4"], rest: ["4"]}),
                module.DutyRow("06-08", {duty: ["4"], rest: ["4"]}),
                module.DutyRow("08-10", {duty: ["12"], rest: []}),
                module.DutyRow("22-24", {duty: ["8"], rest: []}),
            ],
            summary={on_duty: ["4", "8", "12"]},
        )
        tomorrow = module.DutySheet(roc_date="1150628", rows=[module.DutyRow("08-10", {duty: ["12"]})], summary={on_duty: ["12"]})

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150627"), [], tomorrow)
        rest_actions = [
            action
            for action in actions
            if action.kind == "entry_log" and action.target == "4" and action.fields.get(reason_key) in ("休息", "休息返隊", "休息後退勤")
        ]

        self.assertEqual(
            [(action.time, action.actor, action.fields[direction], action.fields[reason_key], action.date_offset) for action in rest_actions],
            [("04:00", "8", "出", "休息後退勤", 1)],
        )
        self.assertEqual(
            [
                (action.time, action.fields[direction], action.fields[reason_key], action.source, action.date_offset)
                for action in actions
                if action.kind == "entry_log" and action.target == "4" and action.date_offset == 1
            ],
            [("04:00", "出", "休息後退勤", "休息後退勤", 1)],
        )

    def test_next_morning_0608_rest_checkout_is_not_duplicated_by_tomorrow_preview(self) -> None:
        module = duty_rehearsal_module()
        duty = "\u503c\u73ed"
        rest = "\u4f11\u606f"
        on_duty = "\u5728\u52e4"
        direction = "\u51fa\u6216\u5165"
        reason_key = "\u9818\u7528\u4e8b\u7531\u53ca\u5730\u9ede"
        today = module.DutySheet(
            roc_date="1150627",
            rows=[
                module.DutyRow("06-08", {duty: ["8"], rest: ["19"]}),
                module.DutyRow("08-10", {duty: ["12"], rest: []}),
                module.DutyRow("22-24", {duty: ["8"], "\u6551\u8b77": ["19"], rest: []}),
            ],
            summary={on_duty: ["8", "12", "19"]},
        )
        tomorrow = module.DutySheet(roc_date="1150628", rows=[module.DutyRow("08-10", {duty: ["12"]})], summary={on_duty: ["12"]})

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150627"), [], tomorrow)
        rest_checkouts = [
            action
            for action in actions
            if action.kind == "entry_log" and action.target == "19" and action.fields.get(reason_key) == "休息後退勤"
        ]

        self.assertEqual(
            [(action.time, action.actor, action.fields[direction], action.fields[reason_key], action.date_offset, action.source) for action in rest_checkouts],
            [("06:00", "8", "出", "休息後退勤", 1, "休息後退勤")],
        )

    def test_continuous_0408_rest_checkout_suppresses_0800_checkout(self) -> None:
        module = duty_rehearsal_module()
        yesterday = module.DutySheet(
            roc_date="1150627",
            rows=[
                module.DutyRow("04-06", {"值班": ["4"], "休息": ["4"]}),
                module.DutyRow("06-08", {"值班": ["4"], "休息": ["4"]}),
                module.DutyRow("08-10", {"值班": ["12"], "休息": []}),
                module.DutyRow("22-24", {"值班": ["4"], "休息": []}),
            ],
            summary={"在勤": ["4", "12"]},
        )
        today = module.DutySheet(roc_date="1150628", rows=[module.DutyRow("08-10", {"值班": ["12"]})], summary={"在勤": ["12"]})

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150628"), [], None)
        checkout_actions = [
            (action.time, action.fields["領用事由及地點"], action.source)
            for action in actions
            if action.kind == "entry_log" and action.target == "4" and action.fields.get("領用事由及地點") in ("退勤", "值退")
        ]

        self.assertNotIn("4", module.rest_starting_at(yesterday, 6, today))
        self.assertEqual(checkout_actions, [])

    def test_dynamic_0700_handoff_from_duty_rows(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150624",
            rows=[
                module.DutyRow("06-07", {"值班": ["28"]}),
                module.DutyRow("07-08", {"值班": ["16"]}),
                module.DutyRow("08-10", {"值班": ["12"]}),
            ],
            summary={"在勤": ["12", "16", "28"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150624"), [], None)
        handoff_0700 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "07:00"
        ]

        self.assertEqual(
            handoff_0700,
            [
                ("entry_log", "07:00", "28", "28", "值退", "值班交接"),
                ("entry_log", "07:00", "28", "16", "值班", "值班交接"),
                ("work_log", "07:00", "28", "28", None, "值班交接"),
            ],
        )

    def test_dynamic_0600_handoff_from_duty_rows(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150624",
            rows=[
                module.DutyRow("04-06", {"值班": ["28"]}),
                module.DutyRow("06-08", {"值班": ["16"]}),
            ],
            summary={"在勤": ["16", "28"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150624"), [], None)
        handoff_0600 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "06:00"
        ]

        self.assertEqual(
            handoff_0600,
            [
                ("entry_log", "06:00", "28", "28", "值退", "值班交接"),
                ("entry_log", "06:00", "28", "16", "值班", "值班交接"),
                ("work_log", "06:00", "28", "28", None, "值班交接"),
            ],
        )

    def test_dynamic_0000_handoff_from_duty_rows(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150624",
            rows=[
                module.DutyRow("00-02", {"值班": ["16"]}),
                module.DutyRow("22-24", {"值班": ["28"]}),
            ],
            summary={"在勤": ["16", "28"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150624"), [], None)
        handoff_0000 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "00:00"
        ]

        self.assertEqual(
            handoff_0000,
            [
                ("entry_log", "00:00", "28", "28", "值退", "值班交接"),
                ("entry_log", "00:00", "28", "16", "值班", "值班交接"),
                ("work_log", "00:00", "28", "28", None, "值班交接"),
            ],
        )

    def test_continuous_overnight_to_morning_duty_skips_0800_handoff(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150624",
            rows=[module.DutyRow("08-10", {"值班": ["28"]})],
            summary={"在勤": ["28"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150623",
            rows=[module.DutyRow("06-08", {"值班": ["28"]})],
            summary={"在勤": ["28"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150624"), [], None)
        handoff_0800 = [action for action in actions if action.source == "值班交接" and action.time == "08:00"]
        self.assertEqual(handoff_0800, [])

    def test_case_query_counts_two_unreturned_ambulances_at_2200_handoff(self) -> None:
        module = duty_rehearsal_module()
        two_people = "choose('" + "(^w^)" * 33 + "H01,L02')"

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[dict[str, object]] | list[str]]:
                return {
                    "headers": ["案件類別", "受理時間", "派遣時間", "返隊時間"],
                    "rows": [
                        {
                            "cells": ["緊急救護", "21:01:00", "21:03:00", "0001/01/01 00:00:00"],
                            "personnel_source": two_people,
                        },
                        {
                            "cells": ["緊急救護", "21:05:00", "21:07:00", "0001/01/01 00:00:00"],
                            "personnel_source": two_people,
                        },
                    ],
                }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(module, "js_click"), mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150806")

        items = module.unreturned_case_vehicle_items(
            cases,
            dict(module.DEFAULT_WORK_LOG_DEFAULTS),
            "1150806",
            before_hour=22,
        )

        self.assertEqual([case.return_time for case in cases], ["", ""])
        self.assertEqual(sum(item["count"] for item in items), 2)

    def test_case_query_counts_four_person_ems_case_as_two_vehicles(self) -> None:
        module = duty_rehearsal_module()
        four_people = "choose('" + "(^w^)" * 33 + "H01,L02,H03,S04')"

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[dict[str, object]] | list[str]]:
                return {
                    "headers": ["案件類別", "受理時間", "派遣時間", "返隊時間"],
                    "rows": [{
                        "cells": ["緊急救護", "21:01:00", "21:03:00", "0001/01/01 00:00:00"],
                        "personnel_source": four_people,
                    }],
                }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(module, "js_click"), mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150806")

        items = module.unreturned_case_vehicle_items(
            cases,
            dict(module.DEFAULT_WORK_LOG_DEFAULTS),
            "1150806",
            before_hour=22,
        )

        self.assertEqual(cases[0].personnel_count, 4)
        self.assertEqual(items[0]["count"], 2)

    def test_case_query_excludes_case_with_return_time_column_value(self) -> None:
        module = duty_rehearsal_module()
        four_people = "choose('" + "(^w^)" * 33 + "H01,L02,H03,S04')"

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[dict[str, object]] | list[str]]:
                return {
                    "headers": ["案件類別", "受理時間", "派遣時間", "返隊時間"],
                    "rows": [{
                        "cells": ["緊急救護", "21:01:00", "21:03:00", "21:45:00"],
                        "personnel_source": four_people,
                    }],
                }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(module, "js_click"), mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150806")

        items = module.unreturned_case_vehicle_items(
            cases,
            dict(module.DEFAULT_WORK_LOG_DEFAULTS),
            "1150806",
            before_hour=22,
        )

        self.assertEqual(cases[0].return_time, "21:45:00")
        self.assertEqual(items, [])

    def test_case_query_uses_aligned_return_value_when_header_index_is_misaligned(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[dict[str, object]] | list[str]]:
                return {
                    "headers": ["案件編號", "受理時間", "到達時間", "案件類別", "地點", "返隊時間"],
                    "rows": [{
                        "cells": [
                            "case-1",
                            "2026/08/12 15:20:35",
                            "2026/08/12 15:50:01",
                            "緊急救護-車禍",
                            "桃園市中壢區示例路",
                            "",
                        ],
                        "personnel_source": "",
                        "return_value": "2026/08/12 15:50:01",
                    }],
                }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(module, "js_click"), mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150812")

        items = module.unreturned_case_vehicle_items(
            cases,
            dict(module.DEFAULT_WORK_LOG_DEFAULTS),
            "1150812",
            before_hour=16,
        )

        self.assertEqual(cases[0].return_time, "15:50:01")
        self.assertEqual(items, [])

    def test_case_query_keeps_blank_aligned_return_value_unreturned(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[dict[str, object]] | list[str]]:
                return {
                    "headers": ["案件編號", "受理時間", "到達時間", "案件類別", "地點", "返隊時間"],
                    "rows": [{
                        "cells": [
                            "case-2",
                            "2026/08/12 15:20:35",
                            "2026/08/12 15:21:10",
                            "緊急救護-車禍",
                            "桃園市中壢區示例路",
                            "",
                        ],
                        "personnel_source": "",
                        "return_value": "",
                    }],
                }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(module, "js_click"), mock.patch.object(
            module.time, "sleep"
        ), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150812")

        items = module.unreturned_case_vehicle_items(
            cases,
            dict(module.DEFAULT_WORK_LOG_DEFAULTS),
            "1150812",
            before_hour=16,
        )

        self.assertEqual(cases[0].return_time, "")
        self.assertEqual(sum(item["count"] for item in items), 1)

    def test_visible_table_rejects_missing_query_controls(self) -> None:
        module = duty_rehearsal_module()

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "wait_for_any_form_control", return_value=False
        ), mock.patch.object(module, "js_set") as set_value, mock.patch.object(
            module, "wait_for_query_completion"
        ):
            with self.assertRaisesRegex(module.TimeoutException, "查詢按鈕"):
                driver = mock.Mock(execute_script=mock.Mock(return_value=[]))
                module.query_visible_table(driver, module.ENTRY_LOG_AP, "1151024")

        set_value.assert_not_called()

    def test_visible_table_rejects_missing_date_fields(self) -> None:
        module = duty_rehearsal_module()
        date_fields = {
            "_txtSDATE", "_txtSdate", "_txtSDate", "_txtEDATE", "_txtEDate",
            "_txtEndDate", "_txtDate", "_txtTaskDate",
        }

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "wait_for_any_form_control", return_value=True
        ), mock.patch.object(module, "suppress_window_open_for_background_query"), mock.patch.object(
            module, "js_set", side_effect=lambda _driver, field_id, _value: field_id not in date_fields
        ), mock.patch.object(module, "wait_for_query_completion"):
            with self.assertRaisesRegex(module.NoSuchElementException, "日期欄位"):
                driver = mock.Mock(execute_script=mock.Mock(return_value=[]))
                module.query_visible_table(driver, module.ENTRY_LOG_AP, "1151024")

    def test_visible_table_rejects_missing_query_button(self) -> None:
        module = duty_rehearsal_module()

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "wait_for_any_form_control", return_value=True
        ), mock.patch.object(module, "suppress_window_open_for_background_query"), mock.patch.object(
            module, "js_set", return_value=True
        ), mock.patch.object(module, "js_click", return_value=False), mock.patch.object(
            module, "wait_for_query_completion"
        ) as wait_for_completion:
            with self.assertRaisesRegex(module.NoSuchElementException, "查詢按鈕"):
                driver = mock.Mock(execute_script=mock.Mock(return_value=[]))
                module.query_visible_table(driver, module.ENTRY_LOG_AP, "1151024")

        wait_for_completion.assert_not_called()

    def test_visible_table_accepts_search_button_with_date_range_fields(self) -> None:
        module = duty_rehearsal_module()
        driver = mock.Mock(execute_script=mock.Mock(return_value=[]))

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "wait_for_any_form_control", return_value=True
        ), mock.patch.object(module, "suppress_window_open_for_background_query"), mock.patch.object(
            module, "js_set", return_value=True
        ), mock.patch.object(
            module, "js_click", side_effect=lambda _driver, field_id: field_id == "_btnSearch"
        ), mock.patch.object(module, "wait_for_query_completion"):
            rows = module.query_visible_table(driver, module.ENTRY_LOG_AP, "1151024")

        self.assertEqual(rows, [])

    def test_duty_login_rejects_redirect_before_entering_credentials(self) -> None:
        module = duty_rehearsal_module()

        for redirected_url in (
            "https://untrusted.example/tyfd119/login119",
            "http://dutymgt.tyfd.gov.tw/tyfd119/login119",
            "https://dutymgt.tyfd.gov.tw/unexpected-login",
        ):
            with self.subTest(url=redirected_url):
                driver = mock.Mock(current_url=redirected_url)
                with mock.patch.object(module, "WebDriverWait") as wait, mock.patch.object(
                    module, "detect_login_error", return_value=""
                ), mock.patch.object(module.time, "sleep"):
                    with self.assertRaisesRegex(RuntimeError, "授權.*登入頁面"):
                        module.login(driver, "test-user", "test-password")
                wait.return_value.until.return_value.send_keys.assert_not_called()
                driver.find_element.return_value.send_keys.assert_not_called()

    def test_visible_table_reads_all_pages_without_joining_page_headers(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def __init__(self, ap_name):
                self.ap_name = ap_name
                self.page = "1"

            def execute_script(self, script, *_args):
                if "Array.from(pageSelect.options" in script:
                    return ["1", "2"]
                indices = range(200) if self.page == "1" else range(200, 201)
                if self.ap_name == module.WORK_LOG_AP:
                    return [
                        ["115/10/24 08:00", "class", "unit", "duty", "", f"work-{i}", "status", "person", "1", "", "", "", ""]
                        for i in indices
                    ]
                rows = [["115/10/24 08:00", "unit", f"person-{i}", "role", "入", "值班"] for i in indices]
                return [["日期", "單位", "人員", "職稱", "出入", "事由"], *rows]

        for ap_name in (module.WORK_LOG_AP, module.ENTRY_LOG_AP):
            with self.subTest(ap_name=ap_name):
                driver = Driver(ap_name)
                with mock.patch.object(module, "open_ap"), mock.patch.object(
                    module, "wait_for_any_form_control", return_value=True
                ), mock.patch.object(module, "suppress_window_open_for_background_query"), mock.patch.object(
                    module, "js_set", return_value=True
                ), mock.patch.object(module, "js_click", return_value=True), mock.patch.object(
                    module, "arm_query_completion_tracking", return_value=100.0
                ), mock.patch.object(module, "wait_for_query_completion"), mock.patch.object(
                    module, "select_case_query_page", side_effect=lambda current, page: setattr(current, "page", page)
                ):
                    rows = module.query_visible_table(driver, ap_name, "1151024")

                self.assertEqual(len(rows), 201)
                if ap_name == module.ENTRY_LOG_AP:
                    self.assertEqual(len(rows[199]), 6)
                    self.assertEqual(rows[200][2], "person-200")
                else:
                    self.assertEqual(rows[200][5], "work-200")
                    self.assertEqual(rows[200][4], "")

    def test_duty_sheet_waits_for_fresh_rows_before_labelling_target_date(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def __init__(self):
                self.result_checks = 0

            def execute_script(self, script, *_args):
                if "return Array.from(document.querySelectorAll('table')).some" in script:
                    self.result_checks += 1
                    return "const tracker" not in script or self.result_checks >= 2
                if "function cellText(cell)" in script:
                    return {
                        "rows": [{"slot": "8-9", "columns": {"值班": "20" if self.result_checks >= 2 else "10"}}],
                        "captured_date": "1151024", "staff": {}, "summary": {}, "checkNums": [],
                    }
                return True

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency):
                self.driver = driver

            def until(self, condition):
                if condition(self.driver) or condition(self.driver):
                    return True
                raise module.TimeoutException()

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "WebDriverWait", ImmediateWait
        ), mock.patch.object(module, "suppress_window_open_for_background_query"), mock.patch.object(
            module, "js_set", return_value=True
        ), mock.patch.object(module, "js_click", return_value=True), mock.patch.object(
            module, "arm_query_completion_tracking", return_value=100.0
        ):
            sheet = module.query_duty_sheet(Driver(), "1151024")

        self.assertEqual(sheet.rows[0].columns["值班"], ["20"])

    def test_duty_sheet_rejects_captured_date_mismatch(self) -> None:
        module = duty_rehearsal_module()
        driver = mock.Mock()
        driver.execute_script.side_effect = lambda script, *_args: {
            "rows": [{"slot": "8-9", "columns": {"值班": "10"}}],
            "captured_date": "1151023", "staff": {}, "summary": {}, "checkNums": [],
        } if "function cellText(cell)" in script else True

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set", return_value=True), mock.patch.object(
            module, "js_click", return_value=True
        ), mock.patch.object(module, "arm_query_completion_tracking", return_value=100.0):
            with self.assertRaisesRegex(RuntimeError, "日期"):
                module.query_duty_sheet(driver, "1151024")

    def test_visible_table_waits_for_completion_and_requests_200_rows(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> list[list[str]]:
                return [["115/08/07 12:00", "新坡分隊", "測試人員", "隊員", "入", "到勤"]]

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set") as set_value, mock.patch.object(
            module, "js_click", return_value=True
        ), mock.patch.object(module, "wait_for_query_completion") as wait_for_completion:
            rows = module.query_visible_table(Driver(), module.ENTRY_LOG_AP, "1150807")

        set_value.assert_any_call(mock.ANY, "_txtPageNum", "200")
        wait_for_completion.assert_called_once_with(mock.ANY, previous_time_origin=None)
        self.assertEqual(len(rows), 1)

    def test_visible_table_accepts_cross_day_query_range(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> list[list[str]]:
                return [
                    ["115/08/13 23:49", "新坡分隊", "測試人員", "隊員", "出", "救護"],
                    ["115/08/13 23:50", "新坡分隊", "測試人員", "隊員", "出", "救護"],
                    ["115/08/14 00:15", "新坡分隊", "測試人員", "隊員", "入", "返隊"],
                    ["115/08/14 00:16", "新坡分隊", "測試人員", "隊員", "出", "救護"],
                ]

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set") as set_value, mock.patch.object(
            module, "js_click", return_value=True
        ), mock.patch.object(module.time, "sleep"), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            rows = module.query_visible_table(
                Driver(),
                module.ENTRY_LOG_AP,
                "1150814",
                start_roc_date="1150813",
                start_time="23:50",
                end_roc_date="1150814",
                end_time="00:15",
            )

        set_value.assert_any_call(mock.ANY, "_txtSDATE", "1150813")
        set_value.assert_any_call(mock.ANY, "_txtEDATE", "1150814")
        set_value.assert_any_call(mock.ANY, "_selSTIMEH", "23")
        set_value.assert_any_call(mock.ANY, "_selSTIMEM", "50")
        set_value.assert_any_call(mock.ANY, "_selETIMEH", "00")
        set_value.assert_any_call(mock.ANY, "_selETIMEM", "15")
        self.assertEqual([row[0] for row in rows], ["115/08/13 23:50", "115/08/14 00:15"])

    def test_external_assignment_state_orders_rows_by_absolute_datetime(self) -> None:
        module = package_module("compare_rehearsal_records")

        rows = module.flatten_rows(
            [
                ["115/08/13", "23:50", "-", "測試員", "出", "救護"],
                ["115/08/14", "00:15", "-", "測試員", "入", "救護返隊"],
            ],
            "1150814",
            start_date="1150813",
            end_date="1150814",
        )
        action = {
            "target": "10",
            "fields": {"出或入": "值退", "領用事由及地點": "退勤"},
        }
        staff = {"10": {"name": "測試員"}}

        self.assertIsNone(
            module.find_open_external_assignment(
                rows,
                "1150814",
                staff,
                action,
                current_at=datetime(2026, 8, 14, 0, 15),
                start_at=datetime(2026, 8, 13, 23, 50),
                end_at=datetime(2026, 8, 14, 0, 15),
            )
        )

    def test_external_assignment_is_closed_by_generic_return_row(self) -> None:
        module = package_module("compare_rehearsal_records")

        rows = module.flatten_rows(
            [
                ["115/09/22", "08:01", "-", "劉家誠", "出", "救護品管會議"],
                ["115/09/22", "12:00", "-", "劉家誠", "入", "返隊"],
            ],
            "1150922",
            start_date="1150922",
            end_date="1150922",
        )
        action = {
            "target": "19",
            "fields": {"出或入": "值班", "領用事由及地點": "值班"},
        }
        staff = {"19": {"name": "劉家誠"}}

        self.assertIsNone(
            module.find_open_external_assignment(
                rows,
                "1150922",
                staff,
                action,
                current_at=datetime(2026, 9, 22, 14, 0),
                start_at=datetime(2026, 9, 22, 8, 0),
                end_at=datetime(2026, 9, 22, 14, 0),
            )
        )

    def test_same_minute_case_return_is_independent_of_row_order(self) -> None:
        module = package_module("compare_rehearsal_records")
        for suffix in ("", " 梯次：2"):
            reason = "案件類別：緊急救護-創傷 案發地點：測試路段" + suffix
            records = [
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "出", reason],
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "入", reason.replace(" ", "&nbsp;")],
            ]
            for raw_rows in (records, list(reversed(records)), records + records):
                with self.subTest(suffix=suffix, rows=raw_rows):
                    self.assertIsNone(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151013"), "1151013",
                        {"10": {"name": "測試員"}}, {"target": "10"},
                        current_at=datetime(2026, 10, 13, 22, 0),
                    ))

    def test_same_minute_case_return_accepts_appended_incident_notes(self) -> None:
        module = package_module("compare_rehearsal_records")
        prefix = "案件類別：火災-其他 案發地點：桃園市觀音區一心路89號"
        for outgoing, incoming in (
            (prefix, prefix + "-電桿起火"),
            (prefix + "-電桿起火", prefix),
            (prefix + "-電桿起火", prefix + "（電線冒煙）"),
            (prefix, prefix + "，電桿起火"),
            (prefix + " 梯次：2", prefix + "－電桿起火 梯次：2"),
        ):
            records = [
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "出", outgoing],
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "入", incoming],
            ]
            for raw_rows in (records, list(reversed(records)), records + records):
                with self.subTest(outgoing=outgoing, incoming=incoming, rows=raw_rows):
                    self.assertIsNone(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151021"), "1151021",
                        {"27": {"name": "林宏為"}}, {"target": "27"},
                        current_at=datetime(2026, 10, 21, 14, 0),
                    ))

    def test_same_minute_case_return_accepts_explicit_fire_alarm_level_prefix(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：火災-一般(集合)住宅 案發地點：桃園市新屋區新文路1227巷123號"
        for prefix in ("一級火警-", "二級火警-", "三級火警-", "四級火警-", "五級火警-",
                       "1級火警-", "2級火警-", "3級火警-", "4級火警-", "5級火警-",
                       "二級火警－", "二級火警—", "二級火警–", "二級火警:", "二級火警："):
            ranked = reason.replace("案發地點：", "案發地點：" + prefix)
            for outgoing, incoming in ((reason, ranked), (ranked, reason),
                                       (reason + "-電線冒煙", ranked + "（屋內冒煙）")):
                records = [
                    ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "出", outgoing],
                    ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "入", incoming],
                ]
                for raw_rows in (records, list(reversed(records)), records + records):
                    with self.subTest(prefix=prefix, outgoing=outgoing, rows=raw_rows):
                        self.assertIsNone(module.find_open_external_assignment(
                            module.flatten_rows(raw_rows, "1151021"), "1151021",
                            {"15": {"name": "測試15番"}}, {"target": "15"},
                            current_at=datetime(2026, 10, 21, 14, 0),
                        ))

    def test_fire_alarm_level_prefix_preserves_address_identity_time_and_pairing_boundaries(self) -> None:
        from itertools import permutations

        module = package_module("compare_rehearsal_records")
        reason = "案件類別：火災-一般(集合)住宅 案發地點：桃園市新屋區新文路1227巷123號"
        ranked = reason.replace("案發地點：", "案發地點：二級火警-")
        for stamp, name, direction, incoming in (
            ("115/10/21 12:41", "另一人", "入", ranked),
            ("115/10/20 12:41", "測試15番", "入", ranked),
            ("115/10/21 12:40", "測試15番", "入", ranked),
            ("115/10/21 14:01", "測試15番", "入", ranked),
            ("115/10/21 12:41", "測試15番", "出", ranked),
            *(("115/10/21 12:41", "測試15番", "入", value) for value in (
                ranked.replace("1227巷", "1228巷"), ranked.replace("1227巷", "1227巷1弄"),
                ranked.replace("123號", "1230號"), ranked.replace("123號", "123-1號"),
                ranked + "之1", ranked + "-二樓", ranked + "-第一梯", ranked + " 梯次：2",
                ranked.replace("火災-一般(集合)住宅", "火災-其他"),
                *(ranked.replace("二級火警-", value) for value in (
                    "六級火警-", "6級火警-", "第二級火警-", "二級火災-", "重大火警-",
                    "二級火警", "備註二級火警-", "二級火警-二級火警-")),
            )),
        ):
            records = [
                ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "出", reason],
                [stamp, "新坡分隊", name, "隊員", direction, incoming],
            ]
            for raw_rows in (records, list(reversed(records))):
                with self.subTest(stamp=stamp, name=name, direction=direction, incoming=incoming):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151021", start_date="1151020"), "1151021",
                        {"15": {"name": "測試15番"}}, {"target": "15"},
                        current_at=datetime(2026, 10, 21, 14, 0),
                    ), "2026-10-21T12:41")
        records = [
            ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "出", reason],
            ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "入", ranked],
            ["115/10/21 12:41", "新坡分隊", "測試15番", "隊員", "出", ranked.replace("二級火警", "三級火警")],
        ]
        for raw_rows in permutations(records):
            self.assertEqual(module.find_open_external_assignment(
                module.flatten_rows(list(raw_rows), "1151021"), "1151021",
                {"15": {"name": "測試15番"}}, {"target": "15"},
                current_at=datetime(2026, 10, 21, 14, 0),
            ), "2026-10-21T12:41")

    def test_same_minute_incident_notes_preserve_address_and_identity_boundaries(self) -> None:
        module = package_module("compare_rehearsal_records")
        prefix = "案件類別：火災-其他 案發地點：桃園市觀音區一心路89號"
        for stamp, name, reason in (
            ("115/10/21 12:41", "林宏為", prefix.replace("89號", "890號") + "-電桿起火"),
            ("115/10/21 12:41", "林宏為", prefix.replace("89號", "89-1號") + "-電桿起火"),
            ("115/10/21 12:41", "林宏為", prefix + "之1-電桿起火"),
            ("115/10/21 12:41", "林宏為", prefix + "-1"),
            ("115/10/21 12:41", "林宏為", prefix + "-一樓"),
            ("115/10/21 12:41", "林宏為", prefix + "-A棟"),
            ("115/10/21 12:41", "林宏為", prefix + "-地下一層"),
            ("115/10/21 12:41", "林宏為", prefix + "（A座）"),
            ("115/10/21 12:41", "林宏為", prefix + "-Ａ座"),
            ("115/10/21 12:41", "林宏為", prefix.replace("一心路", "二心路") + "-電桿起火"),
            ("115/10/21 12:41", "林宏為", prefix.replace("火災-其他", "緊急救護-急病") + "-電桿起火"),
            ("115/10/21 12:41", "另一人", prefix + "-電桿起火"),
            ("115/10/20 12:41", "林宏為", prefix + "-電桿起火"),
            ("115/10/21 12:40", "林宏為", prefix + "-電桿起火"),
            ("115/10/21 12:41", "林宏為", prefix + "-電桿起火 梯次：2"),
            ("115/10/21 12:41", "林宏為", prefix + "-第二梯"),
            ("115/10/21 12:41", "林宏為", prefix + "-電桿起火 梯次：二"),
        ):
            records = [
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "出", prefix],
                [stamp, "新坡分隊", name, "隊員", "入", reason],
            ]
            for raw_rows in (records, list(reversed(records))):
                with self.subTest(stamp=stamp, name=name, reason=reason):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151021", start_date="1151020"), "1151021",
                        {"27": {"name": "林宏為"}}, {"target": "27"},
                        current_at=datetime(2026, 10, 21, 14, 0),
                    ), "2026-10-21T12:41")

    def test_same_minute_incident_notes_preserve_handwritten_wave_markers(self) -> None:
        module = package_module("compare_rehearsal_records")
        prefix = "案件類別：火災-其他 案發地點：桃園市觀音區一心路89號"
        for outgoing, incoming in (
            (prefix + "-第一梯", prefix + "-第二梯"),
            (prefix + "-電桿起火 梯次：一", prefix + "-電桿起火 梯次：二"),
        ):
            records = [
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "出", outgoing],
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "入", incoming],
            ]
            for raw_rows in (records, list(reversed(records))):
                with self.subTest(outgoing=outgoing, incoming=incoming):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151021"), "1151021",
                        {"27": {"name": "林宏為"}}, {"target": "27"},
                        current_at=datetime(2026, 10, 21, 14, 0),
                    ), "2026-10-21T12:41")

    def test_same_minute_incident_note_pairing_rejects_multiple_candidates(self) -> None:
        from itertools import permutations

        module = package_module("compare_rehearsal_records")
        prefix = "案件類別：火災-其他 案發地點：桃園市觀音區一心路89號"
        for direction, time, extra_reason in (
            ("出", "12:41", prefix + "-其他電桿"),
            ("入", "12:41", prefix + "-其他電桿"),
            ("出", "12:40", prefix),
            ("入", "12:40", prefix),
        ):
            records = [
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "出", prefix],
                ["115/10/21 12:41", "新坡分隊", "林宏為", "隊員", "入", prefix + "-電桿起火"],
                ["115/10/21 " + time, "新坡分隊", "林宏為", "隊員", direction, extra_reason],
            ]
            for raw_rows in permutations(records):
                with self.subTest(direction=direction, time=time):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(list(raw_rows), "1151021"), "1151021",
                        {"27": {"name": "林宏為"}}, {"target": "27"},
                        current_at=datetime(2026, 10, 21, 14, 0),
                    ), "2026-10-21T12:41")

    def test_same_minute_return_requires_matching_person_case_and_wave(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段 梯次：2"
        for stamp, name, incoming_reason in (
            ("115/10/13 17:42", "另一人", reason),
            ("115/10/13 17:42", "測試員", reason.replace("創傷", "急病")),
            ("115/10/13 17:42", "測試員", reason.replace("測試路段", "其他路段")),
            ("115/10/13 17:42", "測試員", reason.replace("梯次：2", "梯次：1")),
            ("115/10/13 17:42", "測試員", reason.replace(" 梯次：2", "")),
            ("115/10/12 17:42", "測試員", reason),
            ("115/10/13 17:42", "測試員", "返隊"),
            ("115/10/13 17:42", "測試員", "救護返隊"),
        ):
            records = [
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "出", reason],
                [stamp, "新坡分隊", name, "隊員", "入", incoming_reason],
            ]
            for raw_rows in (records, list(reversed(records))):
                with self.subTest(stamp=stamp, name=name, reason=incoming_reason, rows=raw_rows):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(raw_rows, "1151013", start_date="1151012"), "1151013",
                        {"10": {"name": "測試員"}}, {"target": "10"},
                        current_at=datetime(2026, 10, 13, 22, 0),
                    ), "2026-10-13T17:42")

    def test_same_minute_case_return_preserves_other_departures(self) -> None:
        from itertools import permutations

        module = package_module("compare_rehearsal_records")
        base_reason = "案件類別：緊急救護-創傷 案發地點：測試路段"
        for reason, extra_reason, extra_time in (
            (base_reason, base_reason.replace("測試路段", "其他路段"), "17:42"),
            (base_reason + " 梯次：2", base_reason + " 梯次：3", "17:42"),
            (base_reason, base_reason, "17:43"),
        ):
            records = [
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "出", reason],
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "入", reason],
                [f"115/10/13 {extra_time}", "新坡分隊", "測試員", "隊員", "出", extra_reason],
            ]
            for raw_rows in permutations(records):
                with self.subTest(time=extra_time, reason=extra_reason, rows=raw_rows):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(list(raw_rows), "1151013"), "1151013",
                        {"10": {"name": "測試員"}}, {"target": "10"},
                        current_at=datetime(2026, 10, 13, 22, 0),
                    ), f"2026-10-13T{extra_time}")

    def test_same_minute_case_return_is_conservative_for_repeated_trips(self) -> None:
        from itertools import permutations

        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段"
        for extra_direction, extra_time in (("出", "17:40"), ("入", "17:41")):
            records = [
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "出", reason],
                ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "入", reason],
                [f"115/10/13 {extra_time}", "新坡分隊", "測試員", "隊員", extra_direction, reason],
            ]
            for raw_rows in permutations(records):
                with self.subTest(direction=extra_direction, rows=raw_rows):
                    self.assertEqual(module.find_open_external_assignment(
                        module.flatten_rows(list(raw_rows), "1151013"), "1151013",
                        {"10": {"name": "測試員"}}, {"target": "10"},
                        current_at=datetime(2026, 10, 13, 22, 0),
                    ), "2026-10-13T17:42")

    def test_no_wave_case_does_not_close_an_earlier_return(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段"
        records = [
            ["115/10/13 17:42", "新坡分隊", "測試員", "隊員", "出", reason],
            ["115/10/13 17:41", "新坡分隊", "測試員", "隊員", "入", reason],
        ]
        for raw_rows in (records, list(reversed(records))):
            self.assertEqual(module.find_open_external_assignment(
                module.flatten_rows(raw_rows, "1151013"), "1151013",
                {"10": {"name": "測試員"}}, {"target": "10"},
                current_at=datetime(2026, 10, 13, 22, 0),
            ), "2026-10-13T17:42")

    def test_external_assignment_reversed_case_return_is_returned(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段 梯次：2"
        records = [
            ["115/09/13 06:41", "新坡分隊", "測試員", "隊員", "出", reason],
            ["115/09/13 06:38", "新坡分隊", "測試員", "隊員", "入", reason],
        ]
        for raw_rows in (records, list(reversed(records)), records + records):
            with self.subTest(rows=raw_rows):
                self.assertIsNone(module.find_open_external_assignment(
                    module.flatten_rows(raw_rows, "1150913"), "1150913",
                    {"13": {"name": "測試員"}}, {"target": "13"},
                    current_at=datetime(2026, 9, 13, 8, 1),
                ))

    def test_external_assignment_reversed_return_requires_matching_case(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段 梯次：2"
        incoming_variants = [
            ("115/09/13 06:38", "另一人", reason),
            ("115/09/13 06:38", "測試員", reason.replace("梯次：2", "梯次：1")),
            ("115/09/13 06:38", "測試員", reason.replace("測試路段", "其他路段")),
            ("115/09/13 06:38", "測試員", reason.replace("創傷", "急病")),
            ("115/09/12 06:38", "測試員", reason),
            ("115/09/13 06:38", "測試員", "救護返隊"),
        ]
        for stamp, name, incoming_reason in incoming_variants:
            with self.subTest(stamp=stamp, name=name, reason=incoming_reason):
                rows = module.flatten_rows([
                    ["115/09/13 06:41", "新坡分隊", "測試員", "隊員", "出", reason],
                    [stamp, "新坡分隊", name, "隊員", "入", incoming_reason],
                ], "1150913", start_date="1150912", end_date="1150913")
                self.assertEqual(module.find_open_external_assignment(
                    rows, "1150913", {"13": {"name": "測試員"}}, {"target": "13"},
                    current_at=datetime(2026, 9, 13, 8, 1),
                ), "2026-09-13T06:41")

    def test_external_assignment_reversed_return_preserves_other_departures(self) -> None:
        module = package_module("compare_rehearsal_records")
        reason = "案件類別：緊急救護-創傷 案發地點：測試路段 梯次：2"
        for extra_time, extra_reason, expected in (
            ("06:40", reason.replace("梯次：2", "梯次：3"), "06:40"),
            ("07:00", reason, "07:00"),
        ):
            with self.subTest(time=extra_time):
                rows = module.flatten_rows([
                    ["115/09/13 06:41", "新坡分隊", "測試員", "隊員", "出", reason],
                    ["115/09/13 06:38", "新坡分隊", "測試員", "隊員", "入", reason],
                    [f"115/09/13 {extra_time}", "新坡分隊", "測試員", "隊員", "出", extra_reason],
                ], "1150913")
                self.assertEqual(module.find_open_external_assignment(
                    rows, "1150913", {"13": {"name": "測試員"}}, {"target": "13"},
                    current_at=datetime(2026, 9, 13, 8, 1),
                ), f"2026-09-13T{expected}")

    def test_external_assignment_state_includes_disaster_rescue_departure(self) -> None:
        module = package_module("compare_rehearsal_records")

        rows = module.flatten_rows(
            [["115/09/07", "08:00", "-", "測試員", "出", "災害搶救"]],
            "1150907",
        )
        action = {
            "target": "10",
            "fields": {"出或入": "值退", "領用事由及地點": "退勤"},
        }
        staff = {"10": {"name": "測試員"}}

        self.assertEqual(
            module.find_open_external_assignment(
                rows,
                "1150907",
                staff,
                action,
                current_at=datetime(2026, 9, 7, 8, 10),
            ),
            "2026-09-07T08:00",
        )

    def test_case_query_uses_silent_field_updates_and_native_query_click(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[object]]:
                return {"headers": ["返隊時間"], "rows": []}

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module.time, "sleep"), mock.patch.object(
            module, "js_set", return_value=True
        ) as set_value, mock.patch.object(
            module, "native_click", return_value=True
        ) as native_query, mock.patch.object(module, "wait_for_query_completion"):
            cases = module.query_cases(Driver(), "1150808")

        self.assertEqual(cases, [])
        self.assertEqual(set_value.call_count, 7)
        self.assertTrue(all(call.kwargs == {"dispatch_change": False} for call in set_value.call_args_list))
        native_query.assert_called_once_with(mock.ANY, "_btnQuery")

    def test_case_query_treats_an_empty_result_as_no_cases(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str, *_args: object) -> dict[str, list[object]]:
                return {"headers": [], "rows": []}

        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module.time, "sleep"), mock.patch.object(
            module, "js_set", return_value=True
        ), mock.patch.object(module, "native_click", return_value=True), mock.patch.object(
            module, "wait_for_query_completion"
        ):
            cases = module.query_cases(Driver(), "1150809")

        self.assertEqual(cases, [])

    @unittest.skipUnless(os.environ.get("SINPOSMART_TEST_BROWSER") == "1", "opt-in local Chrome DOM regression")
    def test_page_refresh_ignores_loading_changes_in_control_rows(self) -> None:
        from urllib.parse import quote

        module = duty_rehearsal_module()
        html = """<html><body><p>QUY-000 : 查詢完成</p>
        <table><tr><td id="spinner">ready</td><td>query</td><td>controls</td></tr></table>
        <select name="pageSelect" onchange="document.getElementById('spinner').textContent='loading';
          setTimeout(()=>{const old=document.querySelector('[data-id]').closest('tr');const fresh=old.cloneNode(true);
          fresh.querySelector('[data-id]').textContent='second';old.replaceWith(fresh);},600)">
        <option value="1">1</option><option value="2">2</option></select>
        <table><tr><th>返隊時間</th><th>id</th><th>time</th></tr>
        <tr><td></td><td data-id>first</td><td>12:00:00</td></tr></table></body></html>"""

        with tempfile.TemporaryDirectory(prefix="sinposmart-query-dom-") as temp_dir:
            driver = module.build_driver(
                headless=True,
                profile_root=Path(temp_dir).resolve(),
                prune_profiles=False,
                write_diagnostics=False,
            )
            try:
                driver.get("data:text/html;charset=utf-8," + quote(html))
                previous_origin = module.arm_query_completion_tracking(driver)
                module.select_case_query_page(driver, "2")
                module.wait_for_query_completion(driver, expected_page="2", previous_time_origin=previous_origin)
                self.assertEqual(
                    driver.execute_script("return document.querySelector('[data-id]').textContent"),
                    "second",
                )
            finally:
                module.quit_driver(driver)

    @unittest.skipUnless(os.environ.get("SINPOSMART_TEST_BROWSER") == "1", "opt-in local Chrome DOM regression")
    def test_page_refresh_ignores_loading_changes_inside_existing_record(self) -> None:
        from urllib.parse import quote

        module = duty_rehearsal_module()
        with tempfile.TemporaryDirectory(prefix="sinposmart-record-query-dom-") as temp_dir:
            driver = module.build_driver(
                headless=True, profile_root=Path(temp_dir).resolve(),
                prune_profiles=False, write_diagnostics=False,
            )
            try:
                for identical in (False, True):
                    with self.subTest(identical=identical):
                        update = "" if identical else "fresh.querySelector('[data-id]').textContent='second';"
                        html = """<html><body><p>QUY-000 : 查詢完成</p>
                        <select name="pageSelect" onchange="document.getElementById('spinner').firstChild.data='loading';
                          setTimeout(()=>{const old=document.getElementById('record');const fresh=old.cloneNode(true);
                          fresh.setAttribute('data-page','2');UPDATE old.replaceWith(fresh);},600)">
                        <option value="1">1</option><option value="2">2</option></select>
                        <table><tr><th>返隊時間</th><th>id</th><th>time</th><th>status</th></tr>
                        <tr id="record" data-page="1"><td></td><td data-id>first</td><td>12:00:00</td>
                        <td><span id="spinner">ready</span></td></tr></table></body></html>""".replace("UPDATE", update)
                        driver.get("data:text/html;charset=utf-8," + quote(html))
                        previous_origin = module.arm_query_completion_tracking(driver)
                        module.select_case_query_page(driver, "2")
                        module.wait_for_query_completion(driver, expected_page="2", previous_time_origin=previous_origin)
                        self.assertEqual(
                            driver.execute_script("return document.getElementById('record').getAttribute('data-page')"),
                            "2",
                        )
            finally:
                module.quit_driver(driver)

    def test_page_completion_rejects_old_rows_until_refresh_finishes(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def __init__(self):
                self.checks = 0

            def execute_script(self, _script: str, *_args) -> dict[str, object]:
                self.checks += 1
                return {
                    "completed": True, "page": "2", "hasRows": True,
                    "freshRows": self.checks >= 2, "timeOrigin": 100.0,
                    "readyState": "complete" if self.checks >= 3 else "loading",
                }

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency):
                self.driver = driver

            def until(self, condition):
                for _ in range(3):
                    if condition(self.driver):
                        return True
                raise AssertionError("the refreshed result did not become ready")

        driver = Driver()
        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            module.wait_for_query_completion(driver, expected_page="2", previous_time_origin=100.0)

        self.assertEqual(driver.checks, 3)

    def test_page_completion_accepts_identical_rows_after_ajax_refresh(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def __init__(self):
                self.checks = 0

            def execute_script(self, _script: str, *_args) -> dict[str, object]:
                self.checks += 1
                return {
                    "completed": False, "page": "2", "hasRows": True,
                    "freshRows": self.checks >= 2, "timeOrigin": 100.0, "readyState": "complete",
                }

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency):
                self.driver = driver

            def until(self, condition):
                if condition(self.driver) or condition(self.driver):
                    return True
                raise AssertionError("the identical refreshed rows were rejected")

        driver = Driver()
        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            module.wait_for_query_completion(driver, expected_page="2", previous_time_origin=100.0)

        self.assertEqual(driver.checks, 2)

    def test_page_completion_accepts_a_fresh_no_records_result(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def execute_script(self, _script: str, *_args) -> dict[str, object]:
                return {
                    "completed": True, "page": "", "hasRows": False,
                    "freshCompletion": True, "noRecords": True, "readyState": "complete",
                }

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency):
                self.driver = driver

            def until(self, condition):
                if not condition(self.driver):
                    raise AssertionError("a fresh QUY-500 result must finish the query")
                return True

        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            module.wait_for_query_completion(Driver(), expected_page="2")

    def test_page_completion_accepts_loaded_second_page_without_query_banner(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(_script: str) -> dict[str, object]:
                return {
                    "completed": False,
                    "page": "2",
                    "hasRows": True,
                    "freshCompletion": False,
                    "freshRows": True,
                    "timeOrigin": 100.0,
                }

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency) -> None:
                self.driver = driver

            def until(self, condition) -> bool:
                if not condition(self.driver):
                    raise module.TimeoutException()
                return True

        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            module.wait_for_query_completion(Driver(), expected_page="2")

    def test_query_completion_rejects_old_success_marker_until_this_query_finishes(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            def __init__(self, fresh: bool) -> None:
                self.fresh = fresh

            def execute_script(self, _script: str) -> dict[str, object]:
                return {"completed": True, "page": "", "hasRows": False,
                        "freshCompletion": self.fresh, "timeOrigin": 100.0}

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency) -> None:
                self.driver = driver

            def until(self, condition) -> bool:
                if not condition(self.driver):
                    raise module.TimeoutException()
                return True

        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            with self.assertRaises(module.TimeoutException):
                module.wait_for_query_completion(Driver(False), previous_time_origin=100.0)
            module.wait_for_query_completion(Driver(True), previous_time_origin=100.0)

    def test_entry_query_stops_when_query_button_never_loads(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            pass

        with mock.patch.object(module, "open_ap"), \
             mock.patch.object(module, "wait_for_any_form_control", return_value=False), \
             mock.patch.object(module, "suppress_window_open_for_background_query"), \
             mock.patch.object(module, "js_set", return_value=True), \
             mock.patch.object(module, "js_click", return_value=False), \
             mock.patch.object(module, "wait_for_query_completion"):
            with self.assertRaises(module.TimeoutException):
                module.query_visible_table(Driver(), module.ENTRY_LOG_AP, "1150923")

    def test_entry_query_stops_when_query_button_cannot_be_clicked(self) -> None:
        module = duty_rehearsal_module()
        query_started: list[bool] = []

        class Driver:
            @staticmethod
            def execute_script(script: str) -> dict[str, float] | list[object]:
                if "performance.timeOrigin" in script:
                    return {"timeOrigin": 100.0}
                return []

        with mock.patch.object(module, "open_ap"), \
             mock.patch.object(module, "wait_for_any_form_control", return_value=True), \
             mock.patch.object(module, "suppress_window_open_for_background_query"), \
             mock.patch.object(module, "js_set", return_value=True), \
             mock.patch.object(module, "js_click", return_value=False), \
             mock.patch.object(module, "wait_for_query_completion"):
            with self.assertRaises(module.NoSuchElementException):
                module.query_visible_table(
                    Driver(),
                    module.ENTRY_LOG_AP,
                    "1150923",
                    query_started_callback=lambda: query_started.append(True),
                )
        self.assertEqual(query_started, [])

    def test_entry_query_reports_start_after_successful_click(self) -> None:
        module = duty_rehearsal_module()
        events: list[str] = []

        class Driver:
            @staticmethod
            def execute_script(script: str) -> dict[str, float] | list[object]:
                if "performance.timeOrigin" in script:
                    return {"timeOrigin": 100.0}
                return []

        with mock.patch.object(module, "open_ap"), \
             mock.patch.object(module, "wait_for_any_form_control", return_value=True), \
             mock.patch.object(module, "suppress_window_open_for_background_query"), \
             mock.patch.object(module, "js_set", return_value=True), \
             mock.patch.object(module, "js_click", side_effect=lambda *_args: events.append("click") or True), \
             mock.patch.object(module, "wait_for_query_completion"):
            module.query_visible_table(
                Driver(),
                module.ENTRY_LOG_AP,
                "1150923",
                query_started_callback=lambda: events.append("query_started"),
            )
        self.assertEqual(events, ["click", "query_started"])

    def test_entry_query_keeps_search_button_fallback(self) -> None:
        module = duty_rehearsal_module()
        clicked: list[str] = []

        class Driver:
            @staticmethod
            def execute_script(script: str) -> dict[str, float] | list[object]:
                if "performance.timeOrigin" in script:
                    return {"timeOrigin": 100.0}
                return []

        def click(_driver, button_id: str) -> bool:
            clicked.append(button_id)
            return button_id == "_btnSearch"

        with mock.patch.object(module, "open_ap"), \
             mock.patch.object(module, "wait_for_any_form_control", return_value=True), \
             mock.patch.object(module, "suppress_window_open_for_background_query"), \
             mock.patch.object(module, "js_set", return_value=True), \
             mock.patch.object(module, "js_click", side_effect=click), \
             mock.patch.object(module, "wait_for_query_completion"):
            module.query_visible_table(Driver(), module.ENTRY_LOG_AP, "1150923")

        self.assertEqual(clicked, ["_btnQuery", "_btnSearch"])

    def test_case_query_stops_when_query_form_never_loads(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            pass

        with mock.patch.object(module, "open_ap"), \
             mock.patch.object(module, "wait_for_form_controls", return_value=False), \
             mock.patch.object(module, "suppress_window_open_for_background_query"):
            with self.assertRaises(module.TimeoutException):
                module.query_cases(Driver(), "1150923")

    def test_any_query_button_wait_accepts_search_alias(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            @staticmethod
            def execute_script(script: str, control_ids: tuple[str, ...]) -> bool:
                return ".some(" in script and "_btnSearch" in control_ids

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency) -> None:
                self.driver = driver

            def until(self, condition) -> bool:
                return condition(self.driver)

        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            self.assertTrue(
                module.wait_for_any_form_control(
                    Driver(),
                    ("_btnQuery", "_btnSearch"),
                    timeout=1,
                )
            )

    def test_query_completion_treats_no_records_as_a_finished_query(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            scripts: list[str] = []

            def execute_script(self, script: str) -> dict[str, object]:
                self.scripts.append(script)
                return {
                    "completed": True,
                    "page": "",
                    "hasRows": False,
                    "freshCompletion": True,
                    "freshRows": False,
                    "timeOrigin": 100.0,
                }

        class ImmediateWait:
            def __init__(self, driver, _timeout, *, poll_frequency) -> None:
                self.driver = driver

            def until(self, condition) -> bool:
                return condition(self.driver)

        driver = Driver()
        with mock.patch.object(module, "WebDriverWait", ImmediateWait):
            module.wait_for_query_completion(driver)

        self.assertIn("QUY-500", driver.scripts[0])

    def test_case_query_reads_all_pages(self) -> None:
        module = duty_rehearsal_module()

        class Driver:
            page = "1"

            class Option:
                def __init__(self, owner) -> None:
                    self.owner = owner

                def click(self) -> None:
                    self.owner.page = "2"

            class PageSelect:
                def __init__(self, owner) -> None:
                    self.owner = owner

                def find_element(self, _by, selector: str):
                    if selector == "option[value='2']":
                        return Driver.Option(self.owner)
                    raise module.NoSuchElementException()

            def find_element(self, _by, selector: str):
                if selector == "select[name='pageSelect']":
                    return Driver.PageSelect(self)
                raise module.NoSuchElementException()

            def execute_script(self, script: str, *args: str) -> object:
                if "Array.from(pageSelect.options" in script:
                    return ["1", "2"]
                if self.page == "1":
                    return {
                        "headers": ["案件類別", "受理時間", "派遣時間", "返隊時間"],
                        "rows": [{
                            "cells": ["緊急救護", "21:01:00", "21:03:00", "0001/01/01 00:00:00"],
                            "personnel_source": "",
                        }],
                    }
                return {
                    "headers": ["案件類別", "受理時間", "派遣時間", "返隊時間"],
                    "rows": [{
                        "cells": ["火災", "22:01:00", "22:03:00", "0001/01/01 00:00:00"],
                        "personnel_source": "",
                    }],
                }

        driver = Driver()
        with mock.patch.object(module, "open_ap"), mock.patch.object(
            module, "suppress_window_open_for_background_query"
        ), mock.patch.object(module, "js_set"), mock.patch.object(
            module, "js_click"
        ), mock.patch.object(module, "wait_for_query_completion") as wait_for_completion:
            cases = module.query_cases(driver, "1150807")

        self.assertEqual([case.report_time for case in cases], ["21:01:00", "22:01:00"])
        self.assertEqual(
            wait_for_completion.call_args_list,
            [
                mock.call(driver, previous_time_origin=None),
                mock.call(driver, expected_page="2", previous_time_origin=None),
            ],
        )

    def test_0800_handoff_uses_yesterday_2200_duty_when_today_0608_row_absent(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150626",
            rows=[module.DutyRow("08-10", {"值班": ["12"]})],
            summary={"在勤": ["12", "26"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150625",
            rows=[module.DutyRow("22-24", {"值班": ["26"]})],
            summary={"在勤": ["12", "26"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150626"), [], None)
        handoff_0800 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "08:00"
        ]

        self.assertEqual(
            handoff_0800,
            [
                ("entry_log", "08:00", "26", "26", "值退", "值班交接"),
                ("entry_log", "08:00", "26", "12", "值班", "值班交接"),
                ("work_log", "08:00", "26", "26", None, "值班交接"),
            ],
        )

    def test_next_fire_day_external_preview_keeps_pair_dates_and_unique_keys(self) -> None:
        from dataclasses import asdict
        from datetime import timedelta

        module = duty_rehearsal_module()
        from app_core.duty_task_projection import action_datetime

        for target in (datetime(2026, 10, 21).date(), datetime(2026, 10, 31).date(), datetime(2026, 12, 31).date()):
            for end_slot in ("08-10", "08-24"):
                with self.subTest(target=target, end_slot=end_slot):
                    following = target + timedelta(days=1)
                    today = module.DutySheet(
                        roc_date=module.roc_date(target),
                        rows=[module.DutyRow("08-10", {"值班": ["3"], "外勤支援": ["21"]}),
                              module.DutyRow("10-12", {"值班": ["3"]}),
                              module.DutyRow("22-24", {"值班": ["8"]})],
                        summary={"在勤": ["3", "8", "21"]},
                    )
                    tomorrow = module.DutySheet(
                        roc_date=module.roc_date(following),
                        rows=[module.DutyRow(end_slot, {"值班": ["11"], "外勤支援": ["21"]})]
                             + ([module.DutyRow("10-12", {"值班": ["12"]}),
                                 module.DutyRow("12-14", {"值班": ["12"], "外勤支援": ["22"]})]
                                if end_slot == "08-10" else []),
                        summary={"在勤": ["11", "12", "21", "22"]},
                    )
                    with mock.patch.object(module, "load_work_log_defaults", return_value=dict(module.DEFAULT_WORK_LOG_DEFAULTS)):
                        standalone = module.planned_actions(tomorrow, today, [], following)
                        actions = module.planned_actions(today, None, [], target, [], tomorrow)
                    expected = [action for action in standalone if action.source.startswith("外勤") and action.target == "21"]
                    previews = [action for action in actions if action.source.startswith("外勤") and action.target == "21"
                                and action.date_offset >= 1]
                    self.assertEqual(len(previews), 2)
                    self.assertEqual(previews[0].actor, "8")
                    self.assertEqual({action.staff_roc_date for action in previews}, {tomorrow.roc_date})
                    self.assertEqual(len({action.return_pair_key for action in previews}), 1)
                    self.assertTrue(previews[0].return_pair_key)
                    self.assertEqual([action.duplicate_key for action in previews], [action.duplicate_key for action in expected])
                    self.assertEqual(
                        [action_datetime(asdict(action), today.roc_date) for action in previews],
                        [action_datetime(asdict(action), tomorrow.roc_date) for action in expected],
                    )
                    self.assertEqual(len([action for action in actions if action.source == "外勤簽出" and action.target == "21"]), 2)
                    self.assertFalse(any(action.source == "外勤簽出" and action.target == "22" for action in actions))

    def test_next_fire_day_external_wait_uses_next_roster_without_submitting_early(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        duty_rehearsal_module()
        from PySide6.QtTest import QSignalSpy
        from PySide6.QtWidgets import QApplication
        from app_core.schedule_repository import ScheduleRepository
        from app_core.unreturned_return_queue import UnreturnedReturnQueue
        from qt_app.controllers.duty_controller import DutyController

        app = QApplication.instance() or QApplication(["next-fire-day-preview"])
        for early in (datetime(2026, 10, 21, 22), datetime(2026, 10, 22, 1)):
            current = early

            class Clock(datetime):
                @classmethod
                def now(cls, tz=None):
                    return current.replace(tzinfo=tz)

            with self.subTest(early=early), tempfile.TemporaryDirectory() as temp_dir, mock.patch(
                "qt_app.controllers.duty_controller.datetime", Clock,
            ):
                controller = DutyController(
                    repository=ScheduleRepository(Path(temp_dir)),
                    unreturned_return_queue=UnreturnedReturnQueue(Path(temp_dir)),
                )
                try:
                    controller.set_session_context(1, "night-user")
                    controller.set_actor_no("8")
                    departure = {
                        "kind": "entry_log", "time": "08:00", "date_offset": 1,
                        "actor": "8", "target": "21", "source": "外勤簽出",
                        "duplicate_key": "entry:2026-10-22:8:out:21:外勤支援",
                        "return_pair_key": "return-pair:next-day", "staff_roc_date": "1151022",
                        "fields": {"出或入": "出", "領用事由及地點": "外勤支援"},
                    }
                    returning = {**departure, "time": "10:00", "actor": "11", "source": "外勤簽入",
                                 "duplicate_key": "entry:2026-10-22:10:in:21:返隊:外勤支援",
                                 "fields": {"出或入": "入", "領用事由及地點": "返隊"}}
                    controller.replace_schedule_data({
                        "target_date": "1151021", "today": {"roc_date": "1151021", "staff": {"21": {"name": "今日人員"}}},
                        "tomorrow": {"roc_date": "1151022", "staff_complete": True, "staff": {"21": {"name": "明日人員"}}},
                        "actions": [departure, returning],
                    })
                    row = next(row for row in controller._task_model._tasks if row["taskIndex"] == 0)
                    self.assertIn("明日人員", row["peopleText"])
                    self.assertNotIn("今日人員", row["peopleText"])
                    submitted = QSignalSpy(controller.manualSubmissionRequested)
                    controller.toggleTaskSelection(0)
                    self.assertTrue(controller.canManualSubmitSelected)
                    controller.prepareManualSubmission()
                    self.assertIn("明日人員", controller.manualConfirmationSummary)
                    controller.confirmManualSubmission()
                    self.assertEqual(submitted.count(), 0)
                    self.assertEqual(controller.manual_submission_requests("night-user", "fake", [0], submit_at=early), [])
                    waiting = controller.background_manual_waiting_requests("night-user", "fake", [0])
                    self.assertEqual(len(waiting), 1)
                    self.assertEqual(len(waiting[0].schedule_data["actions"]), 2)
                    current = datetime(2026, 10, 22, 7, 59)
                    controller._refresh_due_tasks()
                    self.assertEqual(submitted.count(), 0)
                    current = datetime(2026, 10, 22, 8)
                    controller._refresh_due_tasks()
                    controller._refresh_due_tasks()
                    self.assertEqual(submitted.count(), 1)
                    requests = controller.manual_submission_requests("night-user", "fake", [0], submit_at=current)
                    self.assertEqual(len(requests), 1)
                    submitted_action = requests[0].schedule_data["actions"][0]
                    self.assertEqual(submitted_action["submit_target_date"], "1151022")
                    self.assertEqual(submitted_action["fields"]["系統寫入時間"], "08:00")
                finally:
                    controller.shutdown()
        self.assertIsNotNone(app)

    def test_next_fire_day_external_submission_uses_original_roster_and_verifies_once(self) -> None:
        from types import SimpleNamespace

        duty_rehearsal_module()
        from app_core.duty_submission_service import DutySubmissionRequest, DutySubmissionService

        with tempfile.TemporaryDirectory() as temp_dir:
            fills = []

            def fill(_driver, action, staff, action_date, *, save):
                fills.append((staff["21"]["name"], action_date, save))
                return {}

            automation = SimpleNamespace(
                WORK_LOG_AP="work", ENTRY_LOG_AP="entry", build_driver=lambda **_kwargs: object(),
                login=lambda *_args: None, quit_driver=lambda _driver: None,
                fill_entry_log_form_for_test=fill,
                query_visible_table=lambda *_args, **_kwargs: (
                    [["115/10/22", "08:00", "-", fills[0][0], "出", "外勤支援"]] if fills else []
                ),
            )
            action = {"kind": "entry_log", "time": "08:00", "date_offset": 1,
                      "actor": "8", "target": "21", "source": "外勤簽出", "staff_roc_date": "1151022",
                      "duplicate_key": "entry:2026-10-22:8:out:21:外勤支援",
                      "fields": {"出或入": "出", "領用事由及地點": "外勤支援", "系統寫入時間": "08:00"}}
            data = {"target_date": "1151021", "actions": [action],
                    "today": {"roc_date": "1151021", "staff": {"21": {"name": "今日人員"}}},
                    "tomorrow": {"roc_date": "1151022", "staff_complete": True, "staff": {"21": {"name": "明日人員"}}}}
            service = DutySubmissionService(
                Path(temp_dir), module_loader=lambda: automation,
                now_factory=lambda: datetime(2026, 10, 22, 8), sleeper=lambda _seconds: None,
            )
            request = DutySubmissionRequest("night-user", "fake", 0, data, trigger_type="manual")
            result = service.execute(request)
            self.assertEqual(result.status, "submitted")
            self.assertEqual(fills, [("明日人員", "1151022", True)])
            self.assertIn("明日人員", " ".join(result.comparison["matched"]))
            repeated = service.execute(request)
            self.assertEqual(repeated.status, "skipped_duplicate")
            self.assertEqual(len(fills), 1)

    def test_next_fire_day_external_missing_roster_stops_before_browser(self) -> None:
        duty_rehearsal_module()
        from app_core.duty_submission_service import (
            DutySubmissionRequest, DutySubmissionService, DutySubmissionValidationError,
        )

        action = {"kind": "entry_log", "time": "08:00", "date_offset": 1,
                  "actor": "8", "target": "21", "source": "外勤簽出", "staff_roc_date": "1151022",
                  "fields": {"出或入": "出", "領用事由及地點": "外勤支援"}}
        for tomorrow in (None, {"roc_date": "1151022", "staff": {}},
                         {"roc_date": "1151022", "staff": {"21": {"name": "部分名冊人員"}}},
                         {"roc_date": "1151022", "staff_complete": False, "staff": {"21": {"name": "部分名冊人員"}}},
                         {"roc_date": "1151023", "staff": {"21": {"name": "錯日人員"}}}):
            with self.subTest(tomorrow=tomorrow), tempfile.TemporaryDirectory() as temp_dir:
                loader = mock.Mock()
                service = DutySubmissionService(Path(temp_dir), module_loader=loader)
                data = {"target_date": "1151021", "actions": [action], "tomorrow": tomorrow,
                        "today": {"roc_date": "1151021", "staff": {"21": {"name": "今日同番號人員"}}}}
                with self.assertRaisesRegex(DutySubmissionValidationError, "名冊"):
                    service.execute(DutySubmissionRequest("night-user", "fake", 0, data, trigger_type="manual"))
                loader.assert_not_called()

    def test_next_fire_day_rebased_background_result_matches_only_its_preview(self) -> None:
        from dataclasses import replace

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        duty_rehearsal_module()
        from PySide6.QtWidgets import QApplication
        from app_core.duty_submission_service import DutySubmissionRequest
        from app_core.schedule_repository import ScheduleRepository
        from app_core.unreturned_return_queue import UnreturnedReturnQueue
        from qt_app.controllers.duty_controller import DutyController

        app = QApplication.instance() or QApplication(["next-fire-day-result"])
        with tempfile.TemporaryDirectory() as temp_dir:
            controller = DutyController(
                repository=ScheduleRepository(Path(temp_dir)),
                unreturned_return_queue=UnreturnedReturnQueue(Path(temp_dir)),
            )
            try:
                controller.set_session_context(1, "night-user")
                controller.set_actor_no("8")
                action = {"kind": "entry_log", "time": "08:00", "date_offset": 1,
                          "actor": "8", "target": "21", "source": "外勤簽出", "staff_roc_date": "1151022",
                          "duplicate_key": "entry:2026-10-22:8:out:21:外勤支援",
                          "fields": {"出或入": "出", "領用事由及地點": "外勤支援"}}
                controller.replace_schedule_data({"target_date": "1151021", "actions": [action]})
                request = DutySubmissionRequest(
                    "night-user", "fake", 0, {"target_date": "1151022", "actions": [{**action, "date_offset": 0}]},
                    trigger_type="manual", session_generation=1, session_actor_no="8", background=True,
                    schedule_generation=controller.schedule_generation, action_key=action["duplicate_key"],
                )
                invalid_requests = [
                    replace(request, background=False),
                    replace(request, session_generation=2),
                    replace(request, session_actor_no="9"),
                    replace(request, user_id="another-user"),
                    replace(request, schedule_data={"target_date": "1151023", "actions": [{**action, "date_offset": 0}]}),
                    replace(request, schedule_data={"target_date": "1151022", "actions": [{**action, "date_offset": 1}]}),
                    replace(request, schedule_data={"target_date": "1151022", "actions": [{**action, "date_offset": 0, "staff_roc_date": ""}]}),
                ]
                for invalid in invalid_requests:
                    with self.subTest(request=invalid):
                        self.assertFalse(controller.handle_submission_request_result(invalid, "submitted", "完成", "result.json"))
                        self.assertEqual(controller.ui_action_key_for_request(invalid), "")
                self.assertEqual(controller.ui_action_key_for_request(request), action["duplicate_key"])
                self.assertTrue(controller.handle_submission_request_result(request, "submitted", "完成", "result.json"))
                self.assertEqual(controller._executed_indices, {0})
                controller.replace_schedule_data({"target_date": "1151022", "actions": [{**action, "date_offset": 0, "staff_roc_date": ""}]})
                self.assertEqual(controller._executed_indices, {0})
                self.assertEqual(controller._comparisons[0]["group"], "done")
            finally:
                controller.shutdown()
        self.assertIsNotNone(app)

    def test_next_fire_day_preview_result_allows_only_bounded_scheduled_catchup(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        duty_rehearsal_module()
        from PySide6.QtWidgets import QApplication
        from app_core.duty_submission_service import DutySubmissionRequest
        from app_core.schedule_repository import ScheduleRepository
        from app_core.unreturned_return_queue import UnreturnedReturnQueue
        from qt_app.controllers.duty_controller import DutyController

        app = QApplication.instance() or QApplication(["next-fire-day-catchup"])
        with tempfile.TemporaryDirectory() as temp_dir:
            controller = DutyController(
                repository=ScheduleRepository(Path(temp_dir)),
                unreturned_return_queue=UnreturnedReturnQueue(Path(temp_dir)),
            )
            try:
                controller.set_session_context(1, "night-user")
                controller.set_actor_no("8")
                action = {"kind": "entry_log", "time": "08:00", "date_offset": 1,
                          "actor": "8", "target": "21", "source": "外勤簽出", "staff_roc_date": "1151022",
                          "duplicate_key": "entry:2026-10-22:8:out:21:外勤支援",
                          "fields": {"出或入": "出", "領用事由及地點": "外勤支援"}}
                controller.replace_schedule_data({"target_date": "1151021", "actions": [action]})
                for actual_time, scheduled_at, expected in (
                    ("08:01", "2026-10-22T08:00:00", True),
                    ("10:00", "2026-10-22T08:00:00", True),
                    ("07:59", "2026-10-22T08:00:00", False),
                    ("10:01", "2026-10-22T08:00:00", False),
                    ("08:01", "2026-10-22T09:00:00", False),
                    ("08:01", "invalid", False),
                    ("08:01", "", False),
                ):
                    with self.subTest(actual_time=actual_time, scheduled_at=scheduled_at):
                        submitted_action = {**action, "time": actual_time, "date_offset": 0,
                                            "scheduled_at": scheduled_at, "submit_target_date": "1151022",
                                            "fields": {**action["fields"], "系統寫入時間": actual_time}}
                        request = DutySubmissionRequest(
                            "night-user", "fake", 0, {"target_date": "1151022", "actions": [submitted_action]},
                            trigger_type="manual", session_generation=1, session_actor_no="8", background=True,
                            schedule_generation=controller.schedule_generation, action_key=action["duplicate_key"],
                        )
                        self.assertEqual(controller.handle_submission_request_result(request, "submitted", "完成", "result.json"), expected)
                        self.assertEqual(controller.ui_action_key_for_request(request), action["duplicate_key"] if expected else "")
            finally:
                controller.shutdown()
        self.assertIsNotNone(app)

    def test_next_fire_day_preview_return_cannot_submit_before_due_for_reused_actor_number(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        duty_rehearsal_module()
        from PySide6.QtTest import QSignalSpy
        from PySide6.QtWidgets import QApplication
        from app_core.schedule_repository import ScheduleRepository
        from app_core.unreturned_return_queue import UnreturnedReturnQueue
        from qt_app.controllers.duty_controller import DutyController

        app = QApplication.instance() or QApplication(["next-fire-day-early-return"])
        current = datetime(2026, 10, 21, 22)

        class Clock(datetime):
            @classmethod
            def now(cls, tz=None):
                return current.replace(tzinfo=tz)

        with tempfile.TemporaryDirectory() as temp_dir, mock.patch(
            "qt_app.controllers.duty_controller.datetime", Clock,
        ):
            controller = DutyController(
                repository=ScheduleRepository(Path(temp_dir)),
                unreturned_return_queue=UnreturnedReturnQueue(Path(temp_dir)),
            )
            try:
                controller.set_session_context(1, "night-user")
                controller.set_actor_no("8")
                departure = {"kind": "entry_log", "time": "08:00", "date_offset": 1,
                             "actor": "8", "target": "21", "source": "外勤簽出", "staff_roc_date": "1151022",
                             "duplicate_key": "next-day:departure", "return_pair_key": "next-day:pair",
                             "fields": {"出或入": "出", "領用事由及地點": "外勤支援"}}
                returning = {**departure, "time": "10:00", "source": "外勤簽入",
                             "duplicate_key": "next-day:return", "fields": {"出或入": "入", "領用事由及地點": "返隊"}}
                controller.replace_schedule_data({
                    "target_date": "1151021", "actions": [departure, returning],
                    "today": {"roc_date": "1151021", "staff": {"8": {"name": "夜班人員"}}},
                    "tomorrow": {"roc_date": "1151022", "staff_complete": True,
                                 "staff": {"8": {"name": "明日同番號人員"}, "21": {"name": "外勤人員"}}},
                })
                self.assertEqual({row["taskIndex"] for row in controller._task_model._tasks}, {0, 1})
                self.assertFalse(controller._is_manual_submission_available(1, now=current))
                self.assertEqual(controller.manual_submission_requests("night-user", "fake", [1], submit_at=current), [])
                submitted = QSignalSpy(controller.manualSubmissionRequested)
                waiting = QSignalSpy(controller.manualWaitingScheduled)
                controller.toggleTaskSelection(1)
                self.assertFalse(controller.canManualSubmitSelected)
                controller.prepareManualSubmission()
                controller.confirmManualSubmission()
                self.assertEqual(submitted.count(), 0)
                self.assertEqual(waiting.count(), 0)
                controller.toggleTaskSelection(1)
                controller.toggleTaskSelection(0)
                self.assertTrue(controller.canManualSubmitSelected)
                controller.prepareManualSubmission()
                controller.confirmManualSubmission()
                self.assertEqual(waiting.count(), 1)
                self.assertEqual(submitted.count(), 0)
                current = datetime(2026, 10, 22, 9, 59)
                self.assertFalse(controller._is_manual_submission_available(1, now=current))
                current = datetime(2026, 10, 22, 10)
                self.assertFalse(controller._is_manual_submission_available(1, now=current))
                self.assertEqual(controller.manual_submission_requests("night-user", "fake", [1], submit_at=current), [])
                controller.replace_schedule_data({
                    "target_date": "1151022", "actions": [{**returning, "date_offset": 0, "staff_roc_date": ""}],
                })
                self.assertTrue(controller._is_manual_submission_available(0, now=current))
            finally:
                controller.shutdown()
        self.assertIsNotNone(app)

    def test_next_morning_0800_handoff_preview_includes_work_log(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150628",
            rows=[module.DutyRow("22-24", {"值班": ["8"]})],
            summary={"在勤": ["8"]},
        )
        tomorrow = module.DutySheet(
            roc_date="1150629",
            rows=[module.DutyRow("08-10", {"值班": ["11"]})],
            summary={"在勤": ["11"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150628"), [], tomorrow)
        next_handoff_0800 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source, action.date_offset)
            for action in actions
            if action.source == "值班交接" and action.time == "08:00" and action.date_offset == 1
        ]

        self.assertEqual(
            next_handoff_0800,
            [
                ("entry_log", "08:00", "8", "8", "值退", "值班交接", 1),
                ("entry_log", "08:00", "8", "11", "值班", "值班交接", 1),
                ("work_log", "08:00", "8", "8", None, "值班交接", 1),
            ],
        )

    def test_0800_handoff_uses_previous_fire_day_not_today_0700(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150626",
            rows=[
                module.DutyRow("07-08", {"值班": ["27"]}),
                module.DutyRow("08-09", {"值班": ["15"]}),
            ],
            summary={"在勤": ["15", "26", "27"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150625",
            rows=[
                module.DutyRow("22-23", {"值班": ["26"]}),
                module.DutyRow("07-08", {"值班": ["26"]}),
            ],
            summary={"在勤": ["15", "26", "27"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150626"), [], None)
        handoff_0800 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "08:00"
        ]

        self.assertEqual(
            handoff_0800,
            [
                ("entry_log", "08:00", "26", "26", "值退", "值班交接"),
                ("entry_log", "08:00", "26", "15", "值班", "值班交接"),
                ("work_log", "08:00", "26", "26", None, "值班交接"),
            ],
        )

    def test_0800_handoff_uses_previous_day_0700_before_2200(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150625",
            rows=[module.DutyRow("08-09", {"值班": ["28"]})],
            summary={"在勤": ["4", "13", "28"]},
        )
        yesterday = module.DutySheet(
            roc_date="1150624",
            rows=[
                module.DutyRow("22-23", {"值班": ["13"]}),
                module.DutyRow("07-08", {"值班": ["4"]}),
            ],
            summary={"在勤": ["4", "13", "28"]},
        )

        actions = module.planned_actions(today, yesterday, [], module.parse_roc_date("1150625"), [], None)
        handoff_0800 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "08:00"
        ]

        self.assertEqual(
            handoff_0800,
            [
                ("entry_log", "08:00", "4", "4", "值退", "值班交接"),
                ("entry_log", "08:00", "4", "28", "值班", "值班交接"),
                ("work_log", "08:00", "4", "4", None, "值班交接"),
            ],
        )

    def test_handoff_uses_continuous_duty_segment_for_work_period(self) -> None:
        module = duty_rehearsal_module()
        today = module.DutySheet(
            roc_date="1150624",
            rows=[
                module.DutyRow("22-24", {"值班": ["13"]}),
                module.DutyRow("0-7", {"值班": ["13"]}),
                module.DutyRow("7-8", {"值班": ["4"]}),
            ],
            summary={"在勤": ["4", "13"]},
        )

        actions = module.planned_actions(today, None, [], module.parse_roc_date("1150624"), [], None)
        handoff_0700 = [
            (action.kind, action.time, action.actor, action.target, action.fields.get("出或入"), action.source)
            for action in actions
            if action.source == "值班交接" and action.time == "07:00"
        ]
        work_log = next(
            action
            for action in actions
            if action.source == "值班交接" and action.time == "07:00" and action.kind == "work_log"
        )

        self.assertEqual(
            handoff_0700,
            [
                ("entry_log", "07:00", "13", "13", "值退", "值班交接"),
                ("entry_log", "07:00", "13", "4", "值班", "值班交接"),
                ("work_log", "07:00", "13", "13", None, "值班交接"),
            ],
        )
        self.assertIn("時間:22-07", work_log.fields["處理情形"])

    def test_update_package_requests_logout_before_stopping_gui(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")

        self.assertIn("function Send-UpdateLogoutEvent", script)
        self.assertIn("update_prepare", script)
        self.assertIn("NamedPipeClientStream", script)
        self.assertIn("TYFD.SinpoSmart.DutyAutomation.Qt", script)
        install_start = script.index("$runningDutyGuiProcesses")
        install_section = script[
            install_start:script.index("Copy-UpdateTree -SourceDir", install_start)
        ]
        self.assertIn('$prepareResult -ne "ready"', install_section)
        self.assertIn("$runningDutyGuiProcesses.Count -ne 1", install_section)
        self.assertIn("Test-IsQtDutyGuiProcess -Process $runningQtProcess", install_section)
        self.assertIn("-ExpectedProcessId $handshakenProcessId", install_section)
        self.assertLess(
            install_section.index("$runningDutyGuiProcesses.Count -ne 1"),
            install_section.index("Wait-UpdateLogoutEvent"),
        )
        self.assertLess(
            install_section.index('$prepareResult -ne "ready"'),
            install_section.index("Stop-RunningDutyGui"),
        )

    def test_update_package_preserves_running_gui_without_ready_handshake(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        install_start = script.index("$runningDutyGuiProcesses")
        install_section = script[
            install_start:script.index("Copy-UpdateTree -SourceDir", install_start)
        ]

        self.assertIn('"busy"', script)
        self.assertIn('"failed"', script)
        self.assertIn('"timeout"', script)
        self.assertIn('$prepareResult -ne "ready"', install_section)
        self.assertIn("Stop-RunningDutyGui -Processes $runningDutyGuiProcesses -Ready", install_section)
        self.assertLess(
            install_section.index('$prepareResult -ne "ready"'),
            install_section.index("Stop-RunningDutyGui"),
        )

    def test_update_package_restarts_gui_when_installation_fails_after_shutdown(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")

        install_start = script.index("$runningDutyGuiProcesses")
        finally_start = script.index("} finally {", install_start)
        install_section = script[install_start:finally_start]
        finally_section = script[finally_start:]

        self.assertIn("$guiStoppedForUpdate = $false", script)
        self.assertIn("$guiRestarted = $false", script)
        self.assertIn("$guiStoppedForUpdate = $true", install_section)
        self.assertIn("$guiRestarted = [bool](Start-DutyGui)", install_section)
        self.assertIn("if ($guiStoppedForUpdate -and -not $guiRestarted)", finally_section)
        self.assertIn("$guiRestarted = [bool](Start-DutyGui)", finally_section)
        self.assertLess(
            finally_section.index("$guiRestarted = [bool](Start-DutyGui)"),
            finally_section.index("if (Test-Path -LiteralPath $tempDir)"),
        )

    def test_update_package_detects_relative_entrypoints_fail_closed(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        helper_start = script.index("function Get-DutyGuiEntrypointToken")
        helper_end = script.index("function Send-UpdateLogoutEvent", helper_start)
        helper_source = script[helper_start:helper_end]
        contract = f"""
$ErrorActionPreference = "Stop"
$packageDir = 'C:\\SinpoSmart Package'
{helper_source}
$script:FakeProcesses = @(
    [pscustomobject]@{{ ProcessId = 11; CommandLine = 'pythonw.exe duty_gui.pyw' }},
    [pscustomobject]@{{ ProcessId = 12; CommandLine = 'python.exe ..\\package\\duty_gui.py' }},
    [pscustomobject]@{{ ProcessId = 13; CommandLine = 'pythonw.exe "C:\\SinpoSmart Package\\duty_gui.pyw"' }},
    [pscustomobject]@{{ ProcessId = 14; CommandLine = 'pythonw.exe "C:\\Other Package\\duty_gui.pyw"' }}
)
function Get-CimInstance {{
    [CmdletBinding()]
    param([Parameter(Position = 0)][string]$ClassName)
    return $script:FakeProcesses
}}
$found = @(Get-RunningDutyGuiProcesses)
$ids = @($found | ForEach-Object {{ [int]$_.ProcessId }})
if ($found.Count -ne 3 -or 11 -notin $ids -or 12 -notin $ids -or 13 -notin $ids) {{
    throw "Relative or local duty_gui process was not detected. IDs: $($ids -join ',')"
}}
if (14 -in $ids) {{
    throw "A confirmed external absolute duty_gui path was incorrectly claimed."
}}
if (-not (Test-IsQtDutyGuiProcess -Process $found[0])) {{
    throw "The relative QML entrypoint was not classified as Qt."
}}
if (Test-IsQtDutyGuiProcess -Process $found[1]) {{
    throw "The relative Tk entrypoint was incorrectly classified as Qt."
}}
"ok"
"""

        result = run_powershell_contract(contract)

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ok", result.stdout)

    def test_update_package_accepts_handshaken_process_natural_exit(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        helper_start = script.index("function Get-DutyGuiEntrypointToken")
        helper_end = script.index("function Send-UpdateLogoutEvent", helper_start)
        stop_start = script.index("function Stop-RunningDutyGui")
        stop_end = script.index("function Start-DutyGui", stop_start)
        contract = f"""
$ErrorActionPreference = "Stop"
$packageDir = 'C:\\SinpoSmart Package'
{script[helper_start:helper_end]}
{script[stop_start:stop_end]}
$missingPid = [int]::MaxValue
$initial = [pscustomobject]@{{ ProcessId = $missingPid; CommandLine = 'pythonw.exe duty_gui.pyw' }}
function Get-RunningDutyGuiProcesses {{ return @() }}
if (-not (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $missingPid -Ready)) {{
    throw "A handshaken process that exited naturally was rejected."
}}
$liveInitial = [pscustomobject]@{{ ProcessId = $PID; CommandLine = 'pythonw.exe duty_gui.pyw' }}
if (Stop-RunningDutyGui -Processes @($liveInitial) -ExpectedProcessId $PID -Ready -ExitTimeoutSeconds 0) {{
    throw "An unclassified but still-live handshaken PID was accepted."
}}
function Get-RunningDutyGuiProcesses {{
    return @([pscustomobject]@{{ ProcessId = 42; CommandLine = 'pythonw.exe duty_gui.pyw' }})
}}
if (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $missingPid -Ready) {{
    throw "A replacement PID was accepted after the handshake."
}}
"ok"
"""

        result = run_powershell_contract(contract)

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ok", result.stdout)

    def test_update_package_waits_for_temporarily_unclassified_gui_to_exit(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        helpers = script[script.index("function Get-DutyGuiEntrypointToken"):script.index("function Send-UpdateLogoutEvent")]
        stop = script[script.index("function Stop-RunningDutyGui"):script.index("function Invoke-HiddenProcess")]
        contract = f"""
$ErrorActionPreference = "Stop"
$packageDir = 'C:/SinpoSmart Contract'
{helpers}
{stop}
function Get-RunningDutyGuiProcesses {{ return @() }}
function Stop-Process {{ throw "The updater must not force-stop the GUI." }}
$child = Start-Process -FilePath (Join-Path $PSHOME 'powershell.exe') -ArgumentList '-NoProfile -NonInteractive -Command Start-Sleep -Milliseconds 500' -WindowStyle Hidden -PassThru
try {{
    $initial = [pscustomobject]@{{ ProcessId = $child.Id; CommandLine = 'pythonw.exe duty_gui.pyw' }}
    if (-not (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $child.Id -Ready)) {{
        throw "A GUI still exiting after ready must be awaited, not rejected immediately."
    }}
    if (-not $child.HasExited) {{ throw "Update was accepted before process exit." }}
}} finally {{ $child.WaitForExit(); $child.Dispose() }}
"ok"
"""
        result = run_powershell_contract(contract)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ok", result.stdout)

    def test_update_package_close_failures_record_bounded_safe_reasons(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        helpers = script[script.index("function Get-DutyGuiEntrypointToken"):script.index("function Send-UpdateLogoutEvent")]
        stop = script[script.index("function Stop-RunningDutyGui"):script.index("function Invoke-HiddenProcess")]
        contract = f"""
$ErrorActionPreference = "Stop"
$packageDir = 'C:/SinpoSmart Contract'
{helpers}
{stop}
function Stop-Process {{ throw "The updater must not force-stop a process." }}
$initial = [pscustomobject]@{{ ProcessId = $PID; CommandLine = 'pythonw.exe duty_gui.pyw' }}
function Get-RunningDutyGuiProcesses {{ return @() }}
if (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $PID -Ready -ExitTimeoutSeconds 0) {{ throw "Live PID accepted." }}
if ($script:guiCloseFailure.code -ne 'gui_exit_timeout') {{ throw "Missing specific timeout diagnostic." }}
if ($script:guiCloseFailure.expected_pid -ne $PID) {{ throw "Missing original PID." }}
if (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $PID -ExitTimeoutSeconds 0) {{ throw "Missing handshake accepted." }}
if ($script:guiCloseFailure.code -ne 'handshake_not_ready') {{ throw "Missing handshake diagnostic." }}
function Get-RunningDutyGuiProcesses {{ return @([pscustomobject]@{{ ProcessId = 42; CommandLine = 'pythonw.exe duty_gui.pyw' }}) }}
if (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $PID -Ready -ExitTimeoutSeconds 0) {{ throw "Replacement GUI accepted." }}
if ($script:guiCloseFailure.code -ne 'gui_process_changed') {{ throw "Missing replacement diagnostic." }}
function Get-RunningDutyGuiProcesses {{ throw 'private-token=must-not-appear' }}
if (Stop-RunningDutyGui -Processes @($initial) -ExpectedProcessId $PID -Ready -ExitTimeoutSeconds 0) {{ throw "Failed query accepted." }}
if ($script:guiCloseFailure.code -ne 'gui_query_failed') {{ throw "Missing query diagnostic." }}
if (($script:guiCloseFailure | ConvertTo-Json) -match 'private-token|must-not-appear') {{ throw "Raw exception leaked." }}
if (-not $script:guiCloseFailure.exception_type) {{ throw "Missing safe exception type." }}
"ok"
"""
        result = run_powershell_contract(contract)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ok", result.stdout)

    def test_update_close_diagnostic_survives_failed_stage_manifest(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        setter = script[script.index("function Set-GuiCloseFailure"):script.index("function Invoke-HiddenProcess")]
        writer = script[script.index("function Write-RemoteStageManifest"):script.index("function Add-PendingUpdateNotice")]
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = str(Path(temp_dir)).replace("'", "''")
            contract = f"""
$ErrorActionPreference = "Stop"
$stagingDir = '{directory}'
$stagingManifestPath = Join-Path $stagingDir 'manifest.json'
$RequestId = '44444444-4444-4444-8444-444444444444'
$packageDir = 'C:/SinpoSmart Contract'
{setter}
{writer}
Set-GuiCloseFailure -Code 'gui_exit_timeout' -Reason 'Timed out waiting for GUI' -ExpectedProcessId 1234 -ObservedCount 0 -WaitedMs 60001
Write-RemoteStageManifest -Status 'failed' -RemoteVersion '2026.10.20.0200' -RemoteSha256 ('a' * 64) -InstalledVersion '2026.10.19.2240' -Detail $script:guiCloseFailure.detail
$record = Get-Content -LiteralPath $stagingManifestPath -Raw -Encoding UTF8 | ConvertFrom-Json
if ($record.close_failure.code -ne 'gui_exit_timeout' -or $record.close_failure.expected_pid -ne 1234 -or $record.close_failure.waited_ms -ne 60001) {{ throw 'Lost failure details in manifest.' }}
if ($record.detail -notmatch 'gui_exit_timeout') {{ throw 'Legacy GUI report lost the specific cause.' }}
Write-RemoteStageManifest -Status 'completed' -RemoteVersion '2026.10.20.0200' -RemoteSha256 ('a' * 64) -InstalledVersion '2026.10.20.0200'
$record = Get-Content -LiteralPath $stagingManifestPath -Raw -Encoding UTF8 | ConvertFrom-Json
if ($record.close_failure) {{ throw 'Stale failure contaminated a successful update.' }}
"ok"
"""
            result = run_powershell_contract(contract)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("ok", result.stdout)

    def test_operational_status_updates_share_lock_and_replace_atomically(self) -> None:
        module = package_module("app_core.operational_sync_service")

        with tempfile.TemporaryDirectory() as temp_dir:
            service = module.OperationalSyncService(Path(temp_dir))
            original_load = service._load_sync_status
            first_loaded = threading.Event()
            release_first = threading.Event()
            load_count = 0
            count_lock = threading.Lock()

            def controlled_load():
                nonlocal load_count
                current = original_load()
                with count_lock:
                    load_count += 1
                    current_count = load_count
                if current_count == 1:
                    first_loaded.set()
                    release_first.wait(2)
                return current

            service._load_sync_status = controlled_load
            event_thread = threading.Thread(
                target=service._record_sync_status,
                args=("event", "ok", "event complete"),
            )
            board_thread = threading.Thread(
                target=service._record_sync_status,
                args=("board", "ok", "board complete"),
            )
            event_thread.start()
            self.assertTrue(first_loaded.wait(1))
            board_thread.start()
            release_first.set()
            event_thread.join(2)
            board_thread.join(2)

            self.assertFalse(event_thread.is_alive())
            self.assertFalse(board_thread.is_alive())
            status = json.loads(service.status_path.read_text(encoding="utf-8"))
            self.assertEqual(status["event"]["state"], "ok")
            self.assertEqual(status["board"]["state"], "ok")
            self.assertEqual(list(service.status_path.parent.glob(f".{service.status_path.name}.*.tmp")), [])

    def test_durable_event_queue_does_not_wait_for_http_delivery(self) -> None:
        module = package_module("app_core.operational_sync_service")

        with tempfile.TemporaryDirectory() as temp_dir:
            service = module.OperationalSyncService(Path(temp_dir))
            payload = service.enqueue_event(
                "logout",
                trigger_type="update",
                actor_no="10",
                durable_only=True,
            )
            pending = service._load_pending_events()

        self.assertEqual([entry["event_id"] for entry in pending], [payload["event_id"]])

    def test_session_timeout_allows_retry_and_ignores_late_first_result(self) -> None:
        package_module("app_core.credential_repository")
        from app_core.credential_repository import CredentialRepository
        from app_core.login_verifier import LoginResult
        from app_core.session import SessionState
        from qt_app.controllers.session_controller import SessionController

        class FakeThread:
            def __init__(self) -> None:
                self.interruption_requested = False
                self.deleted = False

            def requestInterruption(self) -> None:
                self.interruption_requested = True

            def deleteLater(self) -> None:
                self.deleted = True

        with tempfile.TemporaryDirectory() as temp_dir:
            state = SessionState()
            repository = CredentialRepository(Path(temp_dir) / "saved_login.json", "SinpoSmart", None)
            controller = SessionController(state, repository=repository, verifier=object())
            attempt_id = state.begin_login()
            self.assertIsNotNone(attempt_id)
            controller._pending_credentials[attempt_id] = ("user10", "secret", False, "manual")
            thread = FakeThread()
            controller._login_workers[attempt_id] = (thread, object())

            controller._login_timed_out(attempt_id)

            self.assertFalse(controller.isBusy)
            self.assertTrue(controller.hasRunningLoginWorkers)
            self.assertTrue(thread.interruption_requested)
            controller._login_succeeded(
                attempt_id,
                LoginResult(actor_no="10", user_id="user10", actor_name="late"),
            )
            self.assertFalse(controller.isLoggedIn)
            self.assertIsNotNone(state.begin_login())

            controller._finalize_login_thread(attempt_id)
            self.assertFalse(controller.hasRunningLoginWorkers)
            self.assertTrue(thread.deleted)

        source = (
            package_dir() / "qt_app" / "controllers" / "session_controller.py"
        ).read_text(encoding="utf-8")
        tree = ast.parse(source)
        login_method = next(
            node
            for node in ast.walk(tree)
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == "login"
        )
        self.assertNotIn("_login_workers", ast.unparse(login_method.body[0].test))
        self.assertIn("_poll_login_thread_finished", source)
        self.assertIn("QTimer.singleShot", source)

    def test_tray_and_update_guards_reject_busy_stop_without_waiting(self) -> None:
        package_module("app_core.update_repository")
        from app_core.update_repository import UpdateRepository
        from qt_app.controllers.tray_controller import TrayController
        from qt_app.controllers.update_controller import UpdateController

        class FakeApp:
            def __init__(self) -> None:
                self.quit_calls = 0

            def quit(self) -> None:
                self.quit_calls += 1

        block_reason = "勤務登打仍在執行"
        fake_app = FakeApp()
        notices = []
        tray = TrayController(
            fake_app,
            tray_available=False,
            native_notifier=lambda title, message: notices.append((title, message)) or True,
            stop_guard=lambda: block_reason,
        )
        tray.requestQuit()
        self.assertEqual(fake_app.quit_calls, 0)
        self.assertFalse(tray.quitRequested)
        self.assertIn(block_reason, notices[0][1])

        with tempfile.TemporaryDirectory() as temp_dir:
            version_path = Path(temp_dir) / "VERSION.txt"
            version_path.write_text("2026.08.10.0001\n", encoding="utf-8")
            script_path = version_path.with_name("update_package.ps1")
            script_path.write_text("# updater\n", encoding="utf-8")
            launched = []
            update = UpdateController(
                UpdateRepository(version_path),
                process_launcher=lambda path: launched.append(path),
                stop_guard=lambda: block_reason,
            )
            update._update_available = True
            update.launchUpdate()

            self.assertEqual(launched, [])
            self.assertIn(block_reason, update.statusText)

    def test_owned_qt_workers_poll_after_wait_timeout_without_sticky_state(self) -> None:
        controller_root = package_dir() / "qt_app" / "controllers"

        for relative in (
            "app_controller.py",
            "session_controller.py",
            "update_controller.py",
            "duty_controller.py",
            "duty_execution_controller.py",
            "daily_vehicle_controller.py",
            "duty_sheet_controller.py",
            "rest_monthly_controller.py",
            "rescue_video_controller.py",
        ):
            with self.subTest(relative=relative):
                source = (controller_root / relative).read_text(encoding="utf-8")
                self.assertIn("if not thread.wait(5_000):", source)
                self.assertIn("thread.isFinished()", source)
                self.assertIn("QTimer.singleShot", source)

    def test_update_prepare_quits_only_after_ready_response(self) -> None:
        source = (package_dir() / "qt_app" / "main.py").read_text(encoding="utf-8")

        self.assertIn('command == "update_prepare"', source)
        self.assertIn('result not in {"ready", "busy", "failed"}', source)
        self.assertIn('quit_after_response = result == "ready"', source)
        self.assertIn("if quit_after_response:", source)

    def test_update_backup_includes_qt_application_directories(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")

        self.assertIn("$backupDirectories", script)
        self.assertIn('"app_core"', script)
        self.assertIn('"qt_app"', script)

    def test_update_package_requires_qt_windowed_entrypoint(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        start_section = script[
            script.index("function Start-DutyGui"):
            script.index("function Restart-DutyGuiIfRunning")
        ]

        self.assertIn('Join-Path $packageDir "duty_gui.pyw"', start_section)
        self.assertNotIn('Join-Path $packageDir "duty_gui.py"', start_section)
        self.assertNotIn("-Windowed", start_section)
        self.assertIn("CreateNoWindow = $true", start_section)
        self.assertIn("UseShellExecute = $false", start_section)
        self.assertIn('Join-Path $_.FullName "duty_gui.pyw"', script)

    def test_update_package_rejects_incomplete_qt_archive_before_install(self) -> None:
        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")

        self.assertIn("$requiredQtPackageFiles", script)
        for relative in (
            "duty_gui.pyw",
            "update_history.json",
            "qt_app\\main.py",
            "qt_app\\qml\\Main.qml",
            "qt_app\\qml\\components\\AppleButton.qml",
            "qt_app\\qml\\components\\AppleCheckBox.qml",
            "qt_app\\qml\\components\\AppleComboBox.qml",
            "qt_app\\qml\\components\\AppleDialog.qml",
            "qt_app\\qml\\components\\AppleTabButton.qml",
            "qt_app\\qml\\components\\AppleTextArea.qml",
            "qt_app\\qml\\components\\AppleTextField.qml",
            "qt_app\\qml\\components\\AuditSummaryCard.qml",
            "qt_app\\qml\\components\\DataSectionTitle.qml",
            "qt_app\\qml\\components\\DataTableCell.qml",
            "qt_app\\qml\\components\\DangerButton.qml",
            "qt_app\\qml\\components\\DutyActionButton.qml",
            "qt_app\\qml\\components\\DutyTaskCard.qml",
            "qt_app\\qml\\components\\DutyTaskStatusPill.qml",
            "qt_app\\qml\\components\\FormFieldTitle.qml",
            "qt_app\\qml\\components\\PrimaryButton.qml",
            "qt_app\\qml\\components\\SettingsButton.qml",
            "qt_app\\qml\\components\\StrongHeaderTitle.qml",
            "qt_app\\qml\\components\\ToolAddButton.qml",
            "qt_app\\qml\\components\\ToolBrowseButton.qml",
            "qt_app\\qml\\components\\ToolCloseButton.qml",
            "qt_app\\qml\\components\\ToolDateStepButton.qml",
            "qt_app\\qml\\components\\ToolFieldLabel.qml",
            "qt_app\\qml\\components\\ToolFormCard.qml",
            "qt_app\\qml\\components\\ToolMonthCombo.qml",
            "qt_app\\qml\\components\\ToolPanelContent.qml",
            "qt_app\\qml\\components\\ToolPanelHeader.qml",
            "qt_app\\qml\\components\\ToolPanelTitle.qml",
            "qt_app\\qml\\components\\ToolRemoveButton.qml",
            "qt_app\\qml\\components\\ToolRunButton.qml",
            "qt_app\\qml\\components\\ToolSectionTitle.qml",
            "qt_app\\qml\\components\\ToolSidePanel.qml",
            "qt_app\\qml\\components\\ToolStatusBar.qml",
            "qt_app\\qml\\components\\WorkLogValueControl.qml",
            "qt_app\\qml\\components\\qmldir",
            "qt_app\\qml\\dialogs\\AccountManagerWindow.qml",
            "qt_app\\qml\\dialogs\\RescueVideoWindow.qml",
            "qt_app\\qml\\dialogs\\ActionConfirmations.qml",
            "qt_app\\qml\\dialogs\\UpdateHistoryDialog.qml",
            "qt_app\\qml\\dialogs\\ErrorDetailDialog.qml",
            "qt_app\\qml\\dialogs\\qmldir",
            "qt_app\\qml\\pages\\DutySheetToolPanel.qml",
            "qt_app\\qml\\pages\\RestTimeToolPanel.qml",
            "qt_app\\qml\\pages\\MonthlyBaseToolPanel.qml",
            "qt_app\\qml\\pages\\DailyVehicleToolPanel.qml",
            "qt_app\\qml\\pages\\AuditFilterPanel.qml",
            "qt_app\\qml\\pages\\WorkLogSettingsPanel.qml",
            "qt_app\\qml\\pages\\DutyQuickToolsPanel.qml",
            "qt_app\\qml\\pages\\DutyOperationBar.qml",
            "qt_app\\qml\\pages\\DutyTaskArea.qml",
            "qt_app\\qml\\pages\\SessionHeader.qml",
            "qt_app\\qml\\pages\\qmldir",
            "qt_app\\qml\\styles\\Design.qml",
            "qt_app\\qml\\styles\\qmldir",
            "qt_app\\workers\\operational_sync_worker.py",
            "app_core\\operational_sync_service.py",
            "app_core\\credential_repository.py",
        ):
            with self.subTest(relative=relative):
                self.assertIn(f'"{relative}"', script)
        self.assertIn("Update zip is missing required PySide6/QML file", script)
        self.assertLess(
            script.index("$requiredQtPackageFiles"),
            script.index("$runningDutyGuiProcesses"),
        )

    def test_update_package_requires_notice_and_civilpower_runtime_files(self) -> None:
        import re

        script = (package_dir() / "update_package.ps1").read_text(encoding="utf-8-sig")
        required_block = script.split("$requiredQtPackageFiles = @(", 1)[1].split("\n    )", 1)[0]
        required_paths = {value.replace("\\", "/") for value in re.findall(r'"([^"\n]+)"', required_block)}

        for relative in (
            "app_core/notice_collection_service.py",
            "app_core/notice_collection_browser.py",
            "app_core/notice_collection_dom.js",
            "qt_app/controllers/notice_collection_controller.py",
            "app_core/civilpower_service.py",
            "app_core/civilpower_automation.py",
            "qt_app/controllers/civilpower_controller.py",
            "qt_app/qml/pages/CivilpowerToolPanel.qml",
        ):
            with self.subTest(relative=relative):
                self.assertIn(relative, required_paths)

    def test_powershell_scripts_parse(self) -> None:
        root = package_dir()
        scripts = [
            root / "find_winpython.ps1",
            root / "install_startup_shortcut.ps1",
            root / "remove_startup_shortcut.ps1",
            root / "update_package.ps1",
        ]
        parser = (
            "$path = $env:PS_PARSE_PATH; "
            "$tokens = $null; $errors = $null; "
            "[System.Management.Automation.Language.Parser]::ParseFile($path, [ref]$tokens, [ref]$errors) | Out-Null; "
            "if ($errors.Count) { $errors | ForEach-Object { Write-Error $_.Message }; exit 1 }"
        )

        for script in scripts:
            with self.subTest(script=str(script.relative_to(PROJECT_ROOT))):
                env = os.environ.copy()
                env["PS_PARSE_PATH"] = str(script)
                with tempfile.TemporaryFile(mode="w+b") as output:
                    result = subprocess.run(
                        ["powershell", "-NoProfile", "-Command", parser],
                        cwd=PROJECT_ROOT,
                        env=env,
                        stdin=subprocess.DEVNULL,
                        stdout=output,
                        stderr=subprocess.STDOUT,
                        timeout=30,
                    )
                    output.seek(0)
                    diagnostic = output.read().decode("utf-8", errors="replace")
                self.assertEqual(result.returncode, 0, diagnostic)

    def test_duty_board_payload_keeps_full_days_and_stable_hash(self) -> None:
        module = duty_gui_module()
        schedule = self.duty_board_schedule_payload()

        first = module.build_duty_board_payload(schedule)
        second = module.build_duty_board_payload(schedule)

        self.assertEqual(first["schema_version"], 1)
        self.assertEqual([day["roc_date"] for day in first["days"]], ["1150716", "1150717"])
        self.assertEqual(len(first["days"][0]["slots"]), 4)
        self.assertEqual(first["days"][0]["slots"][0]["duty_nos"], ["1", "2"])
        self.assertEqual(first["days"][0]["slots"][0]["names"], ["王小明", "李小華"])
        self.assertEqual(first["days"][0]["slots"][3]["start_hour"], 0)
        self.assertEqual(first["days"][0]["slots"][3]["end_hour"], 1)
        self.assertEqual(first["content_hash"], second["content_hash"])

    def test_post_duty_board_payload_requires_configuration_and_ok_response(self) -> None:
        module = duty_gui_module()
        payload = module.build_duty_board_payload(self.duty_board_schedule_payload())
        previous_url = module.DUTY_BOARD_SYNC_URL
        previous_key = module.DUTY_BOARD_SYNC_KEY
        module.DUTY_BOARD_SYNC_URL = ""
        module.DUTY_BOARD_SYNC_KEY = ""
        try:
            with self.assertRaisesRegex(RuntimeError, "看板同步"):
                module.post_duty_board_payload(payload)
        finally:
            module.DUTY_BOARD_SYNC_URL = previous_url
            module.DUTY_BOARD_SYNC_KEY = previous_key

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def read(self):
                return b'{"ok": true, "changed": false}'

        module.DUTY_BOARD_SYNC_URL = "https://example.invalid/exec"
        module.DUTY_BOARD_SYNC_KEY = "test-key"
        try:
            with mock.patch.object(module.urllib.request, "urlopen", return_value=Response()) as urlopen:
                result = module.post_duty_board_payload(payload)
            request = urlopen.call_args.args[0]
            sent = json.loads(request.data.decode("utf-8"))
            self.assertEqual(result, {"ok": True, "changed": False})
            self.assertEqual(sent["payload"]["content_hash"], payload["content_hash"])
            self.assertNotIn("test-key", request.full_url)
        finally:
            module.DUTY_BOARD_SYNC_URL = previous_url
            module.DUTY_BOARD_SYNC_KEY = previous_key

    def test_hourly_duty_board_sync_reuses_schedule_refresh_once(self) -> None:
        module = duty_gui_module()
        gui = object.__new__(module.DutyGui)
        gui.simple_mode = type("Mode", (), {"get": lambda self: True})()
        gui.session = module.LoginSession(actor_no="1", user_id="user", password="secret", verified=True)
        gui.duty_board_completed_hours = set()
        gui.snapshot_running = False
        scheduled = []
        refreshes = []
        gui.after = lambda delay, callback: scheduled.append((delay, callback)) or "after-id"
        gui.refresh_schedule_background = lambda target, label, target_dates=None: refreshes.append((target, label, target_dates)) or True
        with mock.patch.object(module, "datetime") as mocked_datetime:
            mocked_datetime.now.return_value = datetime(2026, 7, 16, 9, 2)
            mocked_datetime.side_effect = lambda *args, **kwargs: datetime(*args, **kwargs)
            gui.check_hourly_duty_board_sync()
            gui.check_hourly_duty_board_sync()
        self.assertEqual(len(refreshes), 1)
        self.assertEqual(refreshes[0][0], "1150716")
        self.assertEqual(len(scheduled), 2)


class DutyAmbulanceSplitTests(unittest.TestCase):
    def _payloads(self, standby, commander, external, ambulance=None):
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "24號"
        ambulance = ambulance if ambulance is not None else ["16", "28", "1"]
        for column, label in {3: "值班", 4: "外勤", 5: "休息", 6: "備勤緊急救護", 10: "備勤救災"}.items():
            sheet.cell(5, column, label)
        sheet.cell(6, 14, "指揮官")
        for row in range(10, 34):
            hour = (row - 2) % 24
            values = {2: f"{hour:02d}-{(hour + 1) % 24:02d}", 3: "20", 4: ",".join(external),
                5: "30", 6: ambulance[0], 7: ",".join(ambulance[1:]),
                11: ",".join(standby), 14: ",".join(commander)}
            for column, value in values.items():
                sheet.cell(row, column, value)
        normal_staff = list(dict.fromkeys(["20", "30"] + ambulance + standby + commander + external))
        driver = mock.Mock()
        driver.execute_script.return_value = False
        error = mock.Mock()
        def use_mock_browser(_open_browser, verify, **_kwargs):
            verify(driver)
            return driver
        with ExitStack() as stack:
            for name, options in {
                "ThreadPoolExecutor": {}, "log_status": {},
                "retry_duty_number_popup_preflight": {"side_effect": use_mock_browser},
                "preflight_duty_number_popup": {"return_value": "main"},
                "step_config_popups": {}, "wait_for_duty_result_grid": {"return_value": True},
                "step_navigate_to_task_table": {}, "super_js_execute": {"return_value": True},
                "trainee_numbers_from_workbook": {"return_value": set()},
                "expected_on_duty_numbers_from_daily_sheet": {"return_value": normal_staff},
                "build_driver": {"side_effect": AssertionError("測試不得啟動真實瀏覽器")},
            }.items():
                stack.enter_context(mock.patch.object(module, name, **options))
            stack.enter_context(mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook))
            stack.enter_context(mock.patch.object(module.time, "sleep"))
            fill_duty = stack.enter_context(mock.patch.object(module, "step_batch_fill_duty"))
            fill_mission = stack.enter_context(mock.patch.object(module, "step_fill_mission_cells",
                side_effect=RuntimeError("offline-stop-before-save")))
            result = module.start_automation("test-user", "test-password", "1151024", "test.xlsm", {},
                error_callback=error, show_dialogs=False)
        self.assertFalse(result)
        error.assert_called_once_with("offline-stop-before-save")
        return fill_duty.call_args.args[1], fill_mission.call_args.args[1], sheet

    def test_duty_and_mission_use_distinct_ambulance2_fallbacks(self):
        for standby, commander, external, expected_duty, expected_mission, remaining in (
            (["9", "8", "10"], ["2"], ["31"], "9,8", "9,8", "10,2"),
            (["8"], ["2"], ["31"], "8,2", "8,2", ""),
            ([], ["2"], ["31", "32"], "", "2,31", "2"),
        ):
            with self.subTest(standby=standby):
                duty, mission, sheet = self._payloads(standby, commander, external)
                for hour in range(24):
                    self.assertEqual(duty[f"_pln_{hour}_3"], "16,28,1")
                    self.assertEqual(duty[f"_pln_{hour}_4"], expected_duty)
                    self.assertEqual(duty[f"_pln_{hour}_5"], remaining)
                    self.assertEqual(duty[f"_pln_{hour}_6"], "30")
                    self.assertEqual(mission[f"_pln_{hour}_6"], "16,28,1")
                    self.assertEqual(mission[f"_pln_{hour}_7"], expected_mission)
                    all_duty_members = [no for column in range(1, 7)
                        for no in duty[f"_pln_{hour}_{column}"].split(",") if no]
                    self.assertEqual(len(all_duty_members), len(set(all_duty_members)))
                    self.assertTrue(set(mission[f"_pln_{hour}_6"].split(",")).isdisjoint(
                        mission[f"_pln_{hour}_7"].split(",")))
                    self.assertTrue(set(mission[f"_pln_{hour}_1"].split(",")).isdisjoint(
                        mission[f"_pln_{hour}_2"].split(",")))
                self.assertEqual(sheet.cell(10, 11).value, ",".join(standby))
                self.assertEqual(sheet.cell(10, 14).value, ",".join(commander))

    def test_duty_single_candidate_returns_to_standby_without_external_staff(self):
        module = legacy_duty_sheet_module()
        self.assertEqual(module.select_duty_ambulance2_members(["16"], ["2"], ["16", "28", "1"]), [])
        self.assertEqual(module.select_duty_ambulance2_members(["8", "8"], ["8"], ["16", "28", "1"]), [])
        self.assertEqual(module.select_duty_ambulance2_members(["9", "8"], ["2"], ["16", "28", "1"]), ["9", "8"])

    def test_mission_shortage_after_external_candidates_stops_before_writes(self):
        module = legacy_duty_sheet_module()
        workbook = module.openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "24號"
        for column, label in {3: "值班", 4: "休息", 5: "備勤緊急救護", 9: "備勤救災"}.items():
            sheet.cell(5, column, label)
        sheet.cell(6, 13, "指揮官")
        for row in range(10, 34):
            hour = (row - 2) % 24
            for column, value in {2: f"{hour:02d}-{(hour + 1) % 24:02d}", 3: "20", 4: "30",
                5: "16", 6: "28", 13: "2"}.items():
                sheet.cell(row, column, value)
        error = mock.Mock()
        with mock.patch.object(module.openpyxl, "load_workbook", return_value=workbook), mock.patch.object(
            module, "expected_on_duty_numbers_from_daily_sheet", return_value=["20", "30", "16", "28", "2"]
        ), mock.patch.object(module, "ThreadPoolExecutor", side_effect=AssertionError("不得啟動截圖")) as capture, mock.patch.object(
            module, "build_driver", side_effect=AssertionError("不得啟動瀏覽器")
        ) as browser, mock.patch.object(module, "log_status"):
            result = module.start_automation("test-user", "test-password", "1151024", "test.xlsm", {},
                error_callback=error, show_dialogs=False)
        self.assertFalse(result)
        self.assertIn("救災任務編組表救護2車", error.call_args.args[0])
        self.assertIn("外勤", error.call_args.args[0])
        capture.assert_not_called()
        browser.assert_not_called()

    def _schedule(self, split):
        module = duty_rehearsal_module()
        rows = []
        for offset in range(24):
            hour = (8 + offset) % 24
            columns = {"值班": ["10" if hour in (8, 9) else "11"], "救護": ["6", "7", "8"],
                "備勤": ["2", "3", "4"], "休息": [], "防溺車巡": []}
            if hour in (12, 13):
                columns["備勤"].remove("2")
                columns["休息"] = ["2"]
            if hour in (16, 17):
                columns["備勤"].remove("4")
                columns["防溺車巡"] = ["4"]
            if split:
                selected = columns["備勤"][:2] if len(columns["備勤"]) >= 2 else []
                columns["救護2車"] = selected
                columns["備勤"] = [no for no in columns["備勤"] if no not in selected]
            rows.append(module.DutyRow(f"{hour:02d}-{(hour + 1) % 24:02d}", columns))
        return module.DutySheet(roc_date="1151024", rows=rows,
            summary={"在勤": ["2", "3", "4", "6", "7", "8", "10", "11"]})

    def test_rescue_split_preserves_all_scheduled_actions(self):
        module = duty_rehearsal_module()
        with mock.patch.object(module, "load_work_log_defaults", return_value=dict(module.DEFAULT_WORK_LOG_DEFAULTS)):
            before = module.planned_actions(self._schedule(False), None, [], module.parse_roc_date("1151024"))
            after = module.planned_actions(self._schedule(True), None, [], module.parse_roc_date("1151024"))
        self.assertEqual([vars(action) for action in after], [vars(action) for action in before])
        self.assertFalse(any(action.source.startswith("外勤") and
            "救護2車" in action.duplicate_key for action in after))
        self.assertTrue(any(action.source == "外勤簽出" and action.target == "4" for action in after))
        training = [action for action in after if action.source == "在隊訓練"]
        self.assertEqual(len(training), 3)
        for action in training:
            expected = {"2", "3", "4", "6", "7", "8"}
            if action.time == "17:00":
                expected.remove("4")
            self.assertEqual(set(action.fields["服勤人員"]), expected)

    def test_rescue2_before_rest_is_not_an_external_route(self):
        module = duty_rehearsal_module()
        self.assertFalse(module.rest_is_external_route(self._schedule(True), "2", 12, 14))

    def _stale_ambulance_cache(self):
        module = duty_rehearsal_module()
        sheet = self._schedule(True)
        with mock.patch.object(module, "load_work_log_defaults", return_value=dict(module.DEFAULT_WORK_LOG_DEFAULTS)):
            actions = [asdict(action) for action in module.planned_actions(
                sheet, None, [], module.parse_roc_date(sheet.roc_date))]
        completed_key = next(action["duplicate_key"] for action in actions if action["source"] == "在隊訓練")
        for action in actions:
            if action["source"] == "在隊訓練":
                action["target"] = "6,7,8"
                action["fields"]["服勤人員"] = ["6", "7", "8"]
                action["audit_note"] = "保留既有紀錄"
        actions.append({"kind": "entry_log", "time": "08:00", "actor": "10", "target": "2",
            "fields": {"出或入": "出", "領用事由及地點": "救護2車"}, "source": "外勤簽出",
            "duplicate_key": "entry:2026-10-24:8:out:2:救護2車", "date_offset": 0})
        return {"target_date": sheet.roc_date, "today": asdict(sheet), "actions": actions,
            "completed_keys": [completed_key]}, completed_key

    def _assert_refreshed_cache(self, payload, completed_key):
        actions = payload["actions"]
        self.assertFalse(any(action.get("source", "").startswith("外勤") and
            "救護2車" in action.get("duplicate_key", "") for action in actions))
        training = next(action for action in actions if action.get("duplicate_key") == completed_key)
        self.assertEqual(set(training["fields"]["服勤人員"]), {"2", "3", "4", "6", "7", "8"})
        self.assertEqual(training["audit_note"], "保留既有紀錄")
        self.assertEqual(payload["completed_keys"], [completed_key])
        self.assertEqual(len({action["duplicate_key"] for action in actions}), len(actions))

    def test_repository_rebuilds_stale_rescue2_actions_without_overwriting_cache(self):
        duty_rehearsal_module()
        from app_core.schedule_repository import ScheduleRepository
        payload, completed_key = self._stale_ambulance_cache()
        with tempfile.TemporaryDirectory() as directory:
            repository = ScheduleRepository(Path(directory))
            path = repository.schedule_path(payload["target_date"])
            path.parent.mkdir(parents=True)
            original = json.dumps(payload, ensure_ascii=False)
            path.write_text(original, encoding="utf-8")
            snapshot = repository.load_for_date(payload["target_date"])
            self._assert_refreshed_cache(snapshot.data, completed_key)
            self.assertEqual(path.read_text(encoding="utf-8"), original)
            legacy_payload = {"target_date": payload["target_date"],
                "today": asdict(self._schedule(False)), "actions": []}
            path.write_text(json.dumps(legacy_payload, ensure_ascii=False), encoding="utf-8")
            with mock.patch("builtins.__import__", wraps=__import__) as imports:
                legacy_snapshot = repository.load_for_date(payload["target_date"])
            self.assertEqual(legacy_snapshot.data["actions"], [])
            self.assertFalse(any(call.args[0] == "duty_rehearsal" for call in imports.call_args_list))

    def test_legacy_cache_refresh_removes_false_departures_and_keeps_completion_keys(self):
        module = duty_gui_module()
        payload, completed_key = self._stale_ambulance_cache()
        module.DutyGui.ensure_schedule_actions(object(), payload)
        self._assert_refreshed_cache(payload, completed_key)

    def test_duty_fill_explicitly_clears_an_unstaffed_rescue2_cell(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("Node.js is required for the isolated browser-script contract")
        module = legacy_duty_sheet_module()
        driver = mock.Mock()
        module.step_batch_fill_duty(driver, {"_pln_8_4": ""})
        script, payload = driver.execute_script.call_args.args
        harness = """
const cell = {value:'2', focus(){}, dispatchEvent(){}};
global.window = {top:{document:{getElementById(){return cell;}}, frames:[], alert(){}, confirm(){}}};
global.Event = class Event {};
const count = new Function(SCRIPT)(PAYLOAD);
if (cell.value !== '' || count !== 1) throw new Error('Unstaffed ambulance cell was not cleared');
""".replace("SCRIPT", json.dumps(script)).replace("PAYLOAD", json.dumps(payload))
        result = subprocess.run([node, "-e", harness], capture_output=True, text=True, check=False)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_query_reads_old_and_split_headers_and_rejects_missing_header(self):
        from selenium.common.exceptions import JavascriptException
        node = shutil.which("node")
        if not node:
            self.skipTest("Node.js is required for the isolated browser-script contract")
        module = duty_rehearsal_module()
        for split, missing_header in ((False, False), (True, False), (True, True)):
            with self.subTest(split=split, missing_header=missing_header):
                header = ["時段", "值班", "外勤1", "救護"] + (["救護2車"] if split else []) + ["備勤", "休息", "檢核欄"]
                matrix = [] if missing_header else [header]
                for offset in range(24):
                    hour = (8 + offset) % 24
                    matrix.append([f"{hour}-{(hour + 1) % 24}", "10", "11", "6,7"]
                        + (["2,3"] if split else []) + ["4" if split else "2,3,4", "12", ""])
                def execute(script, *args):
                    harness = """
const trs = MATRIX.map(values => ({children:values.map(value => ({innerText:value, querySelectorAll(){return [];}}))}));
const table = {className:'report_list1', querySelectorAll(){return trs;}};
global.document = {body:{innerText:''}, querySelectorAll(){return [table];},
  getElementById(id){return id === '_txtTaskDate' ? {value:'1151024'} : null;}};
global.window = {getComputedStyle(){return {display:'block'};}};
console.log(JSON.stringify(new Function(SCRIPT)(TARGET)));
""".replace("MATRIX", json.dumps(matrix)).replace("SCRIPT", json.dumps(script)).replace("TARGET", json.dumps(args[0]))
                    result = subprocess.run([node, "-e", harness], capture_output=True, text=True,
                        encoding="utf-8", check=False)
                    if result.returncode:
                        raise JavascriptException(result.stderr)
                    return json.loads(result.stdout)
                driver = mock.Mock()
                driver.execute_script.side_effect = execute
                with mock.patch.object(module, "open_ap"), mock.patch.object(module, "WebDriverWait"), mock.patch.object(
                    module, "suppress_window_open_for_background_query"
                ), mock.patch.object(module, "js_set", return_value=True), mock.patch.object(
                    module, "arm_query_completion_tracking", return_value=100.0
                ), mock.patch.object(module, "js_click", return_value=True):
                    if missing_header:
                        with self.assertRaisesRegex(JavascriptException, "勤務表缺少.*標題"):
                            module.query_duty_sheet(driver, "1151024")
                    else:
                        sheet = module.query_duty_sheet(driver, "1151024")
                        self.assertEqual(len(sheet.rows), 24)
                        self.assertEqual(sheet.rows[0].columns["救護"], ["6", "7"])
                        self.assertEqual(sheet.rows[0].columns["休息"], ["12"])
                        if split:
                            self.assertEqual(sheet.rows[0].columns["救護2車"], ["2", "3"])


if __name__ == "__main__":
    unittest.main()
