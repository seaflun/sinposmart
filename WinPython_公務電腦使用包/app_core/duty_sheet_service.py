# -*- coding: utf-8 -*-
"""UI-independent boundary for the legacy duty-sheet automation engine."""

from __future__ import annotations

import importlib.util
import json
import os
import re
import sys
import threading
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from types import ModuleType
from typing import Callable, Iterator

from app_core.duty_task_projection import parse_roc_date
from app_core.schedule_repository import business_roc_date


LEGACY_MODULE_NAME = "_sinposmart_qt_duty_sheet_automation"
LEGACY_SCRIPT_NAME = "sinposmart_1.py"
DUTY_SHEET_WORKDIR_LOCK = threading.Lock()


class DutySheetValidationError(ValueError):
    """A safe request-validation message for the native form."""


class DutySheetExecutionError(RuntimeError):
    """A safe execution failure for the native form."""

    def __init__(self, message: str, *, failure_stage: str = "unknown") -> None:
        super().__init__(message)
        self.failure_stage = failure_stage


@dataclass(frozen=True)
class DutySheetDefaults:
    workbook_path: str
    target_date: str
    attack: str
    stop: str
    amb1: str
    amb2: str
    attack_options: tuple[str, ...]
    stop_options: tuple[str, ...]
    amb_options: tuple[str, ...]
    notification_enabled: bool


@dataclass(frozen=True)
class WorkbookSignature:
    size: int
    modified_ns: int


@dataclass(frozen=True)
class DutySheetRequest:
    user_id: str
    password: str = field(repr=False)
    workbook_path: str
    target_date: str
    attack: str
    stop: str
    amb1: str
    amb2: str
    notification_enabled: bool
    automatic: bool = False
    workbook_signature: WorkbookSignature | None = None

    @property
    def cars_config(self) -> dict[str, str]:
        return {
            "attack": self.attack,
            "stop": self.stop,
            "amb1": self.amb1,
            "amb2": self.amb2,
            "workbook_path": self.workbook_path,
        }


def load_legacy_module(project_dir: Path) -> ModuleType:
    script_path = project_dir / LEGACY_SCRIPT_NAME
    source_mtime = script_path.stat().st_mtime
    existing = sys.modules.get(LEGACY_MODULE_NAME)
    if (
        existing is not None
        and getattr(existing, "__sinposmart_source_path__", None) == str(script_path)
        and getattr(existing, "__sinposmart_source_mtime__", None) == source_mtime
    ):
        return existing
    sys.modules.pop(LEGACY_MODULE_NAME, None)
    importlib.invalidate_caches()
    spec = importlib.util.spec_from_file_location(LEGACY_MODULE_NAME, script_path)
    if spec is None or spec.loader is None:
        raise DutySheetExecutionError("勤務表自動化模組無法載入。")
    module = importlib.util.module_from_spec(spec)
    sys.modules[LEGACY_MODULE_NAME] = module
    sys.path.insert(0, str(project_dir))
    try:
        spec.loader.exec_module(module)
    finally:
        try:
            sys.path.remove(str(project_dir))
        except ValueError:
            pass
    module.__sinposmart_source_path__ = str(script_path)
    module.__sinposmart_source_mtime__ = source_mtime
    return module


@contextmanager
def legacy_workdir(project_dir: Path) -> Iterator[None]:
    previous = Path.cwd()
    with DUTY_SHEET_WORKDIR_LOCK:
        try:
            os.chdir(project_dir)
            yield
        finally:
            os.chdir(previous)


class DutySheetService:
    def __init__(
        self,
        package_root: Path,
        *,
        module_loader: Callable[[Path], ModuleType] = load_legacy_module,
    ) -> None:
        self.package_root = Path(package_root)
        self.project_dir = self.package_root / "duty_sheet_legacy"
        self.config_path = self.project_dir / "config.json"
        self.module_loader = module_loader

    def load_defaults(self, now: datetime | None = None) -> DutySheetDefaults:
        config = self._load_legacy_display_config()
        last = config.get("last_selection", {}) if isinstance(config.get("last_selection"), dict) else {}
        options = config.get("car_options", {}) if isinstance(config.get("car_options"), dict) else {}
        notification = config.get("notification", {}) if isinstance(config.get("notification"), dict) else {}
        workbook = self._resolve_workbook(str(last.get("workbook_path", "") or ""))
        target = parse_roc_date(business_roc_date(now or datetime.now())) + timedelta(days=1)
        return DutySheetDefaults(
            workbook_path=str(workbook) if workbook else "",
            target_date=target.strftime("%Y/%m/%d"),
            attack=str(last.get("attack", "") or ""),
            stop=str(last.get("stop", "") or ""),
            amb1=str(last.get("amb1", "") or ""),
            amb2=str(last.get("amb2", "") or ""),
            attack_options=self._string_options(options.get("attack", [])),
            stop_options=self._string_options(options.get("stop", [])),
            amb_options=self._string_options(options.get("amb", [])),
            notification_enabled=bool(notification.get("enabled", False)),
        )

    @staticmethod
    def _has_complete_saved_display_config(config: dict) -> bool:
        last = config.get("last_selection")
        options = config.get("car_options")
        notification = config.get("notification")
        return (
            isinstance(last, dict)
            and all(key in last for key in ("attack", "stop", "amb1", "amb2"))
            and isinstance(options, dict)
            and all(isinstance(options.get(key), list) for key in ("attack", "stop", "amb"))
            and isinstance(notification, dict)
        )

    def _load_legacy_display_config(self) -> dict:
        """Read complete saved display settings without loading the legacy GUI on first use."""

        saved_config = self._read_config()
        if self._has_complete_saved_display_config(saved_config):
            return saved_config

        try:
            legacy = self.module_loader(self.project_dir)
            with legacy_workdir(self.project_dir):
                config = legacy.load_config()
        except Exception:
            return self._read_config()
        return config if isinstance(config, dict) else self._read_config()

    def add_vehicle_option(self, group: str, code: str, plate: str) -> str:
        group = str(group or "").strip()
        code = str(code or "").strip()
        plate = str(plate or "").strip()
        if group not in ("attack", "stop", "amb"):
            raise DutySheetValidationError("車輛類型不正確。")
        if not code or not plate:
            raise DutySheetValidationError("請輸入車輛代號與車牌號碼。")

        value = f"{code}/{plate}"
        config = self._read_config(for_update=True)
        options = config.setdefault("car_options", {})
        hidden_options = config.setdefault("hidden_car_options", {})
        values = options.setdefault(group, [])
        hidden_values = hidden_options.setdefault(group, [])
        if value in hidden_values:
            hidden_values.remove(value)
        if value not in values:
            values.append(value)
        self._write_config(config)
        return value

    def remove_vehicle_option(self, group: str, value: str) -> str:
        group = str(group or "").strip()
        value = str(value or "").strip()
        if group not in ("attack", "stop", "amb"):
            raise DutySheetValidationError("車輛類型不正確。")
        if not value:
            raise DutySheetValidationError("請選擇要移除的車輛。")

        config = self._read_config(for_update=True)
        options = config.setdefault("car_options", {})
        hidden_options = config.setdefault("hidden_car_options", {})
        values = options.setdefault(group, [])
        if value not in values:
            raise DutySheetValidationError(f"車輛清單中沒有：{value}")
        values.remove(value)
        hidden_values = hidden_options.setdefault(group, [])
        if value not in hidden_values:
            hidden_values.append(value)

        fallback = str(values[0]) if values else ""
        last = config.setdefault("last_selection", {})
        fields = (
            ("attack",)
            if group == "attack"
            else ("stop",)
            if group == "stop"
            else ("amb1", "amb2")
        )
        for field_name in fields:
            if str(last.get(field_name, "") or "").strip() == value:
                last[field_name] = fallback
        self._write_config(config)
        return value

    def validate(self, request: DutySheetRequest) -> DutySheetRequest:
        if not request.user_id.strip() or not request.password:
            raise DutySheetValidationError("請先完成勤務系統登入。")
        workbook = Path(request.workbook_path) if request.automatic else self._resolve_workbook(request.workbook_path)
        if workbook is None or not workbook.is_file() or workbook.suffix.lower() not in (".xlsx", ".xlsm"):
            raise DutySheetValidationError("請選擇有效的勤務表 Excel 檔案（.xlsx 或 .xlsm）。")
        try:
            datetime.strptime(request.target_date.strip(), "%Y/%m/%d")
        except ValueError as exc:
            raise DutySheetValidationError("日期格式必須是 YYYY/MM/DD。") from exc
        missing = [
            label
            for label, value in (
                ("攻擊車", request.attack),
                ("指揮車", request.stop),
                ("救護車 1", request.amb1),
                ("救護車 2", request.amb2),
            )
            if not str(value or "").strip()
        ]
        if missing:
            raise DutySheetValidationError(f"請選擇：{'、'.join(missing)}。")
        return DutySheetRequest(
            user_id=request.user_id.strip(),
            password=request.password,
            workbook_path=str(workbook),
            target_date=request.target_date.strip(),
            attack=request.attack.strip(),
            stop=request.stop.strip(),
            amb1=request.amb1.strip(),
            amb2=request.amb2.strip(),
            notification_enabled=bool(request.notification_enabled),
            automatic=request.automatic,
            workbook_signature=request.workbook_signature,
        )

    def prepare_automatic_request(self, user_id: str, password: str, target_date: str) -> DutySheetRequest:
        """Runs in a worker: use only the existing saved selection, never a bundled sample."""
        config = self._read_config(for_update=True)
        last = config.get("last_selection", {})
        if not isinstance(last, dict):
            raise DutySheetValidationError("請先手動選取一次勤務表。")
        workbook = self.resolve_automatic_workbook(str(last.get("workbook_path", "") or ""), target_date)
        stat = workbook.stat()
        return self.validate(DutySheetRequest(
            user_id, password, str(workbook), parse_roc_date(target_date).strftime("%Y/%m/%d"),
            *(str(last.get(key, "") or "") for key in ("attack", "stop", "amb1", "amb2")),
            notification_enabled=bool(config.get("notification", {}).get("enabled", False)),
            automatic=True, workbook_signature=WorkbookSignature(stat.st_size, stat.st_mtime_ns),
        ))

    def resolve_automatic_workbook(self, saved_path: str, target_date: str) -> Path:
        target = parse_roc_date(target_date)
        if not saved_path.strip():
            raise DutySheetValidationError("尚無前一次選取的勤務表，請手動選取檔案。")
        saved = Path(saved_path)
        if not saved.is_absolute():
            saved = self.project_dir / saved
        candidates = []

        def inspect(path: Path) -> bool:
            try:
                stat = path.stat()
                self._validate_automatic_workbook(path, target)
                if path.stat().st_mtime_ns != stat.st_mtime_ns:
                    raise DutySheetValidationError("檔案正在修改")
            except Exception as exc:
                reason = str(exc) if isinstance(exc, DutySheetValidationError) else "檔案無法讀取"
                candidates.append({"file": str(path), "valid": False, "reason": reason})
                return False
            candidates.append({"file": str(path), "valid": True, "modified_ns": stat.st_mtime_ns})
            return True

        if inspect(saved):
            self._record_workbook_selection(target_date, candidates, saved)
            return saved
        month_match = re.fullmatch(r"0?(\d{1,2})月", saved.parent.name)
        year_match = re.fullmatch(r"(\d{3,4})(年)?", saved.parent.parent.name)
        if month_match and year_match:
            suffix = year_match.group(2) or ""
            root = saved.parent.parent.parent
            folder = root / f"{target.year - 1911}{suffix}" / f"{target.month}月"
            if not folder.is_dir():
                padded = root / f"{target.year - 1911}{suffix}" / f"{target.month:02d}月"
                folder = padded if padded.is_dir() else folder
        else:
            folder = saved.parent
        try:
            paths = sorted(folder.iterdir(), key=lambda path: path.name)
        except OSError as exc:
            self._record_workbook_selection(target_date, candidates, None)
            raise DutySheetValidationError("勤務表資料夾無法讀取，請確認網路路徑或手動選檔。") from exc
        for path in paths:
            if path == saved or not self._automatic_filename_allowed(path):
                continue
            if "新坡" in path.stem and "勤務表" in path.stem:
                inspect(path)
        valid = sorted((row for row in candidates if row["valid"]), key=lambda row: row["modified_ns"], reverse=True)
        selected = None
        error = "找不到年月、日期分頁及必要欄位皆正確的勤務表，請手動選取。"
        if valid:
            if len(valid) > 1 and valid[0]["modified_ns"] == valid[1]["modified_ns"]:
                error = "多份勤務表修改時間相同，請手動選取，未自行猜測。"
            else:
                selected = Path(valid[0]["file"])
        self._record_workbook_selection(target_date, candidates, selected)
        if selected is None:
            raise DutySheetValidationError(error)
        return selected

    @staticmethod
    def _automatic_filename_allowed(path: Path) -> bool:
        return (path.suffix.lower() == ".xlsm" and not path.name.startswith("~$")
                and not re.search(r"備份|舊版|副本|複製|暫存|backup|copy|old|temp", path.stem, re.I))

    def _validate_automatic_workbook(self, path: Path, target: datetime) -> None:
        from openpyxl import load_workbook
        if not self._automatic_filename_allowed(path):
            raise DutySheetValidationError("不是可用的 .xlsm 勤務表或屬於備份檔")
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            try:
                title = f"{target.day}號"
                if title not in workbook.sheetnames:
                    raise DutySheetValidationError("缺少目標日期分頁")
                sheet = workbook[title]
                header = list(sheet.iter_rows(min_row=1, max_row=6, max_col=99, values_only=True))
                texts = [re.sub(r"\s+", "", str(value or "")) for row in header[4:6] for value in row]
                required = ("值班", "休息", "備勤緊急救護", "備勤救災", "指揮官")
                if any(not any(label in text for text in texts) for label in required):
                    raise DutySheetValidationError("缺少勤務表必要欄位")
                periods = []
                for text in [path.stem] + [str(value or "") for row in header[:4] for value in row]:
                    for match in re.finditer(r"(?<!\d)(\d{3,4})年\s*(\d{1,2})月", text):
                        year, month = map(int, match.groups())
                        periods.append((year + 1911 if year < 1911 else year, month))
                for row in header[:4]:
                    for value in row:
                        if isinstance(value, datetime):
                            periods.append((value.year, value.month))
                year_folder = re.fullmatch(r"(\d{3,4})年?", path.parent.parent.name)
                month_folder = re.fullmatch(r"(\d{1,2})月", path.parent.name)
                if year_folder and month_folder:
                    year = int(year_folder.group(1))
                    periods.append((year + 1911 if year < 1911 else year, int(month_folder.group(1))))
                if not periods or any(period != (target.year, target.month) for period in periods):
                    raise DutySheetValidationError("無法確認目標年月或年月不符")
            finally:
                workbook.close()

    def _record_workbook_selection(self, target_date: str, candidates: list[dict], selected: Path | None) -> None:
        path = self.package_root / "runtime_outputs" / "daily_workbook_selection.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps({
            "target_date": target_date, "checked_at": datetime.now().isoformat(),
            "candidates": candidates, "selected": str(selected) if selected else "",
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temporary, path)

    def confirmation_summary(self, request: DutySheetRequest) -> str:
        return (
            f"日期：{request.target_date}\n"
            f"Excel：{Path(request.workbook_path).name}\n"
            f"車輛：{request.attack}、{request.stop}、{request.amb1}、{request.amb2}\n\n"
            "確認後將登入正式勤務系統並執行勤務表登打。"
        )

    def execute(
        self,
        request: DutySheetRequest,
        *,
        status_callback: Callable[[str], None] | None = None,
        stage_callback: Callable[[str], None] | None = None,
    ) -> str:
        stage = "preflight"

        def report_stage(value: str) -> None:
            nonlocal stage
            stage = value
            if stage_callback is not None:
                stage_callback(value)

        report_stage(stage)
        request = self.validate(request)
        if request.automatic:
            stat = Path(request.workbook_path).stat()
            if request.workbook_signature != WorkbookSignature(stat.st_size, stat.st_mtime_ns):
                raise DutySheetValidationError("勤務表選取後已變更，請人工確認後重新登打。")
        if not (self.project_dir / LEGACY_SCRIPT_NAME).is_file():
            raise DutySheetExecutionError("找不到勤務表自動化模組。")
        report_stage("config_load")
        legacy: ModuleType | None = self.module_loader(self.project_dir)
        errors: list[str] = []
        successes: list[str] = []
        try:
            with legacy_workdir(self.project_dir):
                current_config = legacy.load_config()
                notification = dict(current_config.get("notification", {}))
                notification["enabled"] = request.notification_enabled
                if not request.automatic:
                    legacy.save_config(
                        request.cars_config,
                        login_settings=dict(current_config.get("login", {})),
                        notification_settings=notification,
                        car_options=current_config.get("car_options", {}),
                        hidden_car_options=current_config.get("hidden_car_options", {}),
                    )
                selected_date = datetime.strptime(request.target_date, "%Y/%m/%d")
                target_date = legacy.convert_to_minguo(selected_date)
                result = legacy.start_automation(
                    request.user_id,
                    request.password,
                    target_date,
                    request.workbook_path,
                    request.cars_config,
                    status_callback=status_callback,
                    success_callback=successes.append,
                    error_callback=errors.append,
                    show_dialogs=False,
                    close_driver=True,
                    keep_browser_open_on_success=True,
                    raise_errors=True,
                    stage_callback=report_stage,
                )
        except DutySheetValidationError:
            raise
        except Exception as exc:
            raise DutySheetExecutionError(
                str(exc) or "勤務表登打失敗。",
                failure_stage=stage,
            ) from exc
        finally:
            if legacy is not None and hasattr(legacy, "_runtime_status_callback"):
                legacy._runtime_status_callback = None
        if result is not True:
            raise DutySheetExecutionError(
                errors[-1] if errors else "勤務表登打未完成。",
                failure_stage=stage,
            )
        selection_warning = ""
        if request.automatic:
            try:
                with legacy_workdir(self.project_dir):
                    legacy.save_config(request.cars_config)
            except Exception:
                selection_warning = "（登打成功，但前一次檔案設定未能保存）"
        legacy_result = successes[-1] if successes else ""
        if "勤務表截圖保存失敗" in legacy_result:
            suffix = "（截圖保存失敗）"
        elif request.notification_enabled and "LINE 截圖通知失敗" in legacy_result:
            suffix = "（LINE 截圖通知失敗）"
        else:
            suffix = ""
        return f"勤務表已登打完成：{target_date}{suffix}{selection_warning}"

    def _resolve_workbook(self, value: str) -> Path | None:
        value = str(value or "").strip()
        if value:
            candidate = Path(value).expanduser()
            if not candidate.is_absolute():
                candidate = self.project_dir / candidate
            if candidate.is_file():
                return candidate.resolve()
        return next((path.resolve() for path in self.project_dir.glob("*.xlsm")), None)

    def _read_config(self, *, for_update: bool = False) -> dict:
        if not self.config_path.is_file():
            return {}
        try:
            loaded = json.loads(self.config_path.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            if for_update:
                raise DutySheetValidationError("勤務表設定檔無法讀取，未覆寫原檔。") from exc
            return {}
        if isinstance(loaded, dict):
            return loaded
        if for_update:
            raise DutySheetValidationError("勤務表設定檔格式不正確，未覆寫原檔。")
        return {}

    def _write_config(self, config: dict) -> None:
        self.config_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.config_path.with_suffix(f"{self.config_path.suffix}.tmp")
        try:
            temporary.write_text(
                json.dumps(config, ensure_ascii=False, indent=4),
                encoding="utf-8",
            )
            temporary.replace(self.config_path)
        except OSError as exc:
            try:
                temporary.unlink()
            except OSError:
                pass
            raise DutySheetValidationError("勤務表設定檔無法儲存。") from exc

    @staticmethod
    def _string_options(values) -> tuple[str, ...]:
        if not isinstance(values, list):
            return ()
        return tuple(str(value) for value in values if str(value).strip())
